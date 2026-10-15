"""Сквозная проверка в настоящем окне программы (pywebview).

Окно открывается так же, как в `desktop.run`, страницу нажимают через JavaScript, файлы выбираются и
сохраняются через тот же мост, что использует пользователь. Отличие одно: системные окна «Открыть» и
«Сохранить как» подменяются заранее заданными путями, потому что автоматика не может нажимать в них.
Сами системные окна этой проверкой не покрываются.

Запуск: .venv/bin/python tools/e2e_window.py [--shots ПАПКА] [--theme light|dark]
Каждый файл из case/ проходит путь: выбрать → Обезличить → Сохранить как → выбрать сохранённый →
Восстановить → Сохранить как. Результат сверяется с исходником по ячейкам, абзацам и фигурам.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import webview  # noqa: E402

from anonymizer import desktop, helper  # noqa: E402
from anonymizer.vault import Vault  # noqa: E402
from anonymizer.webapp import create_server  # noqa: E402
from tests.helpers import content_units, docx_bytes, docx_text, xlsx_bytes  # noqa: E402

CASE = ROOT / "case"
DONE_TEXT = ("Готово", "Файл готов", "Файл для ИИ готов")


class Driver:
    def __init__(self, shots: Path | None, dark: bool, inproc: bool = False):
        self.shots = shots
        self.dark = dark
        self.home = Path(tempfile.mkdtemp(prefix="anon-e2e-"))
        os.environ["ANONYMIZER_HOME"] = str(self.home)
        self.out = self.home / "saved"
        self.out.mkdir()
        self.queue: list[list[str]] = []
        self.dialog_log: list[tuple] = []
        self.results: list[tuple[str, bool, str]] = []
        self.server, self.app = create_server(Vault(self.home / "vault.bin"), self.home / "work")
        self.port = self.server.server_address[1]
        threading.Thread(target=self.server.serve_forever, daemon=True).start()
        self.holder: dict = {}
        # По умолчанию мост говорит с фоном через RemoteApp, как настоящее окно; --inproc берёт App напрямую.
        self.remote = None
        if not inproc:
            self.client = helper.HelperClient({"port": self.port, "token": self.app.token}, timeout=60)
            self.remote = helper.RemoteApp(self.client)
        self.bridge = desktop.Bridge(self.remote or self.app, lambda: self.holder["window"])
        self.bridge._dialog = self._fake_dialog        # подмена только системных окон

    def _fake_dialog(self, kind, **kwargs):
        self.dialog_log.append((kind, kwargs.get("save_filename", ""), kwargs.get("file_types", ())))
        if not self.queue:
            return []
        answer = self.queue.pop(0)
        return answer

    # -- помощники -----------------------------------------------------------

    def js(self, code: str):
        return self.holder["window"].evaluate_js(code)

    def wait_js(self, condition: str, timeout: float = 90):
        end = time.time() + timeout
        while time.time() < end:
            try:
                if self.js(f"Boolean({condition})"):
                    return True
            except Exception:
                pass
            time.sleep(.3)
        return False

    def check(self, name: str, ok: bool, detail: str = "") -> None:
        self.results.append((name, bool(ok), detail))
        print(("OK   " if ok else "FAIL ") + name + (f" — {detail}" if detail and not ok else ""), flush=True)

    def shot(self, name: str) -> None:
        """Снимок содержимого окна средствами WKWebView (macOS): разрешение на запись экрана не нужно."""
        if not self.shots:
            return
        self.shots.mkdir(parents=True, exist_ok=True)
        try:
            import AppKit
            from PyObjCTools import AppHelper
            from webview.platforms.cocoa import BrowserView
            view = BrowserView.instances[self.holder["window"].uid].webview
            done = threading.Event()

            def handler(image, error):
                try:
                    if image is not None:
                        rep = AppKit.NSBitmapImageRep.imageRepWithData_(image.TIFFRepresentation())
                        data = rep.representationUsingType_properties_(AppKit.NSBitmapImageFileTypePNG, None)
                        data.writeToFile_atomically_(str(self.shots / f"{name}.png"), True)
                finally:
                    done.set()

            AppHelper.callAfter(lambda: view.takeSnapshotWithConfiguration_completionHandler_(None, handler))
            done.wait(10)
        except Exception as exc:
            print("снимок не сделан:", repr(exc))

    def click(self, selector: str) -> None:
        self.js(f"document.querySelector({json.dumps(selector)}).click()")

    def text(self, selector: str) -> str:
        return self.js(f"(document.querySelector({json.dumps(selector)})||{{innerText:''}}).innerText") or ""

    # -- сценарий ------------------------------------------------------------

    def run(self) -> None:
        window = self.holder["window"]
        if not self.wait_js("document.readyState==='complete' && window.pywebview && window.pywebview.api && window.pywebview.api.pick_files"):
            self.check("окно загрузило страницу и мост", False)
            return self.finish()
        self.check("окно загрузило страницу и мост", True)
        self.check("заголовок окна", self.js("document.title") == "Anonymizer")
        self.shot("01-start")
        files = sorted(p for ext in ("xlsx", "docx", "pptx") for p in CASE.glob(f"*.{ext}") if not p.name.startswith("~$"))
        if not getattr(self, "only_project", False):
            for source in files:
                self.one_file(source)
            self.multi(files[:2])
            self.extras()
            self.drop_scenario()
        if not getattr(self, "only_project", False):
            self.chat_features()
        self.project_flow()
        self.finish()

    def one_file(self, source: Path) -> None:
        label = source.name[:40]
        self.js("modes.chat.files.length = 0; modes.chat.folder = null; startDraft(); document.querySelector('#result-chat').innerHTML='';")
        self.queue = [[str(source)]]
        self.click("#pick-chat")
        ok = self.wait_js("modes.chat.files.length===1 && modes.chat.files[0].state==='ready'", 30)
        self.check(f"[{label}] выбор файла через мост", ok)
        if not ok:
            return
        self.click("#btn-chat-anon")
        ok = self.wait_js("document.querySelector('#result-chat .status')", 240)
        status = self.text("#result-chat .status")
        self.check(f"[{label}] обезличивание завершено", ok and status.startswith(DONE_TEXT), status)
        self.shot("02-result-" + re.sub(r"\W", "_", source.stem)[:20])
        target = self.out / f"anon_{source.name}"
        ok = self.save_version("#btn-result-save", target)
        self.check(f"[{label}] сохранение кнопкой «Сохранить файл для ИИ…»", ok and target.exists())
        if not target.exists():
            return
        # Возврат: тот же файл, как будто его вернул ИИ.
        pass
        self.queue = [[str(target)]]
        self.click("#pick-chat")
        ok = self.wait_js("modes.chat.files.length===1 && modes.chat.files[0].state==='ready'", 90)
        state = self.js("JSON.stringify(modes.chat.files.map(f => [f.name, f.state, f.info && f.info.tokens, f.error]))")
        self.check(f"[{label}] выбор обезличенного файла для возврата", ok and self.js("modes.chat.files[0].info.tokens") > 0, state)
        self.check(f"[{label}] главное действие для обезличенного файла «Восстановить»", self.js("document.querySelector('#btn-chat-rest').classList.contains('main')"))
        self.click("#btn-chat-rest")
        ok = self.wait_js("document.querySelector('#result-chat .status')", 240)
        self.check(f"[{label}] возврат завершён", ok, self.text("#result-chat .status"))
        back = self.out / f"back_{source.name}"
        ok = self.save_version("#btn-result-save", back)
        self.check(f"[{label}] сохранение восстановленного файла", ok and back.exists())
        if back.exists():
            original, restored = content_units(source), content_units(back)
            diff = [k for k in original if original[k] != restored.get(k)]
            self.check(f"[{label}] после круга все {len(original)} единиц совпали с исходником", not diff and len(original) == len(restored), f"{len(diff)} отличий: {diff[:3]}")

    def multi(self, files: list[Path]) -> None:
        self.js("modes.chat.files.length = 0; modes.chat.folder = null; startDraft();")
        self.queue = [[str(p) for p in files]]
        self.click("#pick-chat")
        ok = self.wait_js("modes.chat.files.length===2 && modes.chat.files.every(f=>f.state==='ready')", 30)
        self.check("[несколько файлов] выбор двух файлов", ok)
        self.click("#btn-chat-anon")
        ok = self.wait_js("document.querySelector('#result-chat .status')", 240)
        self.check("[несколько файлов] задание завершено", ok)
        folder = self.out / "many"
        folder.mkdir()
        both = all(self.save_version(f"#feed li.step:last-child li:nth-child({n}) [data-savefile][data-role='anon']", folder / f"f{n}.xlsx") for n in (1, 2))
        self.check("[несколько файлов] выгрузка обеих версий", both and len(list(folder.iterdir())) == 2, str(list(folder.iterdir())))
        # Отмена диалога не должна ломать страницу (после U2 «Сохранить как…» спрятано в меню «…», в тексте ленты видна кнопка «Сохранить»).
        self.save_version("#feed li.step:last-child li:nth-child(2) [data-savefile][data-role='anon']", None)
        time.sleep(1)
        self.check("[отмена диалога] страница жива", "Сохранить" in self.text("#feed"))
        self.shot("03-multi")

    # -- дополнительные сценарии ---------------------------------------------

    def pick(self, mode: str, path: Path, expect_ready: bool = True) -> bool:
        self.js(f"modes.chat.files.length = 0; modes.chat.folder = null; startDraft();")
        self.queue = [[str(path)]]
        self.click('#pick-chat')
        return self.wait_js(f"modes.chat.files.length===1 && modes.chat.files[0].state==='ready'", 30)

    def run_and_save(self, mode: str, target: Path) -> str:
        self.click("#btn-chat-" + mode)
        self.wait_js(f"document.querySelector('#result-chat .status')", 240)
        status = self.text(f"#result-chat .status")
        if target and self.js("Boolean(document.querySelector('#feed li.step:last-child [data-savefile]'))"):
            role = "anon" if mode == "anon" else "restored"
            self.save_version(f"#feed li.step:last-child [data-savefile][data-role='{role}']", target)
        return status

    def save_version(self, selector: str, target: Path | None) -> bool:
        """Нажимает «Выгрузить файл…» у версии в ленте; системное окно подменено заданным путём."""
        self.js("document.querySelectorAll('.toast').forEach(t => t.remove())")
        self.queue = [[str(target)]] if target else [[]]
        self.click(selector)
        return self.wait_js("[...document.querySelectorAll('.toast')].some(t => t.textContent.startsWith('Сохранено'))", 30 if target else 2)

    def ask(self, password: str | None) -> None:
        """Ответ в собственном диалоге страницы (пароль или подтверждение)."""
        self.wait_js("document.querySelector('#ask').open", 10)
        if password is not None:
            self.js(f"document.querySelector('#ask-input').value = {json.dumps(password)}")
        self.click("#ask-ok")
        self.wait_js("!document.querySelector('#ask').open", 10)

    def drop_scenario(self) -> None:
        """Перетаскивание: файл, брошенный на окно (не только на рамку), добавляется, страница остаётся на месте."""
        self.js("modes.chat.files.length = 0; modes.chat.folder = null; startDraft();")
        before = self.js("location.href")
        self.js("""(() => {
          const dt = new DataTransfer();
          dt.items.add(new File(["Директор Смирнов Алексей Петрович подписал акт."], "перетащено.txt", {type: "text/plain"}));
          const ev = new DragEvent("drop", {dataTransfer: dt, bubbles: true, cancelable: true});
          window.__dropPrevented = !document.body.dispatchEvent(ev);
        })()""")
        ok = self.wait_js("modes.chat.files.length===1 && modes.chat.files[0].state==='ready'", 15)
        self.check("[перетаскивание] файл, брошенный на окно, добавлен в список", ok)
        self.check("[перетаскивание] страница не ушла на документ", self.js("location.href") == before and self.js("window.__dropPrevented") is True)
        self.click("#btn-chat-anon")
        self.wait_js("document.querySelector('#result-chat .status')", 60)
        self.check("[перетаскивание] перетащенный файл обезличивается", self.text("#result-chat .status").startswith(DONE_TEXT), self.text("#result-chat .status"))

    # -- чаты и таблица замен ------------------------------------------------

    def add_file(self, path: Path) -> bool:
        """Добавляет файл в открытый чат (черновик не создаётся)."""
        self.js("modes.chat.files.length = 0; modes.chat.folder = null; renderFiles('chat');")
        self.queue = [[str(path)]]
        self.click("#pick-chat")
        return self.wait_js("modes.chat.files.length===1 && modes.chat.files[0].state==='ready'", 30)

    def run_action(self, mode: str) -> str:
        self.js("document.querySelector('#result-chat').innerHTML = ''")
        self.click("#btn-chat-" + mode)
        self.wait_js("!modes.chat.running && document.querySelector('#result-chat .status')", 120)
        return self.text("#result-chat .status")

    def blob(self, role: str, index: int = -1) -> str:
        cid = self.js("chatId")
        chat = self.app.chats.get(cid)
        data = self.app.chats.read_blob(cid, chat["files"][index]["id"], role)
        return data.decode("utf-8") if data else ""

    def act(self, base: str, name: str, variant: int = 0, answer: str | None = "") -> None:
        """Действие над кодом в таблице: вкладка «Что вернётся» для правки значения, «Что скрывается» для остального."""
        if name == "edit":
            self.js("setTab('back')"); time.sleep(.3)
            sel = f'input[data-edit="{base}"][data-variant="{variant or 1}"]'
            self.js(f"""(() => {{ const i = document.querySelector('{sel}'); i.focus(); i.value = {json.dumps(answer)};
                i.dispatchEvent(new KeyboardEvent('keydown', {{key: 'Enter', bubbles: true, cancelable: true}})); }})()""")
            time.sleep(1.2)
            return
        self.js("setTab('hide')"); time.sleep(.3)
        self.js(f"""(() => {{ const b = document.querySelector('[data-toggle="{base}"]'); if (b.getAttribute('aria-expanded') !== 'true') b.click(); }})()""")
        time.sleep(.3)
        if name in ("keep", "hide"):
            self.click(f'[data-switch="{base}"]'); time.sleep(1.2); return
        sel = f'[data-act="{name}"][data-base="{base}"]' + (f'[data-variant="{variant}"]' if variant else "")
        if name == "forget":
            self.js("setTab('back')"); time.sleep(.3); sel = f'[data-act="forget"][data-base="{base}"]'
        self.click(sel)
        if name == "merge":
            self.wait_js("document.querySelector('#picker').open", 10)
            self.click("#picker-list li")
            self.wait_js("!document.querySelector('#picker').open", 10)
        elif answer is not None and name in ("split", "forget"):
            self.wait_js("document.querySelector('#ask').open", 10)
            self.click("#ask-ok")
            self.wait_js("!document.querySelector('#ask').open", 10)
        time.sleep(1.2)

    def chat_features(self) -> None:
        vault = self.app.vault
        src = self.out / "caseA.txt"
        src.write_text("Договор подписал Пётр Сидоров. Позвонить Петру Сидорову. Директор Анна Кузнецова согласна.", encoding="utf-8")
        self.js("startDraft()")
        self.check("[чат] файл добавлен", self.add_file(src))
        status = self.run_action("anon")
        self.check("[чат] обезличивание в новом чате", status.startswith(DONE_TEXT), status)
        cid = self.js("chatId")
        anon = self.blob("anon")
        self.check("[чат] имена в обезличенной версии заменены", "Сидоров" not in anon and "Кузнецов" not in anon, anon)
        self.check("[чат] название чата из имени файла", self.js("chatData.title") == "caseA", self.js("chatData.title"))
        self.check("[чат] есть одна подсказка «Скрыть ещё слова»", self.js("document.querySelectorAll('#extra-terms').length") == 1 and self.js("!document.querySelector('#hide-once')"))
        saved = self.out / "caseA-anon.txt"
        self.check("[чат] сохранение версии для ИИ", self.save_version("#btn-result-save", saved) and saved.exists())

        # возврат внутри чата
        self.check("[возврат в чате] файл от ИИ добавлен", self.add_file(saved))
        status = self.run_action("rest")
        self.check("[возврат в чате] выполнен в том же чате", self.js("chatId") == cid and self.js("chatData.steps[chatData.steps.length-1].kind") == "restore", status)
        restored = self.blob("restored")
        self.check("[возврат в чате] значения вернулись", "Пётр Сидоров" in restored and "Анна Кузнецова" in restored, restored)

        # таблица: показать значения, поиск
        self.click("#btn-table")
        self.wait_js("paneOpen", 5)
        self.js("setTab('hide')")
        self.check("[таблица] значения скрыты по умолчанию", "Сидоров" not in self.text("#table-body"))
        self.click("#btn-reveal")
        time.sleep(.5)
        self.check("[таблица] «Показать значения» раскрывает и запоминается", "Сидоров" in self.text("#table-body") and self.text("#btn-reveal") == "Скрыть значения")
        self.shot("10-table")
        self.js("document.querySelector('#pane-search').value='Кузнец'; document.querySelector('#pane-search').dispatchEvent(new Event('input'))")
        time.sleep(.4)
        self.check("[таблица] поиск оставляет одну строку", self.js("document.querySelectorAll('.ent-row').length") == 1)
        self.js("document.querySelector('#pane-search').value=''; document.querySelector('#pane-search').dispatchEvent(new Event('input'))")
        persons = self.js("JSON.stringify(chatData.table.filter(r => r.kind === 'PERSON').map(r => [r.base, r.variants.length]))")
        persons = json.loads(persons)
        first, second = persons[0][0], persons[1][0]

        # изменить значение, возврат старого файла, отмена
        self.act(first, "edit", 1, "Пётр Сидоров-Ларин")
        self.check("[таблица] правка значения отмечена", "изменено" in self.text("#pane-note"), self.text("#pane-note"))
        self.add_file(saved)
        self.run_action("rest")
        self.check("[таблица] после правки возврат даёт новое значение", "Сидоров-Ларин" in self.blob("restored"), self.blob("restored"))
        self.click("#btn-undo")
        time.sleep(1)
        self.add_file(saved)
        self.run_action("rest")
        self.check("[таблица] отмена правки возвращает прежнее значение", "Сидоров-Ларин" not in self.blob("restored") and "Пётр Сидоров" in self.blob("restored"))

        # объединить: старый файл возвращается дословно, новые обезличивания используют общий код
        before = restored
        self.act(second, "merge", answer="")
        self.check("[объединение] коды объединены в хранилище", vault.describe(second)["base"] == first, str(vault.describe(second)["base"]))
        self.add_file(saved)
        self.run_action("rest")
        self.check("[объединение] старый файл возвращается дословно", self.blob("restored") == before)
        self.add_file(src)
        self.run_action("anon")
        fresh = self.blob("anon")
        self.check("[объединение] новое обезличивание выдаёт общий код", second not in fresh and first in fresh, fresh)
        self.click("#btn-undo") if self.js("Boolean(document.querySelector('#btn-undo'))") else None

        # разделить написание
        count = vault.stats()["entities"]
        row = json.loads(self.js(f"JSON.stringify(chatData.table.find(r => r.base === '{first}').variants.length)")) if False else self.js(f"chatData.table.find(r => r.base === '{first}').variants.length")
        if row and row > 1:
            self.act(first, "split", 2)
            self.check("[разделение] написание получило свой код", vault.stats()["entities"] == count + 1)
            self.add_file(saved)
            self.run_action("rest")
            self.check("[разделение] старый файл возвращается дословно", self.blob("restored") == before)

        # не скрывать, скрывать снова, применить правки к файлам чата
        self.act(first, "keep", 1)
        mine = self.app.chats.terms(self.js("chatId"))
        self.check("[не скрывать] решение сохранено только в этом чате, общий список пуст", any("Сидоров" in t for t in mine["keep_terms"]) and not vault.prefs["keep_terms"], str((mine, vault.prefs["keep_terms"])))
        self.check("[не скрывать] показан блок применения правок", self.js("!document.querySelector('#pane-apply').hidden"))
        self.act(first, "hide", 1)
        mine = self.app.chats.terms(self.js("chatId"))
        self.check("[скрывать снова] слово переехало в «Всегда скрывать» этого чата, общие списки пусты", any("Сидоров" in t for t in mine["hide_terms"]) and not mine["keep_terms"] and not vault.prefs["hide_terms"] and not vault.prefs["keep_terms"], str((mine, vault.prefs["hide_terms"])))
        self.act(first, "keep", 1)
        self.js("document.querySelector('#btn-apply').click()")
        self.wait_js("document.querySelector('#apply-state').innerText === 'Файл обновлён'", 90)
        rebuilt = self.blob("anon", 0)
        self.check("[применить правки] слово больше не заменяется", "Пётр Сидоров" in rebuilt, rebuilt)
        self.shot("11-after-apply")

        # ручка ширины должна быть под настоящим указателем (проверка попадания, а не посланных событий)
        self.js("setSide(true, {persist: false}); setPane(false); window.scrollTo(0, 0)")
        time.sleep(.6)
        hit = json.loads(self.js("""(() => { const e = document.getElementById('split-side'), r = e.getBoundingClientRect(); const h = document.elementFromPoint(r.x + r.width / 2, 300); return JSON.stringify([h && h.id, Math.round(r.x), Math.round(document.getElementById('side').getBoundingClientRect().right)]); })()"""))
        self.check("[ширина] ручка списка чатов стоит на краю списка и ловит указатель", hit[0] == "split-side" and abs(hit[1] + 8 - hit[2]) <= 1, str(hit))

        layers = json.loads(self.js("""(() => { const g = s => document.querySelector(s).getBoundingClientRect(), r = s => parseFloat(getComputedStyle(document.querySelector(s)).borderTopLeftRadius);
          const rail = g('#rail'), card = g('#side .side-in'), main = g('main.main'), av = g('#btn-profile');
          return JSON.stringify({railW: Math.round(rail.width), gap1: Math.round(main.left - (card.right)), cardR: r('#side .side-in'), mainR: r('main.main'), avRound: getComputedStyle(document.querySelector('#btn-profile .avatar')).borderRadius,
            avInRail: av.left >= rail.left && av.right <= rail.right + 1 && av.bottom > rail.bottom - 60, bodyScroll: document.documentElement.scrollHeight <= innerHeight + 1}); })()"""))
        self.check("[слои] колонка значков, карточка списка и карточка чата со скруглением, зазор 8, профиль круглый внизу колонки", layers["railW"] == 56 and layers["gap1"] == 8 and layers["cardR"] >= 16 and layers["mainR"] >= 16 and layers["avInRail"] and layers["bodyScroll"], str(layers))
        self.click("#btn-profile"); self.click("#open-settings")
        self.wait_js("document.querySelector('#settings').open", 10)
        tabs = self.js("document.querySelectorAll('.set-tab').length")
        self.js("document.querySelector('.set-tab[data-sec=\"storage\"]').click()")
        shown = self.js("[...document.querySelectorAll('.set-sec')].filter(x => !x.hidden).map(x => x.dataset.sec).join()")
        self.check("[настройки] страница с разделами: семь вкладок, переключение показывает один раздел", tabs == 7 and shown == "storage", f"{tabs} {shown}")
        self.js("document.querySelector('#btn-settings-close').click()")
        self.wait_js("!document.querySelector('#settings').open", 10)
        self.holder["window"].resize(1280, 800)
        time.sleep(.9)
        # окно в фоне не двигает CSS-переходы: для проверки положения они выключаются
        self.js("(() => { const s = document.createElement('style'); s.textContent = '*,*::before,*::after{transition:none!important;animation:none!important}'; document.head.appendChild(s); })()")
        self.js("setPane(true)")
        time.sleep(.8)
        hit = json.loads(self.js("""(() => { const e = document.getElementById('split-pane'), r = e.getBoundingClientRect(); const h = document.elementFromPoint(r.x + r.width / 2, 300); return JSON.stringify([h && h.id, Math.round(r.x), Math.round(document.getElementById('pane').getBoundingClientRect().left)]); })()"""))
        self.check("[ширина] ручка таблицы стоит на краю таблицы и ловит указатель", hit[0] == "split-pane" and abs(hit[1] - hit[2]) <= 1, str(hit))

        # добавить слово, которое не скрылось, прямо из таблицы
        self.js("document.querySelector('#btn-pane-add').click()")
        self.wait_js("document.querySelector('#ask').open", 10)
        self.js("document.querySelector('#ask-input').value = 'Секретный проект Зефир'")
        self.click("#ask-ok")
        self.wait_js("!document.querySelector('#ask').open", 10)
        self.wait_js("document.querySelector('#pane-note') && !document.querySelector('#pane-note').hidden", 10)
        mine = self.app.chats.terms(self.js("chatId"))
        self.check("[таблица] слово добавлено из таблицы и записано в «Всегда скрывать» этого чата, не в общие", any("Зефир" in t for t in mine["hide_terms"]) and not any("Зефир" in t for t in vault.prefs["hide_terms"]), str((mine, vault.prefs["hide_terms"])))
        self.js("setTab('hide')")
        self.check("[таблица] после добавления предложена пересборка", self.js("!document.querySelector('#pane-apply').hidden"))

        # удалить код: возврат оставляет код в файле и сообщает
        other = self.out / "caseB.txt"
        other.write_text("Письмо для Валерия Громова. Валерий Громов подтвердил.", encoding="utf-8")
        self.js("startDraft()")
        self.add_file(other)
        self.run_action("anon")
        cid2 = self.js("chatId")
        self.check("[чаты] второй чат создан, список из двух", cid2 != cid and self.js("chatList.length") >= 2)
        anon2 = self.out / "caseB-anon.txt"
        self.save_version("#btn-result-save", anon2)
        base = self.js("chatData.table.filter(r => r.kind === 'PERSON')[0].base")
        self.click("#btn-table")
        self.wait_js("paneOpen", 5)
        self.act(base, "forget")
        self.add_file(anon2)
        status = self.run_action("rest")
        self.check("[удалить код] возврат оставляет код и сообщает", base in self.blob("restored") and "Валерия Громова" not in self.blob("restored") and status.startswith("Файл восстановлен не полностью"), f"{status} | {self.blob('restored')}")

        # клавиатура и доступность
        self.js("document.querySelector('#btn-pane-close').click()")
        self.js("document.querySelector('#btn-table').focus()")
        self.click("#btn-table")
        self.js("document.dispatchEvent(new KeyboardEvent('keydown', {key: 'Escape', bubbles: true}))")
        time.sleep(.3)
        self.check("[клавиатура] Escape закрывает таблицу и возвращает фокус на кнопку", (not self.js("paneOpen")) and self.js("document.activeElement.id") == "btn-table")
        self.click("#btn-table")
        time.sleep(.3)
        self.js("setTab('hide')"); time.sleep(.3)
        self.check("[клавиатура] строки и действия таблицы — настоящие кнопки с подписями", self.js("[...document.querySelectorAll('#table-body [data-toggle], #table-body [data-act]')].every(b => b.tagName === 'BUTTON')") and self.js("document.querySelector('.ent-main').getAttribute('aria-expanded') !== null"))
        self.js("document.querySelector('.ent-main').click()")
        time.sleep(.3)
        self.check("[клавиатура] у действий есть подписи для экранного диктора", self.js("[...document.querySelectorAll('[data-act]')].every(b => (b.getAttribute('aria-label') || '').length > 5)"))
        self.click(f'[data-act="forget"][data-base="{base}"]') if False else None
        self.js("ask({title: 'Проверка', text: 'Диалог', ok: 'Закрыть'})")
        self.wait_js("document.querySelector('#ask').open", 5)
        self.check("[клавиатура] диалог модальный, фокус внутри", self.js("document.querySelector('#ask').matches(':modal')") and self.js("document.querySelector('#ask').contains(document.activeElement)"))
        self.js("document.querySelector('#ask').dispatchEvent(new Event('cancel')); document.querySelector('#ask').close()")
        self.check("[доступность] есть область объявлений и индикатор хода", self.js("Boolean(document.querySelector('#sr-status[aria-live=polite]') && document.querySelector('[role=progressbar]'))"))

        # тёмная тема
        self.js("document.documentElement.dataset.theme = 'light'")
        time.sleep(.3)
        light_bg = self.js("getComputedStyle(document.body).backgroundColor")
        self.js("document.documentElement.dataset.theme = 'dark'")
        time.sleep(.3)
        bg = self.js("getComputedStyle(document.body).backgroundColor")
        self.check("[тёмная тема] фон тёмный", self.js("document.documentElement.dataset.theme") == "dark" and bg != light_bg, f"{bg} / {light_bg}")
        self.shot("12-dark")
        self.js("document.documentElement.dataset.theme = 'light'")

        # узкое окно: таблица заменяет ленту, «Назад к чату»
        window = self.holder["window"]
        window.resize(900, 780)
        time.sleep(1)
        self.check("[узкое окно] таблица заменяет ленту", self.js("getComputedStyle(document.querySelector('.main')).display") == "none", self.js("window.innerWidth"))
        self.check("[узкое окно] видна кнопка «Назад к чату»", "Назад к чату" in self.text("#btn-pane-close"))
        self.shot("13-narrow")
        self.click("#btn-pane-close")
        time.sleep(.4)
        self.check("[узкое окно] после «Назад» снова видна лента", self.js("getComputedStyle(document.querySelector('.main')).display") != "none")
        window.resize(1280, 800)
        time.sleep(1)

        self.chrome_checks(cid2)

        # переименовать и удалить чат
        self.js("(() => { const t = document.querySelector('#chat-title'); t.value = 'Кейс Громова'; t.dispatchEvent(new Event('change')); })()")
        self.wait_js("chatData.title === 'Кейс Громова'", 10)
        self.check("[чат] переименование", self.js("chatData.title") == "Кейс Громова" and "Кейс Громова" in self.text("#chat-list"))
        self.js("document.querySelector('#chat-settings').hidden = false; document.querySelector('#chat-settings').open = true")
        self.click("#btn-delete-chat")
        self.ask(None)
        self.wait_js("!document.querySelector('#flash').hidden", 10)
        self.check("[удалить чат] чат стёрт, итог остаётся на странице", self.app.chats.get(cid2) is None and "удалён" in self.text("#flash") and self.js("chatId") is None, self.text("#flash"))
        self.check("[удалить чат] другой чат цел", self.app.chats.get(cid) is not None)

    def pointer(self, selector: str, kind: str, x: float) -> None:
        self.js(f"document.querySelector({json.dumps(selector)}).dispatchEvent(new PointerEvent({json.dumps(kind)}, {{clientX: {x}, clientY: 300, button: 0, pointerId: 1, bubbles: true}}))")

    def key(self, selector: str, key: str, shift: bool = False) -> None:
        self.js(f"document.querySelector({json.dumps(selector)}).dispatchEvent(new KeyboardEvent('keydown', {{key: {json.dumps(key)}, shiftKey: {str(shift).lower()}, bubbles: true, cancelable: true}}))")

    def anchors(self) -> list:
        return self.js("JSON.stringify(['#chat-title', '.main', '#btn-table', '.chat-item'].map(s => { const b = document.querySelector(s).getBoundingClientRect(); return [Math.round(b.x), Math.round(b.y), Math.round(b.width), Math.round(b.height)]; }))")

    def chrome_checks(self, keep: str) -> None:
        """Ширина колонок и действия над строкой чата: ручки, меню, подтверждение удаления."""
        for title in ("Тест А", "Тест Б", "Тест В"):
            self.app.chats.create(title)
        self.js("loadChats()")
        self.wait_js("chatList.length >= 5", 10)
        self.holder["window"].resize(1440, 900)    # колонка значков и зазоры карточек занимают 80 px: прежних 1280 не хватает на ручки с запасом
        time.sleep(.9)
        self.js("setSide(true, {persist: false}); setPane(true)")
        time.sleep(.5)
        prop = lambda name: self.js(f"document.querySelector('#shell').style.getPropertyValue('{name}')")
        # --- ручки ---
        self.check("[ручка] у ручек роль separator и значения", self.js("['#split-side', '#split-pane'].every(s => { const e = document.querySelector(s); return e.getAttribute('role') === 'separator' && e.tabIndex === 0 && +e.getAttribute('aria-valuemin') < +e.getAttribute('aria-valuenow') && +e.getAttribute('aria-valuenow') <= +e.getAttribute('aria-valuemax'); })"))
        before = self.js("document.querySelector('.pane-in').getBoundingClientRect().width")
        start = self.js("document.querySelector('#split-pane').getBoundingClientRect().x + 4")
        self.pointer("#split-pane", "pointerdown", start)
        self.pointer("#split-pane", "pointermove", start - 80)
        mid = self.js("document.querySelector('.pane-in').getBoundingClientRect().width")
        self.check("[ручка] перетаскивание без перехода: ширина таблицы меняется сразу и класс resizing стоит", abs(mid - (before + 80)) < 2 and self.js("document.querySelector('#shell').classList.contains('resizing')"), f"{before} {mid}")
        self.pointer("#split-pane", "pointerup", start - 80)
        time.sleep(.6)
        self.check("[ручка] ширина таблицы сохранена в настройках", self.app.vault.prefs.get("pane_width") == round(before + 80), str(self.app.vault.prefs.get("pane_width")))
        self.check("[ручка] после отпускания переход снова доступен", not self.js("document.querySelector('#shell').classList.contains('resizing')"))
        self.js("document.querySelector('#split-pane').focus()")
        self.key("#split-pane", "ArrowRight")
        self.check("[ручка] стрелка вправо сужает таблицу на 16", abs(self.js("+document.querySelector('#split-pane').getAttribute('aria-valuenow')") - (round(before + 80) - 16)) < 1)
        self.key("#split-pane", "Home")
        self.check("[ручка] Home — наименьшая ширина 320", self.js("+document.querySelector('#split-pane').getAttribute('aria-valuenow')") == 320)
        self.key("#split-pane", "End")
        self.check("[ручка] End — наибольшая ширина не выходит за окно", 320 < self.js("+document.querySelector('#split-pane').getAttribute('aria-valuenow')") <= 640 and self.js("document.querySelector('.main').getBoundingClientRect().width") >= 479)
        self.key("#split-pane", "Enter")
        self.check("[ручка] Enter возвращает ширину по умолчанию", prop("--pane-pref") == "")
        self.pointer("#split-side", "pointerdown", 264)
        self.pointer("#split-side", "pointermove", 354)
        self.pointer("#split-side", "pointerup", 354)
        time.sleep(.6)
        self.check("[ручка] список чатов расширен до 354 и сохранён", prop("--side-pref") == "354px" and self.app.vault.prefs.get("sidebar_width") == 354 and abs(self.js("document.querySelector('#side').getBoundingClientRect().width") - 8 - 354) < 1, prop("--side-pref"))
        self.key("#split-side", "ArrowLeft", True)
        self.check("[ручка] Shift+стрелка меняет на 64", prop("--side-pref") == "290px", prop("--side-pref"))
        time.sleep(.5)
        self.js("location.reload()")
        self.wait_js("document.readyState==='complete' && typeof STATE !== 'undefined' && STATE && chatList.length >= 5", 20)
        self.check("[ручка] ширина списка переживает перезагрузку из настроек", self.js("document.querySelector('#shell').style.getPropertyValue('--side-pref')") == "290px")
        self.js("openChat(chatList[0].id)")
        time.sleep(.8)
        self.key("#split-side", "Enter")
        self.check("[ручка] двойной щелчок и Enter сбрасывают ширину", prop("--side-pref") == "" and self.js("document.querySelector('#side').getBoundingClientRect().width") == 272)
        time.sleep(.5)
        self.check("[ручка] сброс сохранён как «по умолчанию»", self.app.vault.prefs.get("sidebar_width") is None)
        # --- меню строки ---
        self.js("setPane(false)")
        time.sleep(.4)
        row = "document.querySelectorAll('.chat-row')[1]"
        btn = row + ".querySelector('.row-menu-btn')"
        rects = self.anchors()
        self.js(f"{btn}.focus(); {btn}.click()")
        time.sleep(.3)
        self.check("[меню] кнопка «⋯» открывает меню с ролями и aria-expanded", self.js(f"document.querySelector('#row-menu').hidden === false && {btn}.getAttribute('aria-expanded') === 'true' && {btn}.getAttribute('aria-haspopup') === 'menu' && document.querySelectorAll('#row-menu [role=menuitem]').length === 2"))
        self.check("[меню] меню целиком в окне", self.js("(() => { const b = document.querySelector('#row-menu').getBoundingClientRect(); return b.left >= 0 && b.top >= 0 && b.right <= innerWidth && b.bottom <= innerHeight; })()"))
        self.check("[меню] открытие меню ничего не сдвигает", self.anchors() == rects)
        self.key("#row-menu", "ArrowDown")
        self.check("[меню] стрелка вниз ставит фокус на пункт", self.js("document.activeElement.getAttribute('role')") == "menuitem")
        self.key("#row-menu", "End")
        self.check("[меню] End — последний пункт «Удалить чат…»", "Удалить" in self.js("document.activeElement.textContent"))
        self.key("#row-menu", "Escape")
        self.check("[меню] Escape закрывает и возвращает фокус на кнопку", self.js(f"document.querySelector('#row-menu').hidden && document.activeElement === {btn}") and self.anchors() == rects)
        self.js(f"{btn}.id = 'probe-btn'")
        self.key("#probe-btn", "ArrowDown")
        self.check("[меню] стрелка вниз на кнопке открывает меню на первом пункте", self.js("!document.querySelector('#row-menu').hidden && document.activeElement.textContent === 'Переименовать'"))
        self.pointer("#chat-title", "pointerdown", 600)
        self.check("[меню] щелчок снаружи закрывает меню", self.js("document.querySelector('#row-menu').hidden"))
        self.js(f"{row}.querySelector('.chat-item').dispatchEvent(new MouseEvent('contextmenu', {{clientX: 100, clientY: 200, bubbles: true, cancelable: true}}))")
        self.check("[меню] правая кнопка открывает меню у курсора", self.js("!document.querySelector('#row-menu').hidden"))
        self.key("#row-menu", "Escape")
        self.check("[меню] после правой кнопки фокус возвращается на строку", self.js(f"document.activeElement === {row}.querySelector('.chat-item')"))
        self.key(".chat-row .chat-item", "ContextMenu")
        self.check("[меню] клавиша ContextMenu открывает меню", self.js("!document.querySelector('#row-menu').hidden"))
        self.key("#row-menu", "Escape")
        # переименование из меню
        self.js(f"{btn}.click()"); self.js("document.querySelector('#row-menu [data-do=rename]').click()")
        self.wait_js("document.activeElement.id === 'chat-title'", 10)
        self.check("[меню] «Переименовать» открывает чат и ставит фокус в название", self.js("document.activeElement.id") == "chat-title")
        self.key("#chat-title", "Escape")
        # удаление из меню
        victim = self.js("chatList.find(c => c.title.startsWith('Тест В')).id")
        ids = self.js("JSON.stringify([...document.querySelectorAll('.chat-item')].map(b => b.dataset.id))")
        order = json.loads(ids); nxt = order[order.index(victim) + 1] if order.index(victim) + 1 < len(order) else order[order.index(victim) - 1]
        vbtn = f"[...document.querySelectorAll('.chat-row')].find(r => r.querySelector('.chat-item').dataset.id === {json.dumps(victim)}).querySelector('.row-menu-btn')"
        self.js(f"{vbtn}.click(); document.querySelector('#row-menu [data-do=delete]').click()")
        self.wait_js("document.querySelector('#ask').open", 10)
        self.check("[удалить из меню] диалог называет, что стирается, фокус на «Отмена»", "исходные файлы" in self.text("#ask-text") and "Коды замен в хранилище останутся" in self.text("#ask-text") and self.js("document.activeElement.id") == "ask-cancel")
        self.click("#ask-cancel")
        self.wait_js("!document.querySelector('#ask').open", 10)
        self.check("[удалить из меню] отмена оставляет чат и возвращает фокус на кнопку строки", self.app.chats.get(victim) is not None)
        self.js(f"{vbtn}.click(); document.querySelector('#row-menu [data-do=delete]').click()")
        self.ask(None)
        self.wait_js("!document.querySelector('#flash').hidden", 10)
        time.sleep(.4)
        self.check("[удалить из меню] чат стёрт, строки нет в списке", self.app.chats.get(victim) is None and not self.js(f"chatList.some(c => c.id === {json.dumps(victim)})"))
        self.check("[удалить из меню] фокус перешёл на следующую строку", self.js("document.activeElement.dataset && document.activeElement.dataset.id") == nxt, str(self.js("document.activeElement.className")))
        # клавиша Delete на строке
        self.js("document.querySelector('.chat-item').focus()")
        self.key(".chat-item", "Delete")
        self.wait_js("document.querySelector('#ask').open", 10)
        self.check("[удалить с клавиатуры] Delete на строке открывает то же подтверждение", self.js("document.querySelector('#ask-title').textContent") == "Удалить чат?")
        self.click("#ask-cancel")
        self.wait_js("!document.querySelector('#ask').open", 10)
        # чат, который вот-вот удалят в следующем шаге, должен остаться открытым в ленте
        self.js(f"openChat({json.dumps(keep)})")
        self.wait_js("chatId", 10)
        time.sleep(.5)

    def extras(self) -> None:
        from openpyxl import load_workbook
        source = CASE / "Таблица.xlsx"
        # A. Числа Excel: включить, скачать, вернуть.
        self.js("document.querySelector('#opt-numbers').checked = true")
        self.pick("anon", source)
        self.js("document.querySelector('#opt-numbers').checked = true")
        anon_numbers = self.out / "numbers_anon.xlsx"
        self.run_and_save("anon", anon_numbers)
        cells = load_workbook(anon_numbers).worksheets[0]
        self.check("[числа Excel] числа заменены в обезличенном файле", cells["G14"].value != 1200 and isinstance(cells["G14"].value, (int, float)), str(cells["G14"].value))
        self.js("document.querySelector('#opt-numbers').checked = false")
        self.pick("rest", anon_numbers)
        back_numbers = self.out / "numbers_back.xlsx"
        self.run_and_save("rest", back_numbers)
        original = load_workbook(source).worksheets[0]
        restored = load_workbook(back_numbers).worksheets[0]
        same = all(original[c].value == restored[c].value for c in ("G14", "G15", "G16", "G17", "G18", "H18"))
        self.check("[числа Excel] числа вернулись как в исходнике", same, f"{[restored[c].value for c in ('G14','G15','G16')]}")

        # B. Уже обезличенный файл в режиме «Обезличить».
        plain = self.out / f"anon_{source.name}"
        self.pick("anon", plain)
        note = self.text("#files-chat")
        self.check("[уже обезличенный файл] страница предлагает восстановить", "можно восстановить" in note and "есть коды" in note, note[:120])
        self.check("[уже обезличенный файл] главное действие «Восстановить»", self.js("document.querySelector('#btn-chat-rest').classList.contains('main') && !document.querySelector('#btn-chat-anon').classList.contains('main')"))

        # C. Чужой файл: обезличен другим хранилищем.
        from tests.helpers import Env
        other = Env()
        foreign_job = other.anonymize(source)
        foreign = self.out / "foreign.xlsx"
        shutil.copy(foreign_job.files[0].out_path, foreign)
        self.pick("rest", foreign)
        status = self.run_and_save("rest", None)
        messages = self.text("#result-chat .msgs")
        self.check("[чужой файл] возврат остановлен с понятной причиной", status.startswith("Файл не восстановлен") and "другим хранилищем" in messages, f"{status} | {messages[:120]}")
        other.close()

        # D. Резервная копия, очистка, загрузка копии.
        pass
        self.click("#btn-profile"); self.click("#open-settings")
        self.wait_js("document.querySelector('#settings').open", 10)
        backup = self.out / "backup.anonkeys"
        self.click("#btn-backup")
        self.wait_js("document.querySelector('#ask').open", 10)
        self.js("document.querySelector('#ask-input').value = 'abc'")   # слишком короткий: диалог сообщает и остаётся открытым
        self.click("#ask-ok")
        self.wait_js("!document.querySelector('#ask-error').hidden", 10)
        short_error = self.text("#ask-error")
        self.check("[резервная копия] короткий пароль отклонён строкой с причиной", "не короче 6" in short_error, short_error)
        self.queue = [[str(backup)]]
        self.ask("пароль-12345")
        self.wait_js("[...document.querySelectorAll('.toast')].some(t => t.textContent.startsWith('Копия сохранена'))", 15)
        self.check("[резервная копия] копия сохранена", backup.exists() and backup.stat().st_size > 100, self.text(".toast"))
        entities_before = self.app.vault.stats()["entities"]
        self.click("#btn-clear")
        self.ask(None)
        self.wait_js("document.querySelector('#vault-stats').innerText.includes('Сохранено значений: 0')", 10)
        self.check("[резервная копия] хранилище очищено", self.app.vault.stats()["entities"] == 0)
        self.js("document.querySelector('#settings').close()")
        self.pick("rest", plain)
        status = self.run_and_save("rest", None)
        self.check("[резервная копия] после очистки файл не восстанавливается", status.startswith("Файл не восстановлен") or status.startswith("Файл восстановлен не полностью"), self.text("#result-chat")[:200])
        self.click("#btn-profile"); self.click("#open-settings")
        self.wait_js("document.querySelector('#settings').open", 10)
        self.queue = [[str(backup)]]
        self.click("#btn-import")
        self.ask("пароль-12345")
        self.wait_js(f"document.querySelector('#vault-stats').innerText.includes('Сохранено значений: {entities_before}')", 15)
        self.check("[резервная копия] копия загружена, значения на месте", self.app.vault.stats()["entities"] == entities_before, str(self.app.vault.stats()))
        self.js("document.querySelector('#settings').close()")
        self.pick("rest", plain)
        back_after = self.out / "after_import.xlsx"
        status = self.run_and_save("rest", back_after)
        same = back_after.exists() and load_workbook(back_after).worksheets[0]["C6"].value == load_workbook(source).worksheets[0]["C6"].value
        self.check("[резервная копия] файл снова восстанавливается", same, status)
        self.shot("04-after-backup")

    def restart_check(self) -> None:
        """Перезапуск: новое хранилище читает тот же файл на диске и возвращает данные из файла, обезличенного до перезапуска."""
        from openpyxl import load_workbook
        from anonymizer.service import Service, Upload
        self.app.vault.commit()
        vault = Vault(self.home / "vault.bin")
        service = Service(vault, self.home / "work2")
        plain = self.out / "anon_Таблица.xlsx"
        job = service.restore([Upload(plain.name, plain.read_bytes())])
        ok = job.files[0].status == "ok" and load_workbook(job.files[0].out_path).worksheets[0]["C6"].value == load_workbook(CASE / "Таблица.xlsx").worksheets[0]["C6"].value
        self.results.append(("[перезапуск] файл восстанавливается новым процессом из хранилища на диске", ok, str(job.files[0].messages)))
        print(("OK   " if ok else "FAIL ") + "[перезапуск] файл восстанавливается новым процессом из хранилища на диске", flush=True)
        if not ok:
            self.exit_code = 1

    # -- папка как проект чата ------------------------------------------------

    def pick_project_folder(self, path: Path, kind: str = "anonymize") -> bool:
        self.js("modes.chat.files.length = 0; modes.chat.folder = null; renderFiles('chat');")
        self.queue = [[str(path)]]
        self.click("#pick-folder-chat" if kind == "anonymize" else "#pick-folder-answer")
        return self.wait_js("modes.chat.folder && modes.chat.folder.id", 30)

    def wait_project_done(self, steps: int, timeout: float = 240) -> bool:
        return self.wait_js(f"chatData && chatData.project && chatData.steps.length>={steps} && !modes.chat.running", timeout)

    def project_flow(self) -> None:
        secret = "Смирнов Алексей Петрович"
        src = self.home / "Клиент Ромашка"
        for d in ("Договоры/2024", "Отчёты и таблицы"):
            (src / d).mkdir(parents=True)
        (src / "Договоры" / "Договор.docx").write_bytes(docx_bytes([f"Директор {secret}, ООО «Ромашка»."]))
        (src / "Договоры" / "2024" / "Договор.docx").write_bytes(docx_bytes([f"Подписал {secret}. Телефон +7 913 555-12-34."]))
        (src / "Отчёты и таблицы" / "Реестр.xlsx").write_bytes(xlsx_bytes({"A1": "ФИО", "A2": secret}))
        (src / "Отчёты и таблицы" / "Заметка.txt").write_text(f"Встреча с {secret}", encoding="utf-8")
        (src / "заметка.txt").write_text(f"Звонил {secret}", encoding="utf-8")
        (src / "картинка.png").write_bytes(b"\x89PNG not really")
        rels = sorted(str(q.relative_to(src)) for q in src.rglob("*") if q.is_file() and q.suffix != ".png")
        self.js("startDraft();")
        chats_before = len(self.client.json("GET", "/api/chats")["items"])
        ok = self.pick_project_folder(src)
        self.check("[проект] папка выбрана через мост", ok)
        if not ok:
            return
        card = self.js("document.querySelector('#files-chat').textContent")
        self.check("[проект] два варианта, проект по умолчанию", "Добавить в чат как проект" in card and "Обезличить отдельно, без чата" in card
                   and self.js("document.querySelector('[data-fmode=project]').getAttribute('aria-checked')") == "true", card)
        self.check("[проект] счётчик и пропущенный файл", "5 файлов" in card and "Не будут обработаны: 1" in card and "картинка.png" in card, card)
        self.shot("30-project-choice")
        self.check("[проект] кнопка называется «Обезличить проект»", self.js("document.querySelector('#btn-chat-anon').textContent") == "Обезличить проект", self.js("document.querySelector('#btn-chat-anon').outerHTML"))
        self.click("#btn-chat-anon")
        ok = self.wait_project_done(1)
        self.check("[проект] задание завершено, чат проекта создан", ok)
        if not ok:
            return
        chats = self.client.json("GET", "/api/chats")["items"]
        self.check("[проект] создан ровно один чат", len(chats) == chats_before + 1 and sum(1 for c in chats if c["project"]) >= 1)
        self.js("loadChats()")
        time.sleep(.6)
        self.check("[проект] в списке чатов «Проект · 5 файлов»", "Проект · 5 файлов" in self.text("#chat-list") and "Проект · 5 файлов" in self.text("#chat-meta"))
        shown = self.js("[...document.querySelectorAll('#feed .tree-box .trow .tname')].map(e => e.title)")
        self.check("[проект] одна запись и дерево с относительными путями без слияния одинаковых имён",
                   self.js("document.querySelectorAll('#feed li.step').length") == 1 and sorted(shown) == rels, str(shown))
        self.check("[проект] картинка показана как не обработанная", "картинка.png" in (self.js("document.querySelector('#feed .skip').textContent") or ""))
        self.shot("31-project-result")
        # Сохранение всего проекта папкой в путь с кириллицей и пробелом.
        parent = self.out / "Мои проекты с пробелом"
        parent.mkdir()
        self.queue = [[str(parent)]]
        self.click('#feed [data-saveproj="anon"]')
        ok = self.wait_js("document.querySelector('[data-projsaved]:not([hidden])')", 60)
        saved = list(parent.iterdir())
        self.check("[проект] «Сохранить весь проект как папку…» пишет новую папку", ok and len(saved) == 1 and "Сохранено: 5 файлов" in self.text("[data-projsaved]:not([hidden])"), self.text("[data-projsaved]"))
        if len(saved) != 1:
            return
        anon_root = saved[0]
        files_out = [q for q in anon_root.rglob("*") if q.is_file()]
        shape = sorted(len(q.relative_to(anon_root).parts) for q in files_out)
        want = sorted(len(Path(r).parts) for r in rels)
        self.check("[проект] структура папок сохранена (вложенность и число файлов)", shape == want and "Ромашка" not in str(anon_root), f"{shape} != {want}")
        leak = any(secret.split()[0] in docx_text(q) for q in files_out if q.suffix == ".docx")
        self.check("[проект] в сохранённых файлах нет исходных имён", not leak)
        # Правка в таблице и пересборка всего проекта.
        self.js("setPane(true); setTab('hide'); renderTable()")
        time.sleep(2)
        self.js("(() => { const mine = new Set(chatData.files.flatMap(f => f.tokens || [])); const sw = [...document.querySelectorAll('[data-switch]')].find(b => mine.has(b.dataset.switch)) || document.querySelector('[data-switch]'); sw.click(); })()")
        ok = self.wait_js("!document.querySelector('#pane-apply').hidden", 15)
        hint = str(self.js("document.querySelector('#apply-hint').textContent") or "")
        self.check("[проект] после выключателя виден баннер, пересборка называет число файлов", ok and "все файлы проекта: 5" in hint, hint + " " + str(self.js("[pendingEdits(chatData), chatData.keep_originals, document.querySelector('#pane-apply').hidden, document.querySelector('#apply-state').textContent, chatData.steps.map(s => s.kind).join(), paneTab, paneOpen, document.querySelectorAll('[data-switch]').length, document.querySelector('[data-switch]') && document.querySelector('[data-switch]').outerHTML, chatData.project, chatData.title]")))
        self.shot("32-project-pane")
        self.click("#btn-apply")
        ok = self.wait_js("appliedChat === chatId && chatData.steps.length>=2", 240)
        last = self.js("(() => { const st = chatData.steps.filter(s => s.kind==='anonymize').pop(); return st.files.length })()")
        self.check("[проект] «Обновить файл для ИИ» пересобрала все файлы", ok and last == 5, str(last))
        self.js("setPane(false)")
        # Возврат: сохранить обезличенный проект, выбрать его как папку с ответом ИИ.
        ai = self.out / "ответ ИИ"
        ai.mkdir()
        self.queue = [[str(ai)]]
        self.click('#feed [data-saveproj="anon"]')
        self.wait_js("document.querySelector('[data-projsaved]:not([hidden])')", 60)
        sent = list(ai.iterdir())
        ok = bool(sent) and self.wait_js("!document.querySelector('#pick-folder-answer').hidden", 5)
        self.check("[проект] в чате проекта есть «Папка с ответом ИИ»", ok, f"{sent} {self.js('[folderPickable, !!chatData.project, document.querySelector(\"#pick-folder-answer\").hidden, document.querySelector(\"[data-projsaved]:not([hidden])\") && document.querySelector(\"[data-projsaved]:not([hidden])\").textContent]')}")
        if not ok:
            return
        ok = self.pick_project_folder(sent[0], "restore")
        self.check("[проект] папка с ответом выбрана, действие «Вернуть значения в проект»", ok and self.js("document.querySelector('#btn-chat-anon').textContent") == "Вернуть значения в проект", self.js("document.querySelector('#btn-chat-anon').outerHTML"))
        before = self.js("chatData.steps.length")
        self.click("#btn-chat-anon")
        ok = self.wait_project_done(before + 1)
        n_files = self.js("chatData.files.length")
        self.check("[проект] возврат без дублей файлов", ok and n_files == 5 and self.js("chatData.steps[chatData.steps.length-1].kind") == "restore", str(n_files))
        back = self.out / "возврат"
        back.mkdir()
        self.queue = [[str(back)]]
        self.click('#feed [data-saveproj="restored"]')
        self.wait_js("document.querySelector('[data-projsaved]:not([hidden])')", 60)
        got = list(back.iterdir())
        same = len(got) == 1
        diffs = []
        if same:
            for r in rels:
                a = src / r
                q = got[0] / Path(r).parent / f"{a.stem} (восстановлено){a.suffix}"
                if not q.is_file():
                    diffs.append(f"нет {r}; есть {sorted(str(z.relative_to(got[0])) for z in got[0].rglob('*'))[:8]}")
                elif r.endswith(".docx") and docx_text(q) != docx_text(a):
                    diffs.append(f"{r}: {docx_text(q)!r} != {docx_text(a)!r}")
                elif r.endswith(".txt") and q.read_text(encoding="utf-8") != a.read_text(encoding="utf-8"):
                    diffs.append(f"{r}: {q.read_text(encoding='utf-8')!r}")
        same = same and not diffs
        self.check("[проект] после возврата содержимое файлов совпадает с исходным", same, str(diffs or got))
        self.shot("33-project-restored")
        # Отдельный режим работает как раньше и не создаёт чат.
        small = self.home / "Отдельная папка"
        small.mkdir()
        (small / "а.txt").write_text(f"Звонил {secret}", encoding="utf-8")
        self.js("startDraft();")
        chats_before = len(self.client.json("GET", "/api/chats")["items"])
        ok = self.pick_project_folder(small)
        self.click("[data-fmode=separate]")
        self.queue = [[str(self.out / "отдельно")]]
        self.click("#btn-chat-anon")
        ok = ok and self.wait_js("document.querySelector('#result-chat .saved')", 120)
        self.check("[отдельно] режим без чата работает, чат не создан", ok and len(self.client.json("GET", "/api/chats")["items"]) == chats_before, self.text("#result-chat"))
        # Защита от огромной папки.
        huge = self.home / "Огромная"
        huge.mkdir()
        for k in range(501):
            (huge / f"{k}.txt").write_text("x", encoding="utf-8")
        self.js("startDraft();")
        ok = self.pick_project_folder(huge)
        dis = self.js("document.querySelector('[data-fmode=project]').disabled")
        self.check("[проект] 501 файл: проект недоступен, причина показана, отдельный режим остаётся",
                   ok and dis and "500" in self.text("#files-chat") and not self.js("document.querySelector('[data-fmode=separate]').disabled"), self.text("#files-chat"))
        self.shot("34-project-huge")
        self.js("modes.chat.folder = null; renderFiles('chat'); startDraft();")

    def finish(self) -> None:
        failed = [r for r in self.results if not r[1]]
        print(f"\nИтого: {len(self.results) - len(failed)} из {len(self.results)} проверок пройдено", flush=True)
        self.exit_code = 1 if failed else 0
        self.holder["window"].destroy()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--shots", type=Path)
    parser.add_argument("--theme", choices=["light", "dark"], default="light")
    parser.add_argument("--only-project", action="store_true", help="только сценарий папки-проекта")
    parser.add_argument("--inproc", action="store_true", help="мост на App напрямую, без RemoteApp")
    args = parser.parse_args()
    driver = Driver(args.shots, args.theme == "dark", args.inproc)
    driver.only_project = args.only_project
    driver.exit_code = 2
    window = webview.create_window(desktop.TITLE, f"http://127.0.0.1:{driver.port}/?t={driver.app.token}", js_api=driver.bridge,
                                   width=920, height=780, min_size=(520, 520), text_select=True)
    driver.holder["window"] = window
    webview.settings["ALLOW_DOWNLOADS"] = False
    webview.start(driver.run, private_mode=True)
    driver.server.shutdown()
    if not driver.only_project:
        driver.restart_check()
    shutil.rmtree(driver.home, ignore_errors=True)
    return driver.exit_code


if __name__ == "__main__":
    raise SystemExit(main())
