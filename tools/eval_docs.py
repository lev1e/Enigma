"""Document-level evaluation: the labelled text corpus wrapped into real files (docx/xlsx/pptx/pdf/txt/csv/html).

For every format: anonymize -> look for gold values in all parts of the output -> restore -> compare with the
original. Reports leaked values, restore accuracy, integrity (output opens) and time per format and category.

Usage: .venv-mac/bin/python tools/eval_docs.py [--n 60] [--json out.json]
Synthetic identities only (tests/eval/text_corpus_v1.jsonl). Only items the plain-text detector fully covers
are used, so what is measured here is the file pipeline, not detection (tools/eval_text.py measures that).
"""
from __future__ import annotations

import argparse
import collections
import io
import json
import re
import sys
import time
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import fitz  # noqa: E402
from docx import Document  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

from tests.helpers import Env  # noqa: E402

FONT = "/System/Library/Fonts/Supplemental/Arial.ttf"
SKIP = {"COUNTRY", "CITY", "ORG", "ADDRESS", "PROJECT", "FILE", "DOMAIN"}      # context dependent, text eval covers them


def load_items(n: int) -> list[dict]:
    corpus = [json.loads(line) for line in open("tests/eval/text_corpus_v1.jsonl", encoding="utf-8")]
    env = Env()
    out = []
    for it in corpus:
        if not it["spans"] or any(sp["cat"] in SKIP for sp in it["spans"]):
            continue
        d = env.service.build_detector(env.service.settings_for({}))
        d.harvest(it["text"], "x.txt")
        auto = [f for f in d.scan(it["text"], "x.txt", "text") if f.decision.value == "AUTO"]
        if all(any(f.start <= sp["start"] and f.end >= sp["end"] for f in auto) for sp in it["spans"]):
            out.append(it)
    step = max(1, len(out) // n)
    return out[::step][:n]


def build(fmt: str, items: list[dict]) -> bytes:
    lines = [it["text"] for it in items]
    if fmt == "txt":
        return "\n".join(lines).encode()
    if fmt == "csv":
        return "\n".join(f'{i};"{t}"' for i, t in enumerate(lines)).encode()
    if fmt == "html":
        return ("<html><body>" + "".join(f"<p>{t}</p>" for t in lines) + "</body></html>").encode()
    if fmt == "docx":
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = lines[0]
        for t in lines[1:len(lines) // 2]:
            doc.add_paragraph(t)
        rest = lines[len(lines) // 2:]
        table = doc.add_table(rows=len(rest), cols=1)
        for r, t in enumerate(rest):
            table.cell(r, 0).text = t
        buf = io.BytesIO(); doc.save(buf); return buf.getvalue()
    if fmt == "xlsx":
        wb = Workbook(); ws = wb.active
        for i, t in enumerate(lines, 1):
            ws.cell(i, 1, t)
        buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
    if fmt == "pptx":
        prs = Presentation()
        for k in range(0, len(lines), 6):
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            for j, t in enumerate(lines[k:k + 6]):
                box = slide.shapes.add_textbox(Inches(.5), Inches(1.2 + j * .9), Inches(9), Inches(.8))
                box.text_frame.text = t
        buf = io.BytesIO(); prs.save(buf); return buf.getvalue()
    if fmt == "pdf":
        doc = fitz.open()
        for k in range(0, len(lines), 28):
            page = doc.new_page()
            for j, t in enumerate(lines[k:k + 28]):
                page.insert_text((40, 40 + j * 26), re.sub(r"\s+", " ", t).strip(), fontname="arial", fontfile=FONT, fontsize=10)
        return doc.tobytes()
    raise ValueError(fmt)


def blob_text(fmt: str, data: bytes) -> str:
    """Everything readable in the output: decoded parts, raw bytes, PDF text views."""
    parts = [data.decode("utf-8", "ignore"), data.decode("utf-16le", "ignore")]
    if fmt in ("docx", "xlsx", "pptx"):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            for n in z.namelist():
                parts.append(z.read(n).decode("utf-8", "ignore"))
    if fmt == "pdf":
        doc = fitz.open(stream=data, filetype="pdf")
        parts += [p.get_text("text") for p in doc]
        parts.append(str(doc.metadata))
    return "\n".join(parts)


def canon(s: str) -> str:
    return re.sub(r"[\s ​\U000e0000-\U000e007f]+", "", s).lower()


def units(fmt: str, data: bytes) -> str:
    if fmt in ("txt", "csv", "html"):
        return data.decode("utf-8", "ignore")
    if fmt == "docx":
        doc = Document(io.BytesIO(data))
        return "\n".join([doc.sections[0].header.paragraphs[0].text] + [p.text for p in doc.paragraphs]
                         + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    if fmt == "xlsx":
        return "\n".join(str(c.value) for row in load_workbook(io.BytesIO(data)).active.iter_rows() for c in row if c.value)
    if fmt == "pptx":
        return "\n".join(sh.text_frame.text for s in Presentation(io.BytesIO(data)).slides for sh in s.shapes if sh.has_text_frame)
    if fmt == "pdf":
        return "\n".join(p.get_text("text") for p in fitz.open(stream=data, filetype="pdf"))
    return ""


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--n", type=int, default=60)
    ap.add_argument("--json", default="/tmp/eval_docs.json")
    ap.add_argument("--formats", default="txt,csv,html,docx,xlsx,pptx,pdf")
    args = ap.parse_args()
    items = load_items(args.n)
    print(f"items: {len(items)}")
    report: dict[str, dict] = {}
    leaks_detail = []
    for fmt in args.formats.split(","):
        data = build(fmt, items)
        env = Env()
        t0 = time.time()
        job = env.anonymize((f"eval.{fmt}", data))
        t_anon = time.time() - t0
        f = job.files[0]
        row = {"size": len(data), "t_anon": round(t_anon, 2), "blocked": not f.out_path}
        if not f.out_path:
            report[fmt] = row
            continue
        out = env.output_bytes(job)
        text = canon(blob_text(fmt, out))
        total = collections.Counter(); leaked = collections.Counter()
        for it in items:
            for sp in it["spans"]:
                total[sp["cat"]] += 1
                if canon(sp["text"]) in text:
                    leaked[sp["cat"]] += 1
                    leaks_detail.append((fmt, sp["cat"], sp["text"]))
        t1 = time.time()
        r = env.restore(env.as_upload(job))
        row["t_restore"] = round(time.time() - t1, 2)
        try:
            back = env.output_bytes(r)
            if fmt == "pdf":      # restored fragments become separate lines: compare the word multiset
                words = lambda b: sorted(re.findall(r"\w+", units(fmt, b).replace("\u00a0", " ").replace("\u00ad", "").lower()))   # noqa: E731
                row["restore_equal"] = words(back) == words(data)
            else:
                row["restore_equal"] = canon(units(fmt, back)) == canon(units(fmt, data))
            row["opens"] = True
        except Exception as ex:                                             # noqa: BLE001
            row["restore_equal"], row["opens"] = False, f"{type(ex).__name__}: {ex}"
        row["gold"] = sum(total.values()); row["leaked"] = sum(leaked.values())
        row["leaked_by_cat"] = dict(leaked)
        report[fmt] = row
        env.close()
    print(f"{'fmt':6}{'gold':>5}{'leak':>5}{'restore':>9}{'anon s':>8}{'rest s':>8}  blocked")
    for fmt, row in report.items():
        print(f"{fmt:6}{row.get('gold', 0):5}{row.get('leaked', 0):5}{str(row.get('restore_equal')):>9}"
              f"{row['t_anon']:8}{row.get('t_restore', ''):>8}  {row['blocked']}")
    for d in leaks_detail[:40]:
        print("LEAK", d)
    json.dump({"report": report, "leaks": leaks_detail}, open(args.json, "w"), ensure_ascii=False, indent=1)


if __name__ == "__main__":
    main()
