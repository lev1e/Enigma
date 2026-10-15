"""Finance masking probe: mask + restore a set of xlsx files and look for every masked-away amount in EVERY part of the output.

    python tools/finance_probe.py [--real DIR] [--synthetic]

Real files: the amounts that were masked (originals recorded in the vault) must not appear in any output part (theme/styles
excluded) unless the same number also stands untouched in the input as an excluded/non-money cell. Restore: every cell value of
the restored workbook must equal the input value (numbers by Decimal, formulas by text).
"""
from __future__ import annotations

import io
import re
import sys
import zipfile
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

from anonymizer import finance as F  # noqa: E402
from tests.helpers import Env  # noqa: E402

SKIP = re.compile(r"(?i)^xl/(theme|styles)|docprops/|printersettings")
OPTS = {"finance_numbers": True, "finance_text_amounts": True}


def has_number(text: str, token: str) -> bool:
    return re.search(r"(?<![\d.,])" + re.escape(token) + r"(?![\d]|[.,]\d)", text) is not None


def forms(value: str) -> set[str]:
    d = Decimal(value)
    plain = format(d.normalize(), "f")
    out = {plain}
    if d == d.to_integral_value():
        i = int(d)
        out |= {f"{i:,}", f"{i:,}".replace(",", " "), f"{i:,}".replace(",", " ")}
    return {o for o in out if len(o.replace(",", "").replace(" ", "").replace(" ", "").lstrip("-")) >= 5}


def cellmap(data: bytes) -> dict:
    wb = load_workbook(io.BytesIO(data))
    out = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    out[(ws.title, c.coordinate)] = getattr(c.value, "text", c.value)
    return out


def probe(name: str, data: bytes, secrets: list | None = None) -> dict:
    env = Env()
    try:
        job = env.anonymize((name + ".xlsx", data), options=OPTS)
        f = job.files[0]
        res = {"name": name, "status": f.status, "emitted": bool(f.out_path), "warnings": [m["text"][:160] for m in f.messages if m["level"] != "info"],
               "masked_cells": 0, "leaks": [], "restore_diff": [], "restore_ok": None}
        rec = next(iter(env.vault.finance.values()), None)
        if rec is None or not f.out_path:
            return res
        res["masked_cells"] = sum(len(e.get("s", {})) for e in rec["cells"].values())
        out_bytes = env.output_bytes(job)
        z = zipfile.ZipFile(io.BytesIO(out_bytes))
        parts = {n: z.read(n).decode("utf-8", "ignore") for n in z.namelist() if not SKIP.search(n)}
        origs = {F.canon_number(o) for e in rec["cells"].values() for (o, _s) in e.get("s", {}).values()}
        origs |= {F.canon_number(c["orig"]) for c in rec.get("consts", [])}
        if secrets:
            origs |= {str(s) for s in secrets}
        # числа, которые остались в исходнике нетронутыми (исключённые / не суммы): совпадение с ними не утечка
        keep = {F.canon_number(item[1]) for e in rec["cells"].values() for item in e.get("snap", {}).values() if item[0] == "n" and item[3] in ("d", "x")}
        masked_new = {F.canon_number(sc) for e in rec["cells"].values() for (_o, sc) in e.get("s", {}).values()}
        for o in sorted(x for x in origs if x and x not in keep and x not in masked_new):
            for token in forms(o):
                for n, t in parts.items():
                    if has_number(t, token):
                        res["leaks"].append(f"{o} in {n}")
                        break
        back = env.restore(env.as_upload(job))
        if not back.files[0].out_path:
            res["restore_ok"] = False
            res["restore_diff"] = [m["text"][:120] for m in back.files[0].messages]
            return res
        a, b = cellmap(data), cellmap(env.output_bytes(back))
        for key in sorted(set(a) | set(b)):
            va, vb = a.get(key), b.get(key)
            if va == vb:
                continue
            try:
                if Decimal(str(va)) == Decimal(str(vb)):
                    continue
            except Exception:                                   # noqa: BLE001
                pass
            if isinstance(va, str) and isinstance(vb, str) and va.replace(" ", "") == vb.replace(" ", ""):
                continue
            res["restore_diff"].append(f"{key[0]}!{key[1]}: {str(va)[:30]!r} -> {str(vb)[:30]!r}")
        res["restore_ok"] = not res["restore_diff"]
        return res
    finally:
        env.close()


def main() -> int:
    args = sys.argv[1:]
    rows = []
    if "--synthetic" in args:
        from tools import make_finance_realistic as R
        for name, item in R.build_all().items():
            rows.append(("synthetic", probe(name, item["bytes"], item["secrets"])))
    if "--real" in args:
        d = Path(args[args.index("--real") + 1])
        for p in sorted(d.glob("*.xlsx")):
            rows.append(("real", probe(p.stem, p.read_bytes())))
    for kind, r in rows:
        print(f"[{kind}] {r['name']}: status={r['status']} emitted={r['emitted']} masked={r['masked_cells']} leaks={len(r['leaks'])} restore_ok={r['restore_ok']}")
        for x in r["leaks"][:8]:
            print("    LEAK", x)
        for x in r["restore_diff"][:6]:
            print("    DIFF", x)
        for x in r["warnings"][:4]:
            print("    WARN", x)
    return 0


if __name__ == "__main__":
    sys.exit(main())
