#!/usr/bin/env python3
"""Opt-in native Office smoke: synthetic fixtures only; emits count-only JSON.

Run serially. Never selects front/active documents, enumerates user documents,
quits an application, installs anything or prints AppleScript stderr.
"""
from __future__ import annotations

import argparse
import io
import json
import pathlib
import platform
import subprocess
import sys
import tempfile
import uuid

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

APPS = {"docx": "Microsoft Word", "xlsx": "Microsoft Excel", "pptx": "Microsoft PowerPoint"}
SOURCE = "Office smoke Name1 original"
EDITED = "Office smoke Name1 edited"
RESTORED = "Office smoke Synthetic Person edited"


def fixture(ext):
    import docx
    import openpyxl
    from pptx import Presentation
    from pptx.util import Inches
    target = io.BytesIO()
    if ext == "docx":
        obj = docx.Document()
        obj.add_paragraph().add_run(SOURCE).bold = True
    elif ext == "xlsx":
        obj = openpyxl.Workbook()
        obj.active.title = "Smoke"
        obj.active["A1"] = SOURCE
        obj.active["A1"].font = openpyxl.styles.Font(bold=True)
        obj.active["B1"] = 17
    else:
        obj = Presentation()
        slide = obj.slides.add_slide(obj.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        shape.text_frame.paragraphs[0].add_run().text = SOURCE
        shape.text_frame.paragraphs[0].runs[0].font.bold = True
    obj.save(target)
    return target.getvalue()


def check_format(data, ext):
    if ext == "docx":
        import docx
        return docx.Document(io.BytesIO(data)).paragraphs[0].runs[0].bold is True
    if ext == "xlsx":
        import openpyxl
        obj = openpyxl.load_workbook(io.BytesIO(data))
        try:
            return obj.active["A1"].font.bold is True and obj.active["B1"].value == 17
        finally:
            obj.close()
    from pptx import Presentation
    return Presentation(io.BytesIO(data)).slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.bold is True


def prepare(root):
    from anonymizer import docedit
    from anonymizer.restore import Restorer, restore_ooxml
    from anonymizer.vault import Vault
    vault = Vault(path=None, key=b"S" * 32)
    assert vault.token_for("PERSON", "synthetic-person", "Synthetic Person") == "Name1"
    prepared = {}
    for ext in APPS:
        data = fixture(ext)
        name = "smoke." + ext
        tree = docedit.outline(data, name)
        element = next(e for c in tree["containers"] for e in c["elements"]
                       if e.get("text") == SOURCE and e.get("editable"))
        operation = {"op": "replace_text", "at": element["addr"], "expect": element["fp"], "text": EDITED}
        try:
            docedit.apply_patch(data, name, {"ops": [{**operation, "text": EDITED.replace("Name1", "Name2")}]})
        except docedit.PatchRejected:
            pass
        else:
            raise AssertionError("label protection")
        edited, _ = docedit.apply_patch(data, name, {"ops": [operation]})
        updated = docedit.outline(edited, name)
        assert any(e.get("text") == EDITED for c in updated["containers"] for e in c["elements"])
        prefix = "AnonymizerOfficeSmoke-" + uuid.uuid4().hex
        source_path = root / (prefix + "-source." + ext)
        edited_path = root / (prefix + "-edited." + ext)
        restored_path = root / (prefix + "-restored." + ext)
        source_path.write_bytes(data)
        edited_path.write_bytes(edited)
        restorer = Restorer(vault)
        assert restore_ooxml(edited_path, restored_path, restorer, numbers=False)
        assert restorer.stats.restored == 1
        restored = restored_path.read_bytes()
        tree = docedit.outline(restored, name)
        assert any(e.get("text") == RESTORED for c in tree["containers"] for e in c["elements"])
        assert all(check_format(payload, ext) for payload in (data, edited, restored))
        prepared[ext] = [(source_path, SOURCE), (edited_path, EDITED), (restored_path, RESTORED)]
    return prepared


def quote(value):
    return '"' + str(value).replace('\\', '\\\\').replace('"', '\\"') + '"'


def script_for(ext, documents):
    # Vocabulary: /usr/bin/sdef /Applications/Microsoft\ <app>.app and the
    # included /System/Library/ScriptingDefinitions/CocoaStandard.sdef.
    app = APPS[ext]
    kind = {"docx": "document", "xlsx": "workbook", "pptx": "presentation"}[ext]
    chunks = [f'tell application {quote(app)}', 'with timeout of 45 seconds', 'set passedCount to 0', 'set closedCount to 0']
    for path, expected in documents:
        if ext == "docx":
            opening = f'open (POSIX file {quote(path)}) read only true add to recent files false'
            reading = 'content of text object of ownDocument'
            comparison = f'actualText contains {quote(expected)}'
        elif ext == "xlsx":
            opening = f'open workbook workbook file name {quote(path)} read only true add to mru false'
            reading = 'value of range "A1" of worksheet "Smoke" of ownDocument'
            comparison = f'actualText is {quote(expected)}'
        else:
            opening = f'open (POSIX file {quote(path)})'
            reading = 'content of text range of text frame of shape 1 of slide 1 of ownDocument'
            comparison = f'actualText contains {quote(expected)}'
        chunks += [
            'set ownDocument to missing value', 'set verifiedOwn to false', 'try', opening,
            f'set ownDocument to {kind} {quote(path.name)}',
            f'if (POSIX path of ((full name of ownDocument) as alias)) is not {quote(path)} then error number -2701',
            'set verifiedOwn to true', f'set actualText to {reading}',
            f'if not ({comparison}) then error number -2702',
            'close ownDocument saving no', 'set closedCount to closedCount + 1',
            'set passedCount to passedCount + 1', 'on error errorMessage number errorNumber',
            'if verifiedOwn then', 'try', 'with timeout of 2 seconds',
            'close ownDocument saving no', 'set closedCount to closedCount + 1',
            'end timeout', 'end try', 'end if',
            'return (passedCount as text) & "," & (closedCount as text) & "," & (errorNumber as text)',
            'end try',
        ]
    chunks += ['return (passedCount as text) & "," & (closedCount as text) & ",0"', 'end timeout', 'end tell']
    return '\n'.join(chunks)


def run_app(ext, documents):
    if not pathlib.Path('/Applications', APPS[ext] + '.app').is_dir():
        return {"passed": 0, "closed": 0, "missing_app": 1}
    try:
        result = subprocess.run(['/usr/bin/osascript', '-'], input=script_for(ext, documents),
                                capture_output=True, text=True, timeout=55)
    except subprocess.TimeoutExpired:
        return {"passed": 0, "closed": 0, "timeout_or_dialog": 1, "cleanup_unconfirmed": 3}
    try:
        if result.returncode:
            # Never emit stderr: host errors can include unrelated paths or text.
            return {"passed": 0, "closed": 0, "script_or_automation_error": 1, "cleanup_unconfirmed": 3}
        passed, closed, code = map(int, result.stdout.strip().split(','))
    except ValueError:
        return {"passed": 0, "closed": 0, "unexpected_response": 1, "cleanup_unconfirmed": 3}
    return {"passed": passed, "closed": closed, "error_number": code,
            "automation_denied": int(code == -1743), "timeout_or_dialog": int(code == -1712)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--prepare-only', action='store_true', help='Validate fixtures without opening Office.')
    args = parser.parse_args()
    if platform.system() != 'Darwin' and not args.prepare_only:
        print(json.dumps({"unsupported_platform": 1}))
        return 2
    # Keep synthetic files on native failure: an app may still hold its own file.
    root = pathlib.Path(tempfile.mkdtemp(prefix='AnonymizerOfficeSmoke-')).resolve()
    try:
        prepared = prepare(root)
    except Exception:
        print(json.dumps({"fixture_failures": 1}))
        return 1
    results = {"fixtures": 9, "protected_label_checks": 3, "format_checks": 9, "apps": {}}
    if not args.prepare_only:
        for ext, documents in prepared.items():
            results['apps'][ext] = run_app(ext, documents)
    success = args.prepare_only or all(r.get('passed') == 3 and r.get('closed') == 3 and r.get('error_number') == 0
                                       for r in results['apps'].values())
    if success:
        import shutil
        shutil.rmtree(root)
    results['retained_synthetic_directories'] = int(not success)
    print(json.dumps(results, sort_keys=True))
    return 0 if success else 1


if __name__ == '__main__':
    raise SystemExit(main())
