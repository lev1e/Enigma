"""Labelled English corpus (people, US/UK addresses, cities in prose, phones, SSN/NI, passports, DOB, orgs, emails, IBAN/cards, secrets).

  python tools/make_en_corpus.py build [--seed 7]   -> tests/eval/en_corpus/{files, gold.json, clean/} + tests/eval/text_corpus_en.jsonl
  python tools/make_en_corpus.py run [--jobs 2]      -> end-to-end per file (anonymize -> leak check on output -> restore equality), per-kind table
  python tools/make_en_corpus.py clean               -> false positives on clean English prose (AUTO findings per 1000 words)

Markup used by the templates: [CAT|text] = must be hidden. Everything else is clean text. All identities are invented
(reserved example domains, 555 phone numbers, Luhn-valid test cards, mod-97-valid IBANs).
Files are written under tests/eval/en_corpus (generated, not committed).
"""
from __future__ import annotations

import argparse
import collections
import io
import json
import random
import re
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tools"))
OUT = ROOT / "tests" / "eval" / "en_corpus"
TEXT_JSONL = ROOT / "tests" / "eval" / "text_corpus_en.jsonl"
MARK = re.compile(r"\[([A-Z_]+)\|([^\]]+)\]")

GIVEN = ["John", "Grace", "Marcus", "Elena", "Priya", "Thomas", "Oliver", "Samantha", "Daniel", "Rebecca", "Andrew", "Natalie", "Kevin", "Laura", "Patrick", "Megan",
         "Raymond", "Olivia", "Gareth", "Fiona", "Terrence", "Alicia", "Dwayne", "Harriet", "Nathaniel", "Imogen"]
NICK = [("Robert", "Bob"), ("William", "Bill"), ("Elizabeth", "Liz"), ("James", "Jimmy"), ("Katherine", "Kate"), ("Richard", "Dick"), ("Margaret", "Maggie"), ("Jonathan", "Jon")]
SURN = ["Whitfield", "Okonkwo", "Lindqvist", "Haverford", "Delacroix", "Abernathy", "Kowalczyk", "Underhill", "Vandermeer", "Castellanos", "Pemberton", "Hargrove", "Thornbury",
        "Fitzgerald", "McAllister", "O'Connell", "Brennan", "Ashworth", "Carmichael", "Dunleavy", "Gallagher", "Hollingsworth", "Ibrahim", "Jankowski", "Kilpatrick", "Montgomery",
        "Nakamura", "Oyelaran", "Prescott", "Quigley", "Rutherford", "Sandoval", "Tillman", "Whitcombe"]
ORGS = ["Northwind Holdings", "Blue Harbor Capital", "Quillfeather Labs", "Redwood Ferris Logistics", "Orchard Lane Partners", "Halcyon Dynamics", "Brightline Systems Ltd",
        "Corvid & Finch LLP", "Ashgrove Pharmaceuticals", "Meridian Freight Inc.", "Kestrel Analytics", "Tidewater Insurance", "Larkfield University", "Stonebridge Bank", "Oakmont Realty",
        "Silverline Telecom plc", "Marlowe & Sons", "Granite Peak Energy"]
STREETS = ["Marchmont Avenue", "Calloway Court", "Fenwick Street", "Larkspur Way", "Brambleton Road", "Elm Grove", "Chestnut Lane", "Kingsway", "Harcourt Terrace", "Wexford Drive",
           "Old Mill Road", "Sycamore Close", "Tavistock Boulevard", "Pinecrest Parkway"]
US_CITIES = [("Springfield", "IL", "62704"), ("Portland", "OR", "97205"), ("Austin", "TX", "78701"), ("Columbus", "OH", "43215"), ("Tulsa", "OK", "74103"), ("Denver", "CO", "80202"),
             ("Boston", "MA", "02108"), ("Raleigh", "NC", "27601"), ("Pittsburgh", "PA", "15222"), ("Sacramento", "CA", "95814")]
UK_CITIES = [("London", "NW1 6XE"), ("Leeds", "LS1 4AP"), ("Manchester", "M1 1AE"), ("Bristol", "BS1 5TR"), ("Sheffield", "S1 2HE"), ("Glasgow", "G2 4QZ"), ("Cardiff", "CF10 3AT"),
             ("Nottingham", "NG1 5FS"), ("Edinburgh", "EH1 1YZ")]
PROSE_CITIES = ["Leeds", "Portland", "Austin", "Tulsa", "Boston", "Denver", "Manchester", "Bristol", "Cincinnati", "Sacramento", "Glasgow", "Sheffield", "Raleigh", "Pittsburgh", "Dunedin", "Halifax"]
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]


class G:
    """value generators bound to one RNG"""

    def __init__(self, seed: str):
        self.r = random.Random(seed)

    def d(self, n: int) -> str:
        return "".join(str(self.r.randrange(10)) for _ in range(n))

    def person_parts(self):
        if self.r.random() < .18:
            f, nick = self.r.choice(NICK)
            return f, self.r.choice(SURN), nick
        return self.r.choice(GIVEN), self.r.choice(SURN), None

    def person(self, style: str | None = None) -> str:
        g, s, nick = self.person_parts()
        style = style or self.r.choice(["full", "full", "mid", "init", "init2", "comma", "caps", "nick", "last_init", "title"])
        mi = self.r.choice("ABCDEJKLMNRST")
        if style == "mid":
            return f"{g} {mi}. {s}"
        if style == "init":
            return f"{g[0]}. {s}"                         # J. Whitfield
        if style == "init2":
            return f"{g[0]}.{mi}. {s}"
        if style == "comma":
            return f"{s}, {g}"
        if style == "caps":
            return f"{g} {s}".upper()
        if style == "nick":
            return f"{nick or g} {s}"
        if style == "last_init":
            return f"{s} {g[0]}."
        if style == "title":
            return f"{self.r.choice(['Mr.', 'Ms.', 'Dr.', 'Mrs.'])} {s}"
        return f"{g} {s}"

    def phone(self) -> str:
        r = self.r
        return r.choice([f"(415) 555-{self.d(4)}", f"+1 212 555 {self.d(4)}", f"+1-312-555-{self.d(4)}", f"617.555.{self.d(4)}", f"+44 20 7946 {self.d(4)}", f"07700 900{self.d(3)}",
                         f"020 7946 {self.d(4)}", f"+44 (0)161 496 {self.d(4)}", f"555-{self.d(3)}-{self.d(4)}"])

    def email(self, p: str | None = None) -> str:
        r = self.r
        g, s = r.choice(GIVEN).lower(), r.choice(SURN).lower().replace("'", "")
        return r.choice([f"{g}.{s}@example.com", f"{g[0]}{s}{self.d(2)}@demo-mail.example.org", f"{s}_{g}@corp-demo.example.net", f"billing{self.d(2)}@halcyon-demo.example.com"])

    def ssn(self) -> str:
        return f"{self.r.randrange(100, 665)}-{self.r.randrange(10, 99)}-{self.r.randrange(1000, 9999)}"

    def ni(self) -> str:
        while True:
            a, b = self.r.choice("ABCEGHJKLMNPRSTWXYZ"), self.r.choice("ABCEGHJKLMNPRSTWXYZ")
            if a + b not in ("BG", "GB", "NK", "KN", "TN", "NT", "ZZ"):
                break
        return self.r.choice([f"{a}{b} {self.d(2)} {self.d(2)} {self.d(2)} {self.r.choice('ABCD')}", f"{a}{b}{self.d(6)}{self.r.choice('ABCD')}"])

    def passport(self) -> str:
        return self.r.choice([self.d(9), f"{self.r.choice('ABCDEFGH')}{self.d(8)}", f"{self.r.choice('XYZ')}{self.d(7)}", f"{self.r.choice('ABCDEFGH')}{self.d(7)}"])

    def dob(self) -> str:
        r = self.r
        y, m, dd = r.randrange(1948, 2004), r.randrange(1, 13), r.randrange(1, 29)
        mon = MONTHS[m - 1]
        return r.choice([f"{m:02d}/{dd:02d}/{y}", f"{dd:02d}/{m:02d}/{y}", f"{y}-{m:02d}-{dd:02d}", f"{mon} {dd}, {y}", f"{dd} {mon} {y}", f"{dd:02d}.{m:02d}.{y}",
                         f"{mon[:3]} {dd}, {y}", f"{dd}th of {mon} {y}" if dd > 3 else f"{dd} {mon[:3]} {y}", f"{m}/{dd}/{str(y)[2:]}"])

    def card(self) -> str:
        pre = self.r.choice(["4111", "5500", "4012", "5105"])
        n = pre + self.d(11)
        dg = [int(c) for c in n]
        s = sum((x * 2 - 9 if x * 2 > 9 else x * 2) if i % 2 == 1 else x for i, x in enumerate(reversed(dg + [0])))
        n += str((10 - s % 10) % 10)
        return self.r.choice([" ".join(n[i:i + 4] for i in range(0, 16, 4)), n, "-".join(n[i:i + 4] for i in range(0, 16, 4))])

    def iban(self) -> str:
        bban = f"{self.r.choice(['NWBK', 'BARC', 'HBUK'])}{self.d(14)}"
        chk = 98 - int("".join(str(int(c, 36)) for c in bban + "GB00")) % 97
        v = f"GB{chk:02d}{bban}"
        return self.r.choice([v, " ".join(v[i:i + 4] for i in range(0, len(v), 4))])

    def ip(self) -> str:
        return f"10.{self.r.randrange(1, 250)}.{self.r.randrange(1, 250)}.{self.r.randrange(2, 250)}"

    def secret(self) -> str:
        return self.r.choice([f"sk_live_{self.d(6)}Zq{self.d(8)}Xw", f"Rv{self.d(3)}!kd{self.d(4)}Pm", f"ghp_{self.d(4)}aB{self.d(6)}cD{self.d(8)}eF", f"Tr0ub4dor&{self.d(2)}x"])

    def street_addr(self) -> str:
        r = self.r
        n, st = r.randrange(2, 4900), r.choice(STREETS)
        if r.random() < .5:
            c, stt, z = r.choice(US_CITIES)
            return r.choice([f"{n} {st}, {c}, {stt} {z}", f"{n} {st}, Suite {r.randrange(2, 900)}, {c}, {stt} {z}", f"{n} {st}, Apt {r.randrange(1, 40)}{r.choice('ABC')}, {c}, {stt}"])
        c, z = r.choice(UK_CITIES)
        return r.choice([f"{n} {st}, {c} {z}", f"Flat {r.randrange(1, 20)}, {n} {st}, {c}", f"{n}{r.choice('AB')} {st}, {c}, {z}"])

    def org(self) -> str:
        return self.r.choice(ORGS)

    def city(self) -> str:
        return self.r.choice(PROSE_CITIES)


# ----------------------------------------------------------------------------------------------------- templates (markup)
def line_templates(g: G):
    P, O, A, C = g.person, g.org, g.street_addr, g.city
    r = g.r
    T = [
        lambda: f"Applicant: [PERSON|{P('full')}], date of birth [BIRTH_DATE|{g.dob()}]",
        lambda: f"Dear [PERSON|{g.r.choice(GIVEN)}],",
        lambda: f"Hi [PERSON|{P('nick')}], thanks for the quick reply.",
        lambda: f"Regards,\n[PERSON|{P('full')}]",
        lambda: f"Signed by [PERSON|{P('mid')}] on behalf of [ORG|{O()}].",
        lambda: f"Please send the signed copy to [PERSON|{P('caps')}] at [EMAIL|{g.email()}].",
        lambda: f"Witness: [PERSON|{P('comma')}], passport no. [PASSPORT|{g.passport()}]",
        lambda: f"Contact: [PERSON|{P('init')}] ([PHONE|{g.phone()}])",
        lambda: f"Contact [PERSON|{P('last_init')}] in accounts about the refund.",
        lambda: f"{r.choice(['Mr.', 'Ms.', 'Dr.'])} [PERSON|{r.choice(SURN)}] will meet us on Tuesday.",
        lambda: f"Our new hire, [PERSON|{P('full')}], joined from [ORG|{O()}] last month.",
        lambda: f"I spoke with [PERSON|{P('nick')}] yesterday; [PERSON|{P('init2')}] agreed to the terms.",
        lambda: f"From: [PERSON|{P('full')}] <[EMAIL|{g.email()}]>",
        lambda: f"EMPLOYEE: [PERSON|{P('caps')}]",
        lambda: f"Patient [PERSON|{P('comma')}], DOB [BIRTH_DATE|{g.dob()}], SSN [SSN|{g.ssn()}]",
        lambda: f"National Insurance number: [NI|{g.ni()}]",
        lambda: f"NI no [NI|{g.ni()}] and tax code 1257L.",
        lambda: f"Social Security Number [SSN|{g.ssn()}] on file for [PERSON|{P('full')}].",
        lambda: f"Passport: [PASSPORT|{g.passport()}], expires 2031",
        lambda: f"Passport number [PASSPORT|{g.passport()}] issued in [CITY|{C()}].",
        lambda: f"Born on [BIRTH_DATE|{g.dob()}] in [CITY|{C()}].",
        lambda: f"DOB: [BIRTH_DATE|{g.dob()}]",
        lambda: f"Ship to [ADDRESS|{A()}], attn [PERSON|{P('full')}].",
        lambda: f"Billing address: [ADDRESS|{A()}]",
        lambda: f"She moved to [CITY|{C()}] last spring and now works from the [CITY|{C()}] office.",
        lambda: f"We ship from our warehouse in [CITY|{C()}] to customers in [CITY|{C()}] and [CITY|{C()}].",
        lambda: f"The flight from [CITY|{C()}] to [CITY|{C()}] was delayed by two hours.",
        lambda: f"The conference takes place in [CITY|{C()}] in October.",
        lambda: f"Residence: [CITY|{C()}]",
        lambda: f"City: [CITY|{C()}]",
        lambda: f"Call [PHONE|{g.phone()}] or write to [EMAIL|{g.email()}].",
        lambda: f"Mobile: [PHONE|{g.phone()}]  Fax: [PHONE|{g.phone()}]",
        lambda: f"Wire the funds to IBAN [IBAN|{g.iban()}], reference INV-{g.d(5)}.",
        lambda: f"Card on file: [CARD|{g.card()}], exp 09/29.",
        lambda: f"Server [IP_ADDRESS|{g.ip()}] is reachable via VPN.",
        lambda: f"api_key = [SECRET|{g.secret()}]",
        lambda: f"password: [SECRET|{g.secret()}]",
        lambda: f"[ORG|{O()}] and [ORG|{O()}] agreed to extend the contract.",
        lambda: f"Vendor: [ORG|{O()}], {r.choice(['net 30', 'net 45'])}.",
        lambda: f"She works for [ORG|{O()}] as an analyst.",
        lambda: f"This agreement is made between [ORG|{O()}] ([ADDRESS|{A()}]) and [PERSON|{P('mid')}], born [BIRTH_DATE|{g.dob()}].",
        lambda: f"Please find attached the invoice from [ORG|{O()}] for [PERSON|{P('full')}].",
    ]
    return T + hard_templates(g)


def hard_templates(g: G):
    """no label right before the value, mixed styles: the part of the corpus that guards against tuning to the easy cues"""
    P, O, A, C = g.person, g.org, g.street_addr, g.city
    r = g.r
    g2 = lambda: r.choice(GIVEN) + " " + r.choice(SURN)                # noqa: E731
    return [
        lambda: f"Yesterday [PERSON|{g2()}] called about the invoice.",
        lambda: f"[PERSON|{g2()}] and [PERSON|{g2()}] attended the quarterly review.",
        lambda: f"According to Dr. [PERSON|{r.choice(SURN)}], the results were inconclusive.",
        lambda: f"Thanks, [PERSON|{r.choice(GIVEN)}] — I will forward it to [PERSON|{P('last_init')}].",
        lambda: f"cc: [PERSON|{P('comma')}]; [PERSON|{P('comma')}]",
        lambda: f"PREPARED BY: [PERSON|{P('caps')}]  APPROVED BY: [PERSON|{P('caps')}]",
        lambda: f"To: [PERSON|{P('full')}] <[EMAIL|{g.email()}]>",
        lambda: f"[PERSON|{g2()}]'s laptop was returned to IT.",
        lambda: f"The contract was countersigned by [PERSON|{r.choice(['Anna-Marie', 'Jean-Paul'])} {r.choice(SURN)}-{r.choice(SURN)}].",
        lambda: f"Meeting with [PERSON|{P('init')}] and [PERSON|{P('mid')}] moved to Friday.",
        lambda: f"Offices: [CITY|{C()}], [CITY|{C()}] and [CITY|{C()}].",
        lambda: f"Based in [CITY|{C()}], UK, the team has grown to fifty people.",
        lambda: f"[CITY|{C().upper()}] — The city council approved the plan on Tuesday.",
        lambda: f"Our [CITY|{C()}] branch closes at five; the [CITY|{C()}] team works remotely.",
        lambda: f"Registered office: [ADDRESS|{A()}]",
        lambda: f"Write to us at [ADDRESS|P.O. Box {r.randrange(1000, 9999)}, {r.choice(US_CITIES)[0]}, TX 78701].",
        lambda: f"Collect from [ADDRESS|{r.randrange(2, 90)} {r.choice(STREETS)}] after noon.",
        lambda: f"Employee {g.d(5)}: DOB [BIRTH_DATE|{g.dob()}], NI [NI|{g.ni()}].",
        lambda: f"Customer reference {g.d(6)}; contact on [PHONE|{g.phone()}] after 5pm.",
        lambda: f"Please wire to [IBAN|{g.iban()}] and quote [PERSON|{P('full')}] as the payer.",
        lambda: f"Her card [CARD|{g.card()}] was declined twice.",
        lambda: f"Login with user admin and password [SECRET|{g.secret()}] on the staging box.",
    ]


def make_doc(g: G, n: int = 16) -> str:
    T = line_templates(g)
    return "\n".join(g.r.choice(T)() for _ in range(n)) + "\n"


def make_table(g: G, n: int = 6) -> tuple[list[str], list[list[str]]]:
    """bare values under column headers (no label inside the cell)"""
    hdr_sets = [["Name", "DOB", "Passport", "City", "Phone"], ["Employee", "Date of Birth", "Passport No", "SSN", "Email"], ["Full Name", "Birth Date", "NI Number", "Town", "Address"],
                ["Customer", "Passport Number", "Date of birth", "City", "Card"], ["Applicant", "DOB", "Tax ID", "City", "IBAN"]]
    hdr = g.r.choice(hdr_sets)
    rows = []
    for _ in range(n):
        row = []
        for h in hdr:
            hl = h.lower()
            if "name" in hl or hl in ("employee", "customer", "applicant"):
                v = ("PERSON", g.person(g.r.choice(["full", "comma", "caps", "mid"])))
            elif "birth" in hl or hl == "dob":
                v = ("BIRTH_DATE", g.dob())
            elif "passport" in hl:
                v = ("PASSPORT", g.passport())
            elif hl == "ssn" or "tax id" in hl:
                v = ("SSN", g.ssn())
            elif "ni " in hl + " ":
                v = ("NI", g.ni())
            elif hl in ("city", "town"):
                v = ("CITY", g.city())
            elif "phone" in hl:
                v = ("PHONE", g.phone())
            elif "email" in hl:
                v = ("EMAIL", g.email())
            elif "address" in hl:
                v = ("ADDRESS", g.street_addr())
            elif "card" in hl:
                v = ("CARD", g.card())
            else:
                v = ("IBAN", g.iban())
            row.append(f"[{v[0]}|{v[1]}]")
        rows.append(row)
    return hdr, rows


# --------------------------------------------------------------------------------------------------------- clean prose
CLEAN = [
    "Quarterly Report: Sales Overview\n\nThe Board will meet on Monday to review the Annual Plan. Please read the Terms of Service and the Privacy Policy before signing. "
    "Chapter Two covers the Project Timeline, the Risk Register and the Steering Committee. Will you mark the box if you agree? Mark the items in red. Grant funding will resume in May.",
    "Getting Started\n\nInstall the package, open the Settings page and press Save. The Control Panel shows System Status, Network Usage and Disk Space. Contact Support if the Service "
    "Desk is closed. Hope this helps. The Rose Garden Cafe opens at nine, and Hunter-gatherer societies are discussed in Section Four.",
    "Terms and Conditions\n\n1. Definitions. In this Agreement, the Customer means the party named in the Order Form. 2. Payment. All Invoices are due within thirty days. "
    "3. Governing Law. This Agreement is governed by the laws of England and Wales. The Supplier shall provide the Services with reasonable care.",
    "Recipe: Lemon Chicken\n\nPreheat the oven. Season the Chicken with Salt and Pepper, add Olive Oil and bake for forty minutes. Serve with Rice or Bread. "
    "Fresh Basil makes a good garnish. Store leftovers in the Fridge for up to three days.",
    "Release Notes 2.4\n\nNew: Dark Mode, Faster Search, Export to PDF. Fixed: Login Page crash on Safari. Known issues: the Dashboard may be slow on Windows. Thanks to all Contributors. "
    "See the Changelog for details, and the Roadmap for what comes Next.",
    "Meeting Minutes\n\nAttendees: the Product Team, the Design Team and Finance. Agenda: Budget Review, Hiring Plan, Office Move. Action Items: update the Roadmap, send the Summary, "
    "book the Conference Room. The next meeting is on Friday. Regards to everyone who joined.",
    "The city council approved the new Budget on Tuesday. The Mayor said the Transit Plan would improve Public Safety. Residents can submit comments online until June. "
    "Opponents argue the Housing Strategy ignores Rural Areas. A vote on the Climate Action Plan is expected in the Autumn.",
    "Travel guide. The best time to visit is Spring or early Autumn. Pack a Raincoat, Sunscreen and Comfortable Shoes. Museums are free on the first Sunday of the Month. "
    "The Old Town has narrow streets, a Market Square and a Clock Tower. Trains run every Hour.",
    "Job Description: Software Engineer\n\nResponsibilities: build Features, review Code, write Tests and mentor Junior Staff. Requirements: Python, SQL, Git and good Communication. "
    "Benefits: Health Insurance, Paid Leave, Remote Work. Apply through the Careers page.",
    "Meeting invite: Weekly Sync. Join by Video or Phone. Please Add your Notes to the Shared Document before Monday. Mark, Grant and Will are on leave in August; Hope and Faith are covering.",
]


# ------------------------------------------------------------------------------------------------------------ file writers
def strip_markup(s: str) -> tuple[str, list[tuple[str, str]]]:
    gold = [(m.group(1), m.group(2)) for m in MARK.finditer(s)]
    return MARK.sub(lambda m: m.group(2), s), gold


def spans_of(s: str):
    out, pos, plain = [], 0, []
    last = 0
    for m in MARK.finditer(s):
        plain.append(s[last:m.start()]); pos += m.start() - last
        out.append({"start": pos, "end": pos + len(m.group(2)), "cat": m.group(1), "text": m.group(2)})
        plain.append(m.group(2)); pos += len(m.group(2)); last = m.end()
    plain.append(s[last:])
    return "".join(plain), out


def write_files(g: G, idx: int, text: str, table, fmt: str) -> tuple[bytes, list[tuple[str, str]], str]:
    """-> (file bytes, gold, plain text)"""
    body, gold = strip_markup(text)
    hdr = rows = None
    if table:
        hdr = table[0]
        rows = [[strip_markup(c)[0] for c in r] for r in table[1]]
        for r in table[1]:
            gold += [x for c in r for x in strip_markup(c)[1]]
    if fmt in ("txt", "md"):
        tab = ""
        if rows:
            tab = ("\n| " + " | ".join(hdr) + " |\n|" + "---|" * len(hdr) + "\n" + "\n".join("| " + " | ".join(r) + " |" for r in rows) + "\n") if fmt == "md" else \
                  ("\n" + "\t".join(hdr) + "\n" + "\n".join("\t".join(r) for r in rows) + "\n")
        return (("# Notes\n\n" if fmt == "md" else "") + body + tab).encode(), gold, body
    if fmt == "csv":
        import csv
        b = io.StringIO(); w = csv.writer(b)
        if rows:
            w.writerow(hdr); w.writerows(rows)
        else:
            [w.writerow([ln]) for ln in body.splitlines()]
        return b.getvalue().encode(), gold, body
    if fmt == "html":
        tab = ""
        if rows:
            tab = "<table><tr>" + "".join(f"<th>{h}</th>" for h in hdr) + "</tr>" + "".join("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>" for r in rows) + "</table>"
        from html import escape
        return ("<html><body>" + "".join(f"<p>{escape(ln)}</p>" for ln in body.splitlines()) + tab + "</body></html>").encode(), gold, body
    if fmt == "docx":
        from docx import Document
        d = Document()
        for ln in body.splitlines():
            d.add_paragraph(ln)
        if rows:
            t = d.add_table(rows=1, cols=len(hdr)); t.rows[0].cells[:] and [setattr(c, "text", h) for c, h in zip(t.rows[0].cells, hdr)]
            for r in rows:
                cells = t.add_row().cells
                for c, v in zip(cells, r):
                    c.text = v
        b = io.BytesIO(); d.save(b); return b.getvalue(), gold, body
    if fmt == "xlsx":
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        if rows:
            ws.append(hdr); [ws.append(r) for r in rows]
        else:
            [ws.append([ln]) for ln in body.splitlines()]
        b = io.BytesIO(); wb.save(b); return b.getvalue(), gold, body
    if fmt == "pptx":
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation(); lines = body.splitlines()
        for k in range(0, max(1, len(lines)), 6):
            s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = f"Slide {k // 6 + 1}"
            s.shapes.add_textbox(Inches(.5), Inches(1.5), Inches(9), Inches(4)).text_frame.text = "\n".join(lines[k:k + 6])
        if rows:
            s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = "Table"
            t = s.shapes.add_table(len(rows) + 1, len(hdr), Inches(.3), Inches(1.5), Inches(9.4), Inches(3)).table
            for j, h in enumerate(hdr):
                t.cell(0, j).text = h
            for i, r in enumerate(rows):
                for j, v in enumerate(r):
                    t.cell(i + 1, j).text = v
        b = io.BytesIO(); prs.save(b); return b.getvalue(), gold, body
    if fmt == "pdf":
        import fitz
        doc = fitz.open(); page = doc.new_page(); y = 60
        allines = body.splitlines() + ([" | ".join(hdr)] + [" | ".join(r) for r in rows] if rows else [])
        for ln in allines:
            if y > 780:
                page = doc.new_page(); y = 60
            page.insert_text((50, y), ln, fontsize=10, fontname="helv"); y += 16
        return doc.tobytes(), gold, body
    raise ValueError(fmt)


FMTS = ["txt", "csv", "xlsx", "docx", "pptx", "html", "md", "pdf"]


def build(seed: str) -> None:
    import shutil
    shutil.rmtree(OUT, ignore_errors=True); (OUT / "files").mkdir(parents=True); (OUT / "clean").mkdir()
    manifest = []
    jsonl = []
    # 1. prose documents x formats (every format gets prose docs; table docs for csv/xlsx/docx/pptx/html/md/pdf)
    k = 0
    for rep in range(3):
        for fmt in FMTS:
            g = G(f"{seed}-{fmt}-{rep}")
            text = make_doc(g, 18)
            data, gold, _ = write_files(g, k, text, None, fmt)
            name = f"prose_{fmt}_{rep}.{fmt}"; (OUT / "files" / name).write_bytes(data)
            manifest.append({"id": name, "fmt": fmt, "gold": gold, "kind": "prose"}); k += 1
            plain, spans = spans_of(text)
            for li, (ln_start, ln) in enumerate(_lines(plain)):
                pass
        for fmt in ("csv", "xlsx", "docx", "pptx", "html", "md", "pdf", "txt"):
            g = G(f"{seed}-tab-{fmt}-{rep}")
            hdr, rows = make_table(g, 6)
            data, gold, _ = write_files(g, k, "", (hdr, rows), fmt)
            name = f"table_{fmt}_{rep}.{fmt}"; (OUT / "files" / name).write_bytes(data)
            manifest.append({"id": name, "fmt": fmt, "gold": gold, "kind": "table"}); k += 1
    # 2. clean prose (false-positive check), also wrapped
    for i, c in enumerate(CLEAN):
        (OUT / "clean" / f"clean_{i}.txt").write_text(c, encoding="utf-8")
    # 3. detector-level jsonl: one item per template instance, dense gold
    g = G(f"{seed}-jsonl")
    T = line_templates(g)
    n = 0
    for rep in range(8):
        for t in T:
            raw = t()
            plain, spans = spans_of(raw)
            jsonl.append({"id": f"en-{n:04d}", "group": "en", "text": plain, "spans": spans, "neg": [], "ok": []}); n += 1
    TEXT_JSONL.write_text("\n".join(json.dumps(x, ensure_ascii=False) for x in jsonl) + "\n", encoding="utf-8")
    (OUT / "gold.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"built {len(manifest)} files, {sum(len(m['gold']) for m in manifest)} gold values, {len(jsonl)} jsonl items, {len(CLEAN)} clean docs -> {OUT}")


def _lines(s):
    pos = 0
    for ln in s.split("\n"):
        yield pos, ln; pos += len(ln) + 1


# -------------------------------------------------------------------------------------------------------------- runners
def one(path: Path, spec: dict) -> dict:
    """worker body (subprocess): anonymize one file, check leaks on output + restore equality"""
    from tests.helpers import Env
    import eval_real as ER
    data = path.read_bytes(); fmt = spec["fmt"]
    res = {"id": spec["id"], "fmt": fmt, "kind": spec["kind"], "gold": len(spec["gold"])}
    env = Env(**({"retain_secrets_for_restore": True} if fmt == "pdf" else {}))
    try:
        job = env.anonymize((path.name, data))
    except Exception as ex:                                        # noqa: BLE001
        res["crash"] = f"{type(ex).__name__}: {str(ex)[:120]}"; return res
    f = job.files[0]
    res["status"] = f.status
    cats = collections.defaultdict(lambda: [0, 0]); leaks = []
    if f.out_path:
        out = Path(f.out_path).read_bytes()
        view = ER.canon(ER.output_views(fmt, out))
        for cat, val in spec["gold"]:
            toks = [t for t in re.findall(r"[A-Za-z0-9'’]+", val)]
            if cat == "PERSON":
                nd = [ER.canon(max(toks, key=len))]
            elif cat in ("CITY",):
                nd = [ER.canon(val)]
            else:
                nd = [ER.canon(val)]
            leaked = any(n and n in view for n in nd)
            cats[cat][0] += 1; cats[cat][1] += leaked
            if leaked:
                leaks.append([cat, val])
        try:
            rj = env.restore(env.as_upload(job)); rf = rj.files[0]
            if rf.out_path:
                res["restore_equal"] = ER.restore_equal(fmt, data, Path(rf.out_path).read_bytes(), [x for c, v in spec["gold"] if c == "SECRET" for x in (v, v.replace("&", "&amp;"))])[1]
            else:
                res["restore_equal"] = False
        except Exception as ex:                                    # noqa: BLE001
            res["restore_equal"] = False; res["restore_err"] = str(ex)[:80]
    else:
        res["blocked"] = True
    res["cats"] = dict(cats); res["leaks"] = leaks; res["natural"] = job.result.get("total", None) if hasattr(job, "result") and isinstance(job.result, dict) else None
    env.close()
    return res


def run(jobs: int) -> None:
    from concurrent.futures import ThreadPoolExecutor
    man = json.loads((OUT / "gold.json").read_text())

    def go(spec):
        p = subprocess.run([sys.executable, "-I", str(Path(__file__).resolve()), "worker", spec["id"]], capture_output=True, timeout=600, cwd=ROOT, env={"PYTHONPATH": str(ROOT)})
        line = [x for x in p.stdout.decode("utf-8", "ignore").splitlines() if x.startswith("RESULT ")]
        return json.loads(line[-1][7:]) if line else {"id": spec["id"], "fmt": spec["fmt"], "kind": spec["kind"], "gold": len(spec["gold"]), "crash": p.stderr.decode("utf-8", "ignore")[-200:]}
    with ThreadPoolExecutor(max_workers=jobs) as ex:
        rows = list(ex.map(go, man))
    (OUT / "results.json").write_text(json.dumps(rows, ensure_ascii=False, indent=1), encoding="utf-8")
    report(rows)


def report(rows=None) -> None:
    rows = rows or json.loads((OUT / "results.json").read_text())
    per = collections.defaultdict(lambda: [0, 0]); fm = collections.defaultdict(lambda: [0, 0]); rest = [0, 0]
    for r in rows:
        for c, (n, l) in r.get("cats", {}).items():
            per[c][0] += n; per[c][1] += l
            fm[r["fmt"]][0] += n; fm[r["fmt"]][1] += l
        if "restore_equal" in r:
            rest[0] += 1; rest[1] += bool(r["restore_equal"])
        if r.get("crash") or r.get("blocked"):
            print("PROBLEM", r["id"], r.get("crash") or "blocked")
    print(f"{'kind':12}{'gold':>6}{'leak':>6}  recall")
    for c, (n, l) in sorted(per.items()):
        print(f"{c:12}{n:6}{l:6}  {1 - l / n:.3f}")
    print("per format:", {k: f"{v[1]}/{v[0]}" for k, v in sorted(fm.items())})
    print(f"restore equal: {rest[1]}/{rest[0]}")
    tot = sum(v[0] for v in per.values()); lk = sum(v[1] for v in per.values())
    print(f"TOTAL gold={tot} leaked={lk} recall={1 - lk / max(1, tot):.3f}")
    for r in rows:
        for c, v in r.get("leaks", []):
            print("LEAK", r["id"], c, v)


def clean() -> None:
    from tests.helpers import Env
    env = Env(); words = fp = 0; hits = []
    docs = list((OUT / "clean").glob("*.txt")) or None
    if not docs:
        print("run build first"); return
    texts = [p.read_text() for p in docs]
    for p in ROOT.glob("tests/eval/real/*.txt"):                     # real English prose too (ignored by git)
        try:
            texts.append(p.read_text(errors="ignore")[:20000])
        except OSError:
            pass
    for t in texts:
        det = env.service.build_detector(env.service.settings_for({}))
        det.harvest(t, "x.txt")
        fs = [f for f in det.scan(t, "x.txt", "text") if f.decision.value == "AUTO"]
        words += len(t.split()); fp += len(fs); hits += [(f.category, f.original) for f in fs]
    print(f"clean prose: {len(texts)} docs, {words} words, AUTO findings={fp} ({1000 * fp / max(1, words):.2f} per 1000 words)")
    print(collections.Counter(hits).most_common(40))


if __name__ == "__main__":
    if len(sys.argv) > 2 and sys.argv[1] == "worker":
        man = {m["id"]: m for m in json.loads((OUT / "gold.json").read_text())}
        print("RESULT " + json.dumps(one(OUT / "files" / sys.argv[2], man[sys.argv[2]]), ensure_ascii=False))
    else:
        ap = argparse.ArgumentParser(); ap.add_argument("cmd", choices=["build", "run", "report", "clean"]); ap.add_argument("--seed", default="7"); ap.add_argument("--jobs", type=int, default=2)
        a = ap.parse_args()
        {"build": lambda: build(a.seed), "run": lambda: run(a.jobs), "report": report, "clean": clean}[a.cmd]()
