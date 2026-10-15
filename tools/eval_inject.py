"""Dense injection of synthetic identities into real documents, plus synthetic complex documents (tools/eval_real.py helper).

inject(fmt, data, pack, traits) -> bytes     put dozens of gold values into awkward spots of an existing file
synth(kind, fmt, pack, seed)  -> bytes       build a long messy document (contract/invoice/resume/ruling/minutes/mail/log/...) in `fmt`
Spots per format: body, tables, headers/footers, comments, text boxes, hyperlinks, document properties, notes, hidden sheets,
numeric cells, annotations, rasterised banners and rotated stamps, values split over runs/lines.
"""
from __future__ import annotations

import csv
import io
import json
import random
import re
from xml.sax.saxutils import escape

FONT = "/System/Library/Fonts/Supplemental/Arial.ttf"

RU_P = ["Стороны договорились о нижеследующем и подтверждают, что обладают всеми необходимыми полномочиями для заключения настоящего соглашения.",
        "Исполнитель обязуется оказать услуги в объёме и в сроки, указанные в приложении, а Заказчик обязуется принять и оплатить их в порядке, установленном разделом 4.",
        "Любые изменения и дополнения действительны только при условии, что они совершены в письменной форме и подписаны уполномоченными представителями сторон.",
        "В случае просрочки платежа Заказчик уплачивает неустойку в размере 0,1 % от неоплаченной суммы за каждый день просрочки, но не более 10 % от цены договора.",
        "Споры, возникающие из настоящего договора, разрешаются путём переговоров, а при недостижении согласия — в Арбитражном суде по месту нахождения ответчика.",
        "Конфиденциальная информация не подлежит раскрытию третьим лицам без предварительного письменного согласия другой стороны в течение пяти лет.",
        "Суд, исследовав материалы дела и заслушав доводы лиц, участвующих в деле, пришёл к выводу, что исковые требования подлежат частичному удовлетворению.",
        "Согласно пункту 3 статьи 309 Гражданского кодекса обязательства должны исполняться надлежащим образом в соответствии с условиями обязательства.",
        "Повестка: 1) утверждение отчёта за квартал; 2) назначение ответственных; 3) разное. Решение принято единогласно, воздержавшихся нет.",
        "Образование: высшее, окончил(а) университет в 2009 году; опыт работы с 2010 года; владение английским языком — свободное; водительские права категории B."]
EN_P = ["The parties agree as follows and confirm that each of them has full authority to enter into this agreement on the terms set out below.",
        "The Supplier shall deliver the goods in the quantities and by the dates listed in Schedule A, and the Buyer shall pay the invoiced amounts within thirty days of receipt.",
        "No amendment is effective unless it is made in writing and signed by an authorised representative of each party.",
        "Late payments bear interest at 1.5 % per month or the maximum rate permitted by law, whichever is lower, from the due date until payment in full.",
        "The Court, having reviewed the record and heard the arguments of counsel, concludes that the motion should be granted in part and denied in part.",
        "Each party shall keep Confidential Information secret and use it solely for the purposes of this agreement for a period of five years after termination.",
        "Action items: (1) finalise the budget; (2) circulate the revised timeline; (3) schedule the follow-up review. Approved unanimously.",
        "Professional experience: eight years in operations and logistics; led a team of twelve; fluent in English and German; full driving licence."]


def blocks_for(kind: str, pack, seed: str) -> list[tuple]:
    r = random.Random(seed)
    P = RU_P if pack.lang == "ru" else EN_P
    ru = pack.lang == "ru"
    title = {"contract": "ДОГОВОР ОКАЗАНИЯ УСЛУГ № {n}" if ru else "SERVICES AGREEMENT No. {n}", "invoice": "СЧЁТ НА ОПЛАТУ № {n}" if ru else "INVOICE {n}",
             "resume": "РЕЗЮМЕ" if ru else "CURRICULUM VITAE", "ruling": "РЕШЕНИЕ № {n}" if ru else "OPINION OF THE COURT No. {n}",
             "minutes": "ПРОТОКОЛ СОБРАНИЯ № {n}" if ru else "MEETING MINUTES No. {n}", "mail": "Re: " + ("Реквизиты и согласование" if ru else "Details and approval")}[kind]
    b: list[tuple] = [("h", title.format(n=r.randrange(100, 9999)))]
    prose, lines = list(pack.prose), list(pack.lines)
    r.shuffle(lines)
    for i in range(14 if kind != "invoice" else 5):
        b.append(("p", f"{i + 1}. " + r.choice(P)))
        if prose and i % 3 == 1:
            b.append(("p", prose.pop(0)))
        if lines and i % 2 == 0:
            b.append(("p", lines.pop()))
    if kind == "invoice":
        hdr = ["№", "Наименование", "Кол-во", "Цена", "Сумма"] if ru else ["#", "Description", "Qty", "Unit price", "Amount"]
        rows = [[str(i + 1), f"{'Услуга' if ru else 'Item'} {r.choice(['A', 'B', 'C'])}-{r.randrange(100, 999)}", str(r.randrange(1, 40)), f"{r.randrange(100, 9999)}.{r.randrange(10, 99)}", f"{r.randrange(1000, 99999)}.{r.randrange(10, 99)}"] for i in range(45)]
        b.append(("t", hdr, rows))
    b.append(("t", *pack.rows))
    b += [("p", ln) for ln in lines]
    b.append(("p", ("Подписи сторон" if ru else "Signatures") + ": ______ / " + (pack.lines[6].split(": ", 1)[-1][:60])))
    return b


# ---------------------------------------------------------------- renderers for synthetic blocks
def _text_of(blocks, md=False) -> str:
    out = []
    for blk in blocks:
        if blk[0] == "h":
            out.append(("# " if md else "") + blk[1] + "\n")
        elif blk[0] == "p":
            out.append(blk[1] + "\n")
        else:
            if md:
                out.append("| " + " | ".join(blk[1]) + " |\n|" + "---|" * len(blk[1]))
                out += ["| " + " | ".join(r) + " |" for r in blk[2]]
            else:
                out.append("\t".join(blk[1])); out += ["\t".join(r) for r in blk[2]]
            out.append("")
    return "\n".join(out)


def synth(kind: str, fmt: str, pack, seed: str) -> bytes:
    b = blocks_for(kind, pack, seed)
    if fmt in ("txt", "md", "log"):
        txt = _text_of(b, md=fmt == "md")
        if fmt == "log":
            txt = "\n".join(f"2026-10-0{1 + i % 8} 1{i % 10}:{i % 6}{i % 9}:0{i % 7} INFO  {ln}" for i, ln in enumerate(txt.splitlines()) if ln.strip())
        return txt.encode("utf-8")
    if fmt in ("html", "htm"):
        body = "".join(f"<h1>{escape(x[1])}</h1>" if x[0] == "h" else f"<p>{escape(x[1])}</p>" if x[0] == "p" else
                       "<table><tr>" + "".join(f"<th>{escape(c)}</th>" for c in x[1]) + "</tr>" + "".join("<tr>" + "".join(f"<td>{escape(c)}</td>" for c in rw) + "</tr>" for rw in x[2]) + "</table>" for x in b)
        return f"<!doctype html><html lang='{pack.lang}'><head><meta charset='utf-8'><title>{escape(b[0][1])}</title></head><body>{body}</body></html>".encode()
    if fmt in ("csv", "tsv"):
        buf = io.StringIO(); w = csv.writer(buf, delimiter="," if fmt == "csv" else "\t")
        for x in b:
            if x[0] == "t":
                w.writerow(x[1]); w.writerows(x[2])
            else:
                w.writerow([x[1]])
        return buf.getvalue().encode("utf-8")
    if fmt == "json":
        return json.dumps({"title": b[0][1], "paragraphs": [x[1] for x in b if x[0] == "p"], "tables": [{"header": x[1], "rows": x[2]} for x in b if x[0] == "t"]}, ensure_ascii=False, indent=1).encode()
    if fmt == "xml":
        body = "".join(f"<para>{escape(x[1])}</para>" if x[0] != "t" else "<table>" + "".join("<row>" + "".join(f"<cell>{escape(c)}</cell>" for c in rw) + "</row>" for rw in x[2]) + "</table>" for x in b)
        return f"<?xml version='1.0' encoding='utf-8'?><document>{body}</document>".encode()
    if fmt in ("yaml", "yml"):
        return ("title: " + json.dumps(b[0][1], ensure_ascii=False) + "\nitems:\n" + "".join(f"  - {json.dumps(x[1], ensure_ascii=False)}\n" for x in b if x[0] == "p")).encode()
    if fmt == "ini":
        return ("[document]\ntitle = " + b[0][1] + "\n[body]\n" + "".join(f"line{i} = {x[1]}\n" for i, x in enumerate(y for y in b if y[0] == "p"))).encode()
    if fmt == "docx":
        from docx import Document
        d = Document()
        for x in b:
            if x[0] == "h":
                d.add_heading(x[1], 1)
            elif x[0] == "p":
                d.add_paragraph(x[1])
            else:
                t = d.add_table(rows=1, cols=len(x[1])); t.style = "Table Grid"
                for i, c in enumerate(x[1]):
                    t.rows[0].cells[i].text = c
                for rw in x[2]:
                    cells = t.add_row().cells
                    for i, c in enumerate(rw):
                        cells[i].text = c
        buf = io.BytesIO(); d.save(buf); return buf.getvalue()
    if fmt == "xlsx":
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active; ws.title = "Документ" if pack.lang == "ru" else "Document"
        for x in b:
            if x[0] == "t":
                ws.append(x[1]); [ws.append(rw) for rw in x[2]]
            else:
                ws.append([x[1]])
        buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
    if fmt == "pptx":
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation(); ps = [x for x in b if x[0] != "t"]
        for k in range(0, len(ps), 4):
            s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = ps[k][1][:80]
            for j, x in enumerate(ps[k + 1:k + 4]):
                s.shapes.add_textbox(Inches(.5), Inches(1.5 + j * 1.6), Inches(9), Inches(1.5)).text_frame.text = x[1]
        buf = io.BytesIO(); prs.save(buf); return buf.getvalue()
    if fmt in ("pdf", "pdf_scan"):
        import fitz
        doc = fitz.open(); page = None; y = 800
        for x in b:
            lines = [x[1]] if x[0] != "t" else ["  |  ".join(x[1])] + ["  |  ".join(rw) for rw in x[2]]
            for ln in lines:
                n = max(1, len(ln) // 95 + 1)
                if page is None or y - 12 * n < 50:
                    page = doc.new_page(); y = 800
                page.insert_textbox(fitz.Rect(40, 842 - y, 555, 842 - y + 12 * n + 4), ln, fontname="arial", fontfile=FONT, fontsize=9)
                y -= 12 * n + 3
        if fmt == "pdf_scan":
            out = fitz.open()
            for pg in doc:
                pix = pg.get_pixmap(dpi=170); q = out.new_page(width=595, height=842); q.insert_image(q.rect, stream=pix.tobytes("png"))
            return out.tobytes()
        return doc.tobytes()
    if fmt in ("png", "jpg", "tiff", "webp", "bmp"):
        return render_page_image(b, fmt, seed)
    raise ValueError(fmt)


def render_page_image(blocks, fmt: str, seed: str) -> bytes:
    from PIL import Image, ImageDraw, ImageFont, ImageFilter
    r = random.Random(seed)
    f = ImageFont.truetype(FONT, 22)
    lines = []
    for x in blocks:
        if x[0] == "t":
            lines += ["  |  ".join(x[1])] + ["  |  ".join(rw) for rw in x[2][:10]]
        else:
            t = x[1]
            while len(t) > 78:
                cut = t.rfind(" ", 0, 78); cut = cut if cut > 0 else 78
                lines.append(t[:cut]); t = t[cut:].strip()
            lines.append(t)
    lines = lines[:44]
    img = Image.new("RGB", (1240, 60 + 28 * len(lines)), (250, 248, 240)); d = ImageDraw.Draw(img)
    for i, ln in enumerate(lines):
        d.text((40, 30 + 28 * i), ln, font=f, fill=(25, 25, 30))
    for _ in range(300):
        d.point((r.randrange(img.width), r.randrange(img.height)), fill=(200, 200, 195))
    img = img.rotate(r.uniform(-1.2, 1.2), expand=True, fillcolor=(250, 248, 240)).filter(ImageFilter.GaussianBlur(0.4))
    buf = io.BytesIO(); img.save(buf, {"jpg": "JPEG", "tiff": "TIFF", "webp": "WEBP", "bmp": "BMP", "png": "PNG"}[fmt]); return buf.getvalue()


# ---------------------------------------------------------------- banners
def banner(lines: list[str], width: int = 1000, size: int = 26, rotate: float = 0.0):
    from PIL import Image, ImageDraw, ImageFont
    f = ImageFont.truetype(FONT, size)
    img = Image.new("RGB", (max(width, 900), 40 + int(size * 1.7) * len(lines)), "white")
    d = ImageDraw.Draw(img)
    for i, ln in enumerate(lines):
        d.text((24, 20 + i * int(size * 1.7)), ln, font=f, fill=(15, 15, 15))
    return img.rotate(rotate, expand=True, fillcolor="white") if rotate else img


def _png(img) -> bytes:
    b = io.BytesIO(); img.save(b, "PNG"); return b.getvalue()


# ---------------------------------------------------------------- injection into existing files
def inject(fmt: str, data: bytes, pack, traits: str = "") -> bytes:
    L, M = pack.lines, pack.meta
    hdr, rows = pack.rows
    if fmt == "pdf":
        import fitz
        doc = fitz.open(stream=data, filetype="pdf")
        scanned = "scanned" in traits
        if scanned:
            for chunk, rot in ((L[:12], 0), (L[12:], 0)):
                img = banner(chunk, 900, 22); page = doc.new_page(width=595, height=842)
                page.insert_image(fitz.Rect(30, 30, 565, 30 + 535 * img.height / img.width), stream=_png(img))
            stamp = banner(L[6:9], 700, 24, rotate=7)
            page.insert_image(fitz.Rect(60, 600, 540, 600 + 480 * stamp.height / stamp.width), stream=_png(stamp))
        else:
            page = doc.new_page(width=595, height=842)
            page.insert_textbox(fitz.Rect(40, 40, 555, 520), "\n".join(L[:14]), fontname="arial", fontfile=FONT, fontsize=10)
            y = 540
            for rw in [hdr] + rows:
                page.insert_text((40, y), "   ".join(c[:28] for c in rw[:4]), fontname="arial", fontfile=FONT, fontsize=8.5); y += 13
            page.insert_textbox(fitz.Rect(40, 640, 555, 800), "\n".join(L[14:]), fontname="arial", fontfile=FONT, fontsize=10)
            p0 = doc[0]
            for i, ln in enumerate(L[:3]):
                p0.insert_text((30, p0.rect.height - 18 - (2 - i) * 11), ln, fontname="arial", fontfile=FONT, fontsize=8)
            try:
                a = p0.add_text_annot(fitz.Point(50, 50), "; ".join(L[3:5])); a.set_info(title=M["author"], subject=L[2][:60]); a.update()
            except Exception:                                               # noqa: BLE001
                pass
        doc.set_metadata({"author": M["author"], "title": M["title"], "subject": M["subject"], "keywords": M["keywords"]})
        return doc.tobytes()
    if fmt in ("png", "jpg", "tiff", "webp", "bmp"):
        from PIL import Image
        base = Image.open(io.BytesIO(data)).convert("RGB")
        ban = banner(L[:13], base.width, 24); st = banner(L[13:17], 760, 22, rotate=-6); sm = banner(L[17:], 900, 16)
        W = max(base.width, ban.width, st.width, sm.width)
        out = Image.new("RGB", (W, base.height + ban.height + st.height + sm.height), "white")
        y = 0
        for part in (base, ban, st, sm):
            out.paste(part, (0, y)); y += part.height
        buf = io.BytesIO()
        out.save(buf, {"jpg": "JPEG", "png": "PNG", "tiff": "TIFF", "webp": "WEBP", "bmp": "BMP"}[fmt], **({"quality": 90} if fmt in ("jpg", "webp") else {}))
        return buf.getvalue()
    if fmt == "docx":
        from docx import Document
        from docx.oxml import parse_xml
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.oxml.ns import qn
        doc = Document(io.BytesIO(data))
        sec = doc.sections[0]
        sec.header.paragraphs[0].text = L[0]; sec.footer.paragraphs[0].text = L[1] + " | " + L[3]
        sec.header.add_paragraph(L[2])
        for ln in L[4:12]:
            doc.add_paragraph(ln)
        p = doc.add_paragraph()                                              # a value split over runs with different formatting
        for k, ln in enumerate(L[12:15]):
            cut = max(3, len(ln) // 2)
            p.add_run(ln[:cut]).bold = bool(k % 2); p.add_run(ln[cut:]).italic = True; p.add_run().add_break()
        t = doc.add_table(rows=1, cols=len(hdr)); t.style = "Table Grid" if "Table Grid" in [s.name for s in doc.styles] else None
        for i, c in enumerate(hdr):
            t.rows[0].cells[i].text = c
        for rw in rows:
            cells = t.add_row().cells
            for i, c in enumerate(rw):
                cells[i].text = c
        for ln in L[15:]:
            doc.add_paragraph(ln)
        for pr in pack.prose:
            doc.add_paragraph(pr)
        email = next(v for k, v in pack.gold if k == "EMAIL")
        hp = doc.add_paragraph("Contact: ")
        rid = doc.part.relate_to("mailto:" + email, RT.HYPERLINK, is_external=True)
        hp._p.append(parse_xml(f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="{rid}"><w:r><w:t>{escape(email)}</w:t></w:r></w:hyperlink>'))
        tb = doc.add_paragraph()
        tb._p.append(parse_xml('<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape style="width:300pt;height:60pt"><v:textbox><w:txbxContent>'
                               f'<w:p><w:r><w:t>{escape(L[5])}</w:t></w:r></w:p><w:p><w:r><w:t>{escape(L[8])}</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>'))
        try:
            doc.add_comment(runs=[doc.paragraphs[-1].runs[0]] if doc.paragraphs[-1].runs else [p.runs[0]], text=f"{L[6]}; {L[7]}", author=M["author"], initials="EN")
        except Exception:                                                   # noqa: BLE001
            pass
        cp = doc.core_properties
        cp.author, cp.title, cp.subject, cp.keywords, cp.comments, cp.last_modified_by = M["author"], M["title"], M["subject"], M["keywords"], M["comments"], L[1][:60]
        buf = io.BytesIO(); doc.save(buf); return buf.getvalue()
    if fmt == "xlsx":
        from openpyxl import load_workbook
        from openpyxl.comments import Comment
        wb = load_workbook(io.BytesIO(data)); ws = wb.active
        row = ws.max_row + 2
        for i, ln in enumerate(L[:10]):
            ws.cell(row + i, 1, ln)
        c = ws.cell(row, 3, L[10]); c.comment = Comment(L[11], M["author"])
        ws.cell(row + 1, 3, f"{L[12]}\n{L[13]}")
        ws.merge_cells(start_row=row + 3, start_column=3, end_row=row + 3, end_column=6); ws.cell(row + 3, 3, L[14])
        try:
            ws.oddHeader.center.text = L[2][:70]; ws.oddFooter.left.text = L[3][:70]
        except Exception:                                                   # noqa: BLE001
            pass
        ct = wb.create_sheet("Contacts"); ct.append(hdr)
        for rw in rows:
            ct.append([int(v) if re.fullmatch(r"\d{10,12}", v) else v for v in rw])
        for k, v in pack.gold:
            if k == "INN":
                ct.append(["INN", int(v)])
        hid = wb.create_sheet("Служебный"); hid.sheet_state = "hidden"
        for ln in L[15:]:
            hid.append([ln])
        for pr in pack.prose:
            hid.append([pr])
        wb.properties.creator, wb.properties.title, wb.properties.subject, wb.properties.keywords, wb.properties.description = M["author"], M["title"], M["subject"], M["keywords"], M["comments"]
        buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
    if fmt == "pptx":
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation(io.BytesIO(data)); lay = prs.slide_layouts[min(6, len(prs.slide_layouts) - 1)]
        s = prs.slides.add_slide(lay)
        for j, ln in enumerate(L[:11]):
            s.shapes.add_textbox(Inches(.4), Inches(.2 + j * .62), Inches(9), Inches(.5)).text_frame.text = ln
        s.notes_slide.notes_text_frame.text = "\n".join(L[11:14])
        s2 = prs.slides.add_slide(lay)
        tbl = s2.shapes.add_table(len(rows) + 1, len(hdr), Inches(.3), Inches(.4), Inches(9.4), Inches(3)).table
        for i, c in enumerate(hdr):
            tbl.cell(0, i).text = c
        for ri, rw in enumerate(rows, 1):
            for i, c in enumerate(rw):
                tbl.cell(ri, i).text = c
        s2.notes_slide.notes_text_frame.text = L[14]
        s3 = prs.slides.add_slide(lay)
        for j, ln in enumerate(L[15:] + pack.prose):
            s3.shapes.add_textbox(Inches(.4), Inches(.2 + j * .55), Inches(9), Inches(.5)).text_frame.text = ln
        s4 = prs.slides.add_slide(lay)
        s4.shapes.add_picture(io.BytesIO(_png(banner(L[3:9], 900, 26))), Inches(.4), Inches(1), width=Inches(9))
        cp = prs.core_properties
        cp.author, cp.title, cp.subject, cp.keywords, cp.comments = M["author"], M["title"], M["subject"], M["keywords"], M["comments"]
        buf = io.BytesIO(); prs.save(buf); return buf.getvalue()
    # ---- text formats
    text = data.decode("utf-8", "replace")
    nl = "" if text.endswith("\n") else "\n"
    if fmt in ("txt", "md", "log", "ini", "yaml", "yml"):
        if fmt == "md":
            blk = "\n\n## Contacts\n\n" + "\n".join(f"- **{ln}**" for ln in L[:8]) + "\n\n| " + " | ".join(hdr) + " |\n|" + "---|" * len(hdr) + "\n" + "\n".join("| " + " | ".join(r) + " |" for r in rows) + \
                "\n\n" + "\n".join(L[8:]) + "\n\n```\n" + "\n".join(L[-3:]) + "\n```\n" + "\n".join(pack.prose) + f"\n\n[mail](mailto:{next(v for k, v in pack.gold if k == 'EMAIL')})\n"
        elif fmt == "log":
            blk = "\n".join(f"2026-10-08 12:{i:02d}:07 INFO user action {ln}" for i, ln in enumerate(L)) + "\n" + "\n".join(f"2026-10-08 13:0{i}:11 WARN {p}" for i, p in enumerate(pack.prose)) + "\n"
        elif fmt == "ini":
            blk = "\n[contacts]\n" + "\n".join(f"line{i} = {ln}" for i, ln in enumerate(L)) + "\n[secrets]\n" + "\n".join(f"note{i} = {p}" for i, p in enumerate(pack.prose)) + "\n"
        elif fmt in ("yaml", "yml"):
            blk = "\ncontacts:\n" + "".join(f"  - {json.dumps(ln, ensure_ascii=False)}\n" for ln in L) + "notes: |\n" + "".join(f"  {p}\n" for p in pack.prose)
        else:
            blk = "\n\n" + "\n".join(L) + "\n\n" + "\t".join(hdr) + "\n" + "\n".join("\t".join(r) for r in rows) + "\n\n" + "\n\n".join(pack.prose) + "\n"
        return (text + nl + blk).encode("utf-8")
    if fmt in ("csv", "tsv"):
        buf = io.StringIO(); w = csv.writer(buf, delimiter="," if fmt == "csv" else "\t", lineterminator="\n")
        w.writerow([]); w.writerow(hdr); w.writerows(rows)
        for ln in L:
            w.writerow([ln, ln.split(":")[-1].strip()])
        w.writerow(["Комментарий" if pack.lang == "ru" else "Comment", pack.prose[0] + "\n" + pack.prose[1]])
        return (text + nl + buf.getvalue()).encode("utf-8")
    if fmt in ("html", "htm"):
        tbl = "<table><tr>" + "".join(f"<th>{escape(c)}</th>" for c in hdr) + "</tr>" + "".join("<tr>" + "".join(f"<td>{escape(c)}</td>" for c in r) + "</tr>" for r in rows) + "</table>"
        em = next(v for k, v in pack.gold if k == "EMAIL"); ph = next(v for k, v in pack.gold if k == "PHONE")
        # No <img> here: a reference to a missing image is (correctly) an unverifiable image for the product and blocks the case,
        # which made every html case 'blocked' without measuring leaks. The attribute value is kept on a non-image element.
        blk = (f"<!-- {escape(L[4])} -->" + "".join(f"<p>{escape(x)}</p>" for x in L) + tbl + "".join(f"<p>{escape(x)}</p>" for x in pack.prose) +
               f'<a href="mailto:{em}" title="{escape(L[6])}">{em}</a> <a href="tel:{re.sub(r"[^+0-9]", "", ph)}">{escape(ph)}</a><span role="img" aria-label="{escape(L[7])}"></span>'
               f'<div data-owner="{escape(L[8])}" hidden>{escape(L[9])}</div><script type="application/ld+json">{{"name":{json.dumps(M["author"], ensure_ascii=False)},"email":"{em}"}}</script>')
        meta = f'<meta name="author" content="{escape(M["author"])}"><meta name="keywords" content="{escape(M["keywords"])}">'
        s = text.replace("</head>", meta + "</head>", 1) if "</head>" in text else meta + text
        return (s.replace("</body>", blk + "</body>", 1) if "</body>" in s else s + blk).encode("utf-8")
    if fmt == "json":
        d = json.loads(data)
        out = {"results": d, "contacts": [dict(zip(hdr, r)) for r in rows], "notes": L, "paragraphs": pack.prose, "owner": {"name": M["author"], "meta": M},
               "credentials": [v for k, v in pack.gold if k == "SECRET"]}
        return json.dumps(out, ensure_ascii=False, indent=1).encode("utf-8")
    if fmt == "xml":
        blk = (f"<!-- {escape(L[0])} -->" + "<contacts>" + "".join("<person " + " ".join(f'f{i}="{escape(c)}"' for i, c in enumerate(r)) + f"><name>{escape(r[0])}</name></person>" for r in rows) + "</contacts>" +
               "".join(f"<note>{escape(x)}</note>" for x in L) + f"<raw><![CDATA[{' | '.join(pack.prose)}]]></raw>")
        k = text.rfind("</")
        return (text[:k] + blk + text[k:] if k > 0 else text + blk).encode("utf-8")
    raise ValueError(fmt)
