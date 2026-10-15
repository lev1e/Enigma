"""OCR: скан и изображения внутри PDF обезличиваются по пикселям, метки видны текстом, возврат восстанавливает значения."""
from __future__ import annotations

import unittest
from pathlib import Path

import fitz

from anonymizer import ocr
from tests import pdf_probe as P
from tests.helpers import Env
from tests.test_pdf_audit import page_doc

READY = Path(P.FONT).exists() and ocr.available()


def scan_of(src: fitz.Document, dpi=200) -> fitz.Document:
    scan = fitz.open()
    for page in src:
        out = scan.new_page(width=page.rect.width, height=page.rect.height)
        out.insert_image(out.rect, pixmap=page.get_pixmap(dpi=dpi))
    return scan


def ocr_text(data: bytes) -> str:
    doc = fitz.open(stream=data, filetype="pdf")
    words = []
    for page in doc:
        found, _ = ocr.ocr_page_words(page)
        words.extend(w.text for w in found)
    return " ".join(words)


@unittest.skipUnless(READY, "нужны Arial и движок распознавания")
class PdfOcr(unittest.TestCase):
    def setUp(self):
        self.env = Env()
        self.addCleanup(self.env.close)

    def run_job(self, doc):
        job = self.env.anonymize(("scan.pdf", doc.tobytes()))
        f = job.files[0]
        return job, f, (Path(f.out_path).read_bytes() if f.out_path else None)

    def test_scan_is_anonymized_in_pixels(self):
        src = page_doc(P.LINE, "Телефон: " + P.PHONE)
        job, f, out = self.run_job(scan_of(src))
        self.assertIsNotNone(out, [m["text"] for m in f.messages])
        seen = ocr_text(out)
        for secret in ("Иванов", "Петрович", P.INN, "ivan.petrov"):
            self.assertNotIn(secret, seen)
        text = "\n".join(p.get_text() for p in fitz.open(stream=out, filetype="pdf"))
        self.assertNotIn(P.INN, text)
        self.assertRegex(text, r"Name\d|Имя\d|\w+\d")

    def test_scan_round_trip(self):
        src = page_doc(P.LINE)
        scan = scan_of(src)
        job, f, out = self.run_job(scan)
        self.assertIsNotNone(out)
        rj = self.env.restore(self.env.as_upload(job))
        rf = rj.files[0]
        self.assertTrue(rf.out_path, [m["text"] for m in rf.messages])
        # Страница-скан возвращается пикселями (а не дорисованным текстом): картинка равна исходной.
        def picture(data):
            d = fitz.open(stream=data, filetype="pdf")
            return fitz.Pixmap(d.extract_image(d[0].get_images(full=True)[0][0])["image"]).samples
        self.assertEqual(picture(Path(rf.out_path).read_bytes()), picture(scan.tobytes()))

    def test_picture_inside_text_page(self):
        doc = page_doc("Обычный текст страницы без данных.")
        small = page_doc(P.LINE)[0].get_pixmap(dpi=200, clip=fitz.Rect(60, 80, 560, 115))
        doc[0].insert_image(fitz.Rect(72, 300, 522, 331), pixmap=small)
        job, f, out = self.run_job(doc)
        self.assertIsNotNone(out, [m["text"] for m in f.messages])
        seen = ocr_text(out)
        self.assertNotIn(P.INN, seen)
        self.assertNotIn("ivan.petrov", seen)


if __name__ == "__main__":
    unittest.main()


def picture_png() -> bytes:
    src = page_doc(P.LINE)
    return src[0].get_pixmap(dpi=200, clip=fitz.Rect(60, 80, 560, 115)).tobytes("png")


def docx_with_picture(png: bytes) -> bytes:
    import io
    import docx
    d = docx.Document()
    d.add_paragraph("Обычный абзац без данных.")
    d.add_picture(io.BytesIO(png))
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def media_of(data: bytes) -> dict[str, bytes]:
    import io
    import zipfile
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist() if "/media/" in n}


@unittest.skipUnless(READY, "нужны Arial и движок распознавания")
class OfficeImageOcr(unittest.TestCase):
    def setUp(self):
        self.env = Env()
        self.addCleanup(self.env.close)

    def test_docx_picture_is_redacted_and_restored_exactly(self):
        png = picture_png()
        job = self.env.anonymize(("pic.docx", docx_with_picture(png)))
        f = job.files[0]
        self.assertTrue(f.out_path, [m["text"] for m in f.messages])
        out = Path(f.out_path).read_bytes()
        images = media_of(out)
        self.assertEqual(len(images), 1)
        (new,) = images.values()
        self.assertNotEqual(new, png)
        seen = " ".join(w.text for w in ocr.ocr_page_words(self._page(new))[0])
        for secret in (P.INN, "ivan.petrov", "Петрович"):
            self.assertNotIn(secret, seen)
        rj = self.env.restore(self.env.as_upload(job))
        rf = rj.files[0]
        self.assertTrue(rf.out_path, [m["text"] for m in rf.messages])
        (back,) = media_of(Path(rf.out_path).read_bytes()).values()
        self.assertEqual(back, png)

    @staticmethod
    def _page(data: bytes):
        doc = fitz.open()
        pix = fitz.Pixmap(data)
        page = doc.new_page(width=pix.width, height=pix.height)
        page.insert_image(page.rect, stream=data)
        return page

    def test_picture_without_data_is_untouched(self):
        d = page_doc("Просто текст без данных, ничего личного.")
        png = d[0].get_pixmap(dpi=200, clip=fitz.Rect(60, 80, 400, 115)).tobytes("png")
        job = self.env.anonymize(("pic.docx", docx_with_picture(png)), options={"image_mode": "redact_text"})
        f = job.files[0]
        self.assertTrue(f.out_path)
        (kept,) = media_of(Path(f.out_path).read_bytes()).values()
        self.assertEqual(kept, png)


def png_with_exif_like_chunks(png: bytes) -> bytes:
    import struct
    import zlib
    chunk = b"tEXt" + b"Author\x00Ivan Petrov GPS 55.7,37.6"
    block = struct.pack(">I", len(chunk) - 4) + chunk + struct.pack(">I", zlib.crc32(chunk))
    return png[:33] + block + png[33:]          # после IHDR (8 + 25 байт)


@unittest.skipUnless(READY, "нужны Arial и движок распознавания")
class StandaloneImages(unittest.TestCase):
    def setUp(self):
        self.env = Env()
        self.addCleanup(self.env.close)

    def _seen(self, data):
        return " ".join(w.text for w in ocr.ocr_page_words(OfficeImageOcr._page(data))[0])

    def test_png_redacted_and_restored_exactly(self):
        png = picture_png()
        job = self.env.anonymize(("scan.png", png))
        f = job.files[0]
        self.assertTrue(f.out_path, [m["text"] for m in f.messages])
        out = Path(f.out_path).read_bytes()
        self.assertTrue(out.startswith(b"\x89PNG"))
        self.assertNotEqual(out, png)
        self.assertEqual(fitz.Pixmap(out).width, fitz.Pixmap(png).width)
        seen = self._seen(out)
        for secret in (P.INN, "ivan.petrov", "Петрович"):
            self.assertNotIn(secret, seen)
        rf = self.env.restore(self.env.as_upload(job)).files[0]
        self.assertTrue(rf.out_path, [m["text"] for m in rf.messages])
        self.assertEqual(Path(rf.out_path).read_bytes(), png)

    def test_jpeg_redacted_and_restored_exactly(self):
        jpg = fitz.Pixmap(picture_png()).tobytes("jpeg", jpg_quality=95)
        job = self.env.anonymize(("scan.jpg", jpg))
        f = job.files[0]
        self.assertTrue(f.out_path, [m["text"] for m in f.messages])
        out = Path(f.out_path).read_bytes()
        self.assertEqual(out[:3], b"\xff\xd8\xff")
        self.assertNotIn(P.INN, self._seen(out))
        rf = self.env.restore(self.env.as_upload(job)).files[0]
        self.assertEqual(Path(rf.out_path).read_bytes(), jpg)

    def test_picture_without_data_keeps_pixels_drops_metadata(self):
        d = page_doc("Просто текст без данных, ничего личного.")
        png = d[0].get_pixmap(dpi=200, clip=fitz.Rect(60, 80, 400, 115)).tobytes("png")
        dirty = png_with_exif_like_chunks(png)
        job = self.env.anonymize(("plain.png", dirty), options={"image_mode": "redact_text"})
        f = job.files[0]
        out = Path(f.out_path).read_bytes()
        self.assertNotIn(b"Ivan Petrov", out)
        self.assertEqual(fitz.Pixmap(out).samples, fitz.Pixmap(png).samples)
        rf = self.env.restore(self.env.as_upload(job)).files[0]
        self.assertEqual(Path(rf.out_path).read_bytes(), dirty)

    def test_fake_image_is_rejected_not_copied(self):
        job = self.env.anonymize(("fake.png", b"\x89PNG not really"))
        f = job.files[0]
        self.assertFalse(f.out_path)
        self.assertEqual(f.status, "error")

    def test_jpeg_metadata_segments_removed(self):
        from anonymizer.formats import strip_image_metadata
        jpg = fitz.Pixmap(picture_png()).tobytes("jpeg")
        exif = b"Exif\x00\x00GPS-SECRET"
        seg = b"\xff\xe1" + (len(exif) + 2).to_bytes(2, "big") + exif
        dirty = jpg[:2] + seg + jpg[2:]
        clean = strip_image_metadata(dirty, ".jpg")
        self.assertNotIn(b"GPS-SECRET", clean)
        self.assertEqual(fitz.Pixmap(clean).samples, fitz.Pixmap(jpg).samples)


class OfficeImageMetadata(unittest.TestCase):
    def test_docx_picture_loses_png_text_chunks_and_restores(self):
        env = Env()
        self.addCleanup(env.close)
        d = page_doc("Просто текст без данных, ничего личного.")
        png = d[0].get_pixmap(dpi=100, clip=fitz.Rect(60, 80, 400, 115)).tobytes("png")
        dirty = png_with_exif_like_chunks(png)
        job = env.anonymize(("pic.docx", docx_with_picture(dirty)))
        f = job.files[0]
        self.assertTrue(f.out_path, [m["text"] for m in f.messages])
        (new,) = media_of(Path(f.out_path).read_bytes()).values()
        self.assertNotIn(b"Ivan Petrov", new)
        rf = env.restore(env.as_upload(job)).files[0]
        (back,) = media_of(Path(rf.out_path).read_bytes()).values()
        self.assertEqual(back, dirty)


class LatinInnLabel(unittest.TestCase):
    def test_inn_with_latin_label_is_hidden(self):
        env = Env()
        self.addCleanup(env.close)
        job = env.anonymize(("a.txt", "inn = 7707083893\nINN/KPP 7707083893/770701001\n"))
        out = Path(job.files[0].out_path).read_text(encoding="utf-8")
        self.assertNotIn("7707083893", out)
        self.assertNotIn("770701001", out)


class WindowsOcrMerge(unittest.TestCase):
    """Слияние слов русского и английского движков Windows (сам движок на этой машине недоступен)."""

    def test_mixed_script_word_loses_to_pure_one(self):
        rows = [
            {"l": "ru", "t": "Иvаnоv", "x0": 10, "y0": 10, "x1": 80, "y1": 30},
            {"l": "en", "t": "Ivanov", "x0": 11, "y0": 10, "x1": 81, "y1": 30},
            {"l": "ru", "t": "Привет", "x0": 100, "y0": 10, "x1": 160, "y1": 30},
            {"l": "en", "t": "Hpnaet", "x0": 101, "y0": 11, "x1": 160, "y1": 30},
            {"l": "ru", "t": "ИНН", "x0": 10, "y0": 60, "x1": 50, "y1": 80},
        ]
        words = ocr.merge_language_words(rows)
        self.assertEqual([w.text for w in words], ["Ivanov", "Привет", "ИНН"])
        self.assertEqual(words[0].line, words[1].line)
        self.assertNotEqual(words[0].line, words[2].line)

    def test_unavailable_off_windows(self):
        import sys
        if sys.platform != "win32":
            self.assertFalse(ocr.WindowsProvider().available())
