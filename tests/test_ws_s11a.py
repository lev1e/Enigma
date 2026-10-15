"""S11A: file preview endpoint for the chat side panel (data only, nothing executes, no original in the anonymized tab)."""
import http.client
import io
import json
import re
import threading
import unittest
from pathlib import Path

from anonymizer import preview
from anonymizer.webapp import create_server

from .helpers import Env

WEB = Path(__file__).resolve().parents[1] / "anonymizer" / "web"
ORIG_SECRET = "Петровский"
ANON_TEXT = "Договор с Name1, email Email1.\nСрок 10 дней.\n"
ORIG_TEXT = f"Договор с {ORIG_SECRET}, email ivan@secret-test.org.\nСрок 10 дней.\n"


def flat(value):
    return json.dumps(value, ensure_ascii=False)


class S11A(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.env = Env()
        cls.server, cls.app = create_server(cls.env.vault, cls.env.base / "work")
        cls.port = cls.server.server_address[1]
        threading.Thread(target=cls.server.serve_forever, daemon=True).start()
        cls.cookie = f"anon_t={cls.app.token}"
        cls.chat = cls.app.chats.create("s11a", keep_originals=True)["id"]

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown(); cls.server.server_close(); cls.env.close()

    def call(self, method, path, payload=None):
        conn = http.client.HTTPConnection("127.0.0.1", self.port, timeout=30)
        head = {"Host": f"127.0.0.1:{self.port}", "Cookie": self.cookie, "X-Requested-With": "anonymizer",
                "Content-Type": "application/json"}
        conn.request(method, path, json.dumps(payload).encode() if payload is not None else None, head)
        resp = conn.getresponse(); data = resp.read(); conn.close()
        return resp.status, (json.loads(data) if data else None)

    def add(self, name, anon, orig=None):
        rec = self.app.chats.add_file(self.chat, name, "text", anon=anon, orig=orig, names={"anon": name})
        return rec["id"]

    def get(self, fid, role="anon", page=None):
        suffix = f"/{page}" if page else ""
        return self.call("GET", f"/api/chats/{self.chat}/files/{fid}/{role}/preview{suffix}")

    def test_text_numbered_with_marks_and_no_original(self):
        fid = self.add("a.txt", ANON_TEXT.encode(), ORIG_TEXT.encode())
        status, data = self.get(fid)
        self.assertEqual(status, 200)
        self.assertEqual(data["kind"], "text")
        self.assertEqual(data["total"], 2)
        self.assertNotIn(ORIG_SECRET, flat(data))
        self.assertNotIn("secret-test", flat(data))
        marks = [seg[0] for line in data["lines"] if isinstance(line, list) for seg in line if seg[1]]
        self.assertIn("Name1", marks)
        self.assertFalse(data["local_only"])

    def test_original_tab_is_local_and_unmarked(self):
        fid = self.add("o.txt", ANON_TEXT.encode(), ORIG_TEXT.encode())
        status, data = self.get(fid, "orig")
        self.assertEqual(status, 200)
        self.assertTrue(data["local_only"])
        self.assertIn(ORIG_SECRET, flat(data))
        self.assertTrue(all(isinstance(line, str) for line in data["lines"]))

    def test_html_returned_as_plain_text_never_markup(self):
        html = "<html><script>alert(1)</script><img src=x onerror=alert(2)>Name1</html>"
        fid = self.add("p.html", html.encode())
        _, data = self.get(fid)
        self.assertEqual(data["kind"], "text")
        self.assertIn("<script>", flat(data))          # shown as source text

    def test_svg_is_source_not_image(self):
        fid = self.add("v.svg", b"<svg onload='alert(1)'><text>Name1</text></svg>")
        _, data = self.get(fid)
        self.assertEqual(data["kind"], "text")

    def test_csv_table_limit(self):
        rows = "\n".join(f"{i};Name{i};x" for i in range(300))
        fid = self.add("t.csv", rows.encode())
        _, data = self.get(fid)
        self.assertEqual(data["kind"], "table")
        self.assertEqual(len(data["rows"]), preview.MAX_ROWS)
        self.assertEqual(data["total"], 300)

    def test_xlsx_table(self):
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["Имя", "Сумма"]); ws.append(["Name1", 10])
        buf = io.BytesIO(); wb.save(buf)
        fid = self.add("t.xlsx", buf.getvalue())
        _, data = self.get(fid)
        self.assertEqual(data["kind"], "table")
        self.assertEqual(data["rows"][0], ["Имя", "Сумма"])

    def test_docx_blocks(self):
        import docx
        d = docx.Document(); d.add_heading("Договор", 1); d.add_paragraph("Сторона Name1")
        t = d.add_table(rows=1, cols=2); t.rows[0].cells[0].text = "Org1"; t.rows[0].cells[1].text = "5"
        buf = io.BytesIO(); d.save(buf)
        fid = self.add("d.docx", buf.getvalue())
        _, data = self.get(fid)
        self.assertEqual(data["kind"], "doc")
        kinds = [b["t"] for b in data["blocks"]]
        self.assertEqual(kinds, ["h", "p", "table"])

    def test_pptx_slides(self):
        import pptx
        deck = pptx.Presentation(); s = deck.slides.add_slide(deck.slide_layouts[1])
        s.shapes.title.text = "Отчёт Org1"; s.placeholders[1].text = "Name1"
        buf = io.BytesIO(); deck.save(buf)
        fid = self.add("s.pptx", buf.getvalue())
        _, data = self.get(fid)
        self.assertEqual(data["kind"], "slides")
        self.assertEqual(data["slides"][0]["n"], 1)

    def test_pdf_page_png_and_navigation(self):
        import fitz
        doc = fitz.open()
        for i in range(3):
            doc.new_page().insert_text((72, 72), f"Page {i + 1} Name1")
        fid = self.add("p.pdf", doc.tobytes())
        _, data = self.get(fid)
        self.assertEqual((data["kind"], data["pages"], data["page"]), ("pdf", 3, 1))
        self.assertTrue(data["image"].startswith("data:image/png;base64,"))
        _, last = self.get(fid, page=9)
        self.assertEqual(last["page"], 3)

    def test_image_data_uri(self):
        import fitz
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 4, 4), False)
        fid = self.add("i.png", pix.tobytes("png"))
        _, data = self.get(fid)
        self.assertEqual(data["kind"], "image")

    def test_size_limit_fails_closed(self):
        result = preview.render("big.txt", b"a" * (preview.MAX_BYTES + 1))
        self.assertEqual(result["kind"], "none")
        self.assertIn("слишком большой", result["reason"])

    def test_broken_and_unknown_fail_closed(self):
        self.assertEqual(preview.render("x.docx", b"not a zip")["kind"], "none")
        self.assertEqual(preview.render("x.exe", b"MZ")["kind"], "none")

    def test_missing_role_is_404(self):
        fid = self.add("m.txt", ANON_TEXT.encode())
        self.assertEqual(self.get(fid, "orig")[0], 404)
        self.assertEqual(self.get(fid, "bogus")[0], 404)

    def test_panel_code_never_writes_file_content_as_markup(self):
        src = (WEB / "index.html").read_text(encoding="utf-8")
        block = src[src.index("/* S11A preview panel */"):src.index("/* end S11A */")]
        self.assertNotIn("innerHTML", block.replace("innerHTML = \"\"", ""))
        self.assertNotIn("fetch(\"http", block)
        self.assertIn('role="complementary"', src)


if __name__ == "__main__":
    unittest.main()
