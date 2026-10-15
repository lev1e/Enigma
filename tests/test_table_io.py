"""Таблица замен: выгрузка, загрузка правок одним шагом и «вернуть заново»."""
import base64
import csv
import io
import json
import unittest

from openpyxl import Workbook, load_workbook

from anonymizer import chat_api, table_io
from anonymizer.webapp import App

from .helpers import Env

TEXT = "Договор подписал Пётр Сидоров. Позвонить Петру Сидорову. Анна Кузнецова согласна."


def sheet_bytes(rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def csv_bytes(rows, delimiter=";", encoding="utf-8-sig"):
    out = io.StringIO(newline="")
    csv.writer(out, delimiter=delimiter, lineterminator="\r\n").writerows(rows)
    return out.getvalue().encode(encoding)


class TableBase(unittest.TestCase):
    def setUp(self):
        self.env = Env()
        self.app = App(self.env.vault, self.env.base / "work", "tok")
        self.chat = self.anonymize("a.txt", TEXT)

    def tearDown(self):
        self.app.cleanup_open()
        self.env.close()

    def finish(self, job_id):
        self.app.runners[job_id].join(30)
        return self.app.job_public(job_id)

    def anonymize(self, name, text, chat_id=None):
        upload = self.app.add_upload(name, text.encode("utf-8"))["id"]
        done = self.finish(self.app.start_job("anonymize", [upload], {}, chat_id=chat_id))
        self.assertEqual(done["state"], "done", done)
        return done["chat"]

    def call(self, method, path, payload=None):
        return chat_api.handle(self.app, method, path, lambda: payload or {})

    def table(self, chat=None):
        return self.app.chats.table(chat or self.chat, self.app.vault)

    def rows(self):
        return table_io.export_rows(self.app, self.chat)

    def code_of(self, fragment):
        return next(r[0] for r in self.rows() if fragment in r[2])

    def preview(self, rows, fmt="csv"):
        data = csv_bytes(rows) if fmt == "csv" else sheet_bytes(rows)
        return table_io.import_table(self.app, self.chat, "т." + fmt, data)

    def apply(self, token):
        return table_io.import_table(self.app, self.chat, "", b"", apply=True, token=token)


class ExportTests(TableBase):
    def test_columns_and_one_row_per_variant(self):
        kind, data, name = self.call("GET", f"/api/chats/{self.chat}/table.csv")
        self.assertEqual((kind, name), ("file", "замены.csv"))
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"))
        self.assertIn(b"\r\n", data)
        lines = list(csv.reader(io.StringIO(data.decode("utf-8-sig")), delimiter=";"))
        self.assertEqual(lines[0], ["Код", "Вид", "Значение", "Скрывать", "Объединить с"])
        variants = sum(len(r["variants"]) for r in self.table())
        self.assertEqual(len(lines) - 1, variants)
        self.assertTrue(any("Сидоров" in line[2] for line in lines[1:]))
        self.assertTrue(all(line[3] in ("да", "нет") for line in lines[1:]))

    def test_xlsx_and_kept_flag(self):
        row = next(r for r in self.table() if r["variants"][0]["original"].startswith("Пётр"))
        self.app.vault.add_term("keep_terms", row["variants"][0]["original"])
        kind, data, name = self.call("GET", f"/api/chats/{self.chat}/table.xlsx")
        self.assertEqual(name, "замены.xlsx")
        sheet = load_workbook(io.BytesIO(data)).worksheets[0]
        values = [[c.value for c in r] for r in sheet.iter_rows()]
        self.assertEqual(values[0], ["Код", "Вид", "Значение", "Скрывать", "Объединить с"])
        kept = [v for v in values[1:] if v[3] == "нет"]
        self.assertEqual(len(kept), 1)
        self.assertEqual(kept[0][0], row["variants"][0]["token"])

    def test_merged_source_row_names_its_target(self):
        names = [r for r in self.table() if r["kind"] == "PERSON"]
        self.assertGreaterEqual(len(names), 2)
        self.app.vault.merge(names[1]["base"], names[0]["base"])
        merged = [r for r in self.rows() if r[4]]
        self.assertTrue(merged)
        self.assertEqual(merged[0][4], names[0]["base"])

    def test_table_file_writes_a_temp_file(self):
        path = self.app.table_file(self.chat, "xlsx")
        self.assertEqual(path.name, "замены.xlsx")
        self.assertTrue(path.read_bytes().startswith(b"PK"))
        self.assertEqual(self.app.table_file(self.chat, "csv").name, "замены.csv")
        with self.assertRaises(ValueError):
            self.app.table_file(self.chat, "doc")
        with self.assertRaises(KeyError):
            self.app.table_file("nochat1234", "csv")


class ParseTests(TableBase):
    def test_cyrillic_headers_any_case_and_delimiters(self):
        head = ["  КОД ", "значение"]
        for delimiter in (";", ",", "\t"):
            rows = table_io.parse_table("x.csv", csv_bytes([head, ["Name1", "Иванов"]], delimiter))
            self.assertEqual((rows[0]["code"], rows[0]["value"], rows[0]["row"]), ("Name1", "Иванов", 2))

    def test_cp1251_and_bom_and_xlsx(self):
        rows = [["Код", "Значение", "Скрывать"], ["Name1", "Иван Петров", "да"]]
        self.assertEqual(table_io.parse_table("a.csv", csv_bytes(rows, ";", "cp1251"))[0]["value"], "Иван Петров")
        self.assertEqual(table_io.parse_table("a.csv", csv_bytes(rows))[0]["hide"], "да")
        self.assertEqual(table_io.parse_table("a.xlsx", sheet_bytes(rows))[0]["value"], "Иван Петров")

    def test_garbage_is_a_plain_error(self):
        for name, data in (("a.xlsx", b"PK\x03\x04junk"), ("a.csv", b"\xff\xfe\x00\x00" * 5), ("a.csv", b""),
                           ("a.csv", csv_bytes([["Имя", "Ещё"], ["1", "2"]]))):
            with self.assertRaises(ValueError):
                table_io.parse_table(name, data)
        with self.assertRaises(ValueError):
            table_io.parse_table("big.csv", b"x" * (table_io.MAX_TABLE_BYTES + 1))


class ImportTests(TableBase):
    def head(self):
        return ["Код", "Значение", "Скрывать", "Объединить с"]

    def test_preview_changes_nothing_and_text_has_no_values(self):
        code = self.code_of("Сидоров")
        before = json.dumps(self.table(), sort_keys=True)
        plan = self.preview([self.head(), [code, "Секретов", "нет", ""]])
        self.assertEqual(json.dumps(self.table(), sort_keys=True), before)
        self.assertEqual(plan["restore_changes"], [f"{code}: значение изменится"])
        self.assertEqual(plan["ai_changes"], [f"{code}: скрывать -> не скрывать"])
        self.assertTrue(plan["text"].startswith("Изменит файл для ИИ: 1. Изменит только возврат: 1."))
        self.assertNotIn("Секретов", json.dumps(plan, ensure_ascii=False))
        self.assertNotIn("Не распознано", plan["text"])

    def test_identical_rows_are_unchanged(self):
        plan = self.preview([self.head(), *[[r[0], r[2], r[3], ""] for r in self.rows()]], fmt="xlsx")
        self.assertEqual((plan["ai_changes"], plan["restore_changes"], plan["unrecognized"]), ([], [], []))
        self.assertEqual(plan["unchanged"], len(self.rows()))

    def test_unrecognized_rows_are_reported_with_sheet_row_numbers(self):
        other = self.anonymize("b.txt", "Письмо для Ольги Громовой.")
        foreign = next(r["variants"][0]["token"] for r in self.table(other) if "Ольг" in r["variants"][0]["original"])
        good = self.code_of("Сидоров")
        plan = self.preview([self.head(), ["Name999", "Икс", "", ""], [foreign, "Чужой", "", ""], [good, "", "", ""],
                             ["не код", "Икс", "", ""], [good, "Х", "может", ""], [good, "Х", "", "Company1"]])
        self.assertEqual([u["row"] for u in plan["unrecognized"]], [2, 3, 4, 5, 6, 7])
        self.assertTrue(plan["text"].endswith("Не распознано строк: 6 (номера: 2, 3, 4, 5, 6, 7)."))
        self.assertEqual(plan["restore_changes"], [])

    def test_apply_is_one_step_and_batch_undo_restores_everything(self):
        names = [r for r in self.table() if r["kind"] == "PERSON"]
        a, b = names[0], names[1]
        code = a["variants"][0]["token"]
        before = json.dumps(self.table(), sort_keys=True)
        hide_before = list(self.app.vault.prefs["hide_terms"]), list(self.app.vault.prefs["keep_terms"])
        steps_before = len(self.app.chats.get(self.chat)["steps"])
        plan = self.preview([self.head(), [code, "Новое", "нет", b["base"]]])
        self.assertEqual(plan["merges"], [f"{a['base']}: объединить с {b['base']}"])
        done = self.apply(plan["token"])
        self.assertEqual(done["applied"], {"ai": 1, "restore": 1, "merges": 1})
        self.assertTrue(done["ai_changed"])
        self.assertEqual(done["undo"]["of"], "batch")
        self.assertEqual(len(done["undo"]["items"]), 3)
        steps = self.app.chats.get(self.chat)["steps"]
        self.assertEqual(len(steps), steps_before + 1)
        self.assertEqual((steps[-1]["kind"], steps[-1]["action"]), ("edit", "import"))
        self.assertNotIn("Новое", json.dumps(steps, ensure_ascii=False))
        self.assertNotEqual(json.dumps(self.table(), sort_keys=True), before)
        self.assertEqual(self.app.vault.describe(a["base"])["base"], b["base"])
        _, undone = self.call("POST", f"/api/chats/{self.chat}/entities/{a['base']}", {"action": "undo", **done["undo"]})
        self.assertIn("table", undone)
        # Слияние снимается, а написания, перенесённые в целевую метку, остаются (так устроена отмена слияния).
        again = {r["base"]: r for r in self.table()}
        self.assertIn(a["base"], again)
        self.assertEqual(again[a["base"]]["variants"], a["variants"])
        self.assertEqual(again[b["base"]]["variants"][:len(b["variants"])], b["variants"])
        self.assertEqual((self.app.vault.prefs["hide_terms"], self.app.vault.prefs["keep_terms"]), hide_before)

    def test_table_undo_route_and_import_route_with_base64(self):
        code = self.code_of("Сидоров")
        data = base64.b64encode(csv_bytes([self.head(), [code, "Другой"]])).decode()
        _, plan = self.call("POST", f"/api/chats/{self.chat}/table/import", {"name": "т.csv", "data": data, "apply": False})
        _, done = self.call("POST", f"/api/chats/{self.chat}/table/import", {"apply": True, "token": plan["token"]})
        self.assertEqual(done["applied"]["restore"], 1)
        self.assertEqual(self.app.vault.describe(code)["variants"][0]["value"], "Другой")
        self.call("POST", f"/api/chats/{self.chat}/table/undo", {"undo": done["undo"]})
        self.assertIn("Сидоров", self.app.vault.describe(code)["variants"][0]["value"])

    def test_token_is_single_use_expires_and_is_per_chat(self):
        code = self.code_of("Сидоров")
        plan = self.preview([self.head(), [code, "Другой"]])
        other = self.anonymize("c.txt", "Письмо для Ольги Громовой.")
        with self.assertRaises(ValueError) as caught:
            table_io.import_table(self.app, other, "", b"", apply=True, token=plan["token"])
        self.assertEqual(str(caught.exception), "Просмотр устарел. Загрузите таблицу заново.")
        plan = self.preview([self.head(), [code, "Другой"]])
        self.app.table_plans[plan["token"]]["time"] -= table_io.PLAN_TTL + 1
        with self.assertRaises(ValueError):
            self.apply(plan["token"])
        plan = self.preview([self.head(), [code, "Другой"]])
        self.apply(plan["token"])
        with self.assertRaises(ValueError):
            self.apply(plan["token"])
        for _ in range(12):
            self.preview([self.head(), [code, "Другой"]])
        self.assertLessEqual(len(self.app.table_plans), table_io.MAX_PLANS)

    def test_merge_of_different_kind_or_conflicting_rows_is_not_applied(self):
        persons = [r for r in self.table() if r["kind"] == "PERSON"]
        other_kind = self.anonymize("d.txt", "Звонил Олег Громов из ООО «Ромашка».", self.chat)
        orgs = [r for r in self.table() if r["kind"] == "ORG"]
        self.assertTrue(orgs)
        plan = self.preview([self.head(), [persons[0]["variants"][0]["token"], persons[0]["variants"][0]["value"], "", orgs[0]["base"]]])
        self.assertEqual(plan["merges"], [])
        self.assertEqual(len(plan["unrecognized"]), 1)


class RestoreAgainTests(TableBase):
    def restore(self):
        data = self.app.chats.read_blob(self.chat, self.app.chats.get(self.chat)["files"][0]["id"], "anon")
        upload = self.app.add_upload("ответ.txt", data)["id"]
        done = self.finish(self.app.start_job("restore", [upload], {}, chat_id=self.chat))
        self.assertEqual(done["state"], "done", done)
        return done

    def test_fresh_restore_is_not_stale_and_edit_makes_it_stale(self):
        self.restore()
        self.assertEqual(self.call("GET", f"/api/chats/{self.chat}")[1]["restore"], {"stale": 0, "redo": False, "reason": ""})
        code = self.code_of("Сидоров")
        self.app.vault.edit_value(code, 1, "Другой Человек")
        state = self.call("GET", f"/api/chats/{self.chat}")[1]["restore"]
        self.assertEqual((state["stale"], state["redo"], state["reason"]), (1, True, ""))
        record = self.app.chats.get(self.chat)["files"][0]
        self.assertIn("ret", record["roles"])
        self.assertNotIn("Сидоров", json.dumps(record, ensure_ascii=False))

    def test_rerestore_records_a_new_step_and_uses_new_values(self):
        self.restore()
        steps = len(self.app.chats.get(self.chat)["steps"])
        self.app.vault.edit_value(self.code_of("Сидоров"), 1, "Другой Человек")
        _, started = self.call("POST", f"/api/chats/{self.chat}/rerestore")
        done = self.finish(started["job"])
        self.assertEqual(done["state"], "done", done)
        chat = self.app.chats.get(self.chat)
        self.assertEqual(len(chat["steps"]), steps + 1)
        self.assertEqual(chat["steps"][-1]["kind"], "restore")
        self.assertEqual(len(chat["files"]), 1)
        self.assertIn("Другой Человек".encode(), self.app.chats.read_blob(self.chat, chat["files"][0]["id"], "restored"))
        self.assertEqual(self.call("GET", f"/api/chats/{self.chat}")[1]["restore"]["stale"], 0)

    def test_rerestore_without_stale_or_without_source(self):
        with self.assertRaises(ValueError):
            self.call("POST", f"/api/chats/{self.chat}/rerestore")
        self.restore()
        record = self.app.chats.get(self.chat)["files"][0]
        chat = self.app.chats.get(self.chat)
        self.app.chats._shred(self.app.chats._blob_path(self.chat, record["id"], "ret"))
        data = self.app.chats._read(self.chat)
        del data["files"][0]["roles"]["ret"]
        self.app.chats._write(data)
        self.app.vault.edit_value(self.code_of("Сидоров"), 1, "Другой Человек")
        state = self.call("GET", f"/api/chats/{self.chat}")[1]["restore"]
        self.assertEqual((state["stale"], state["redo"], state["reason"]), (1, False, table_io.NO_RETURNED))
        with self.assertRaises(ValueError) as caught:
            self.call("POST", f"/api/chats/{self.chat}/rerestore")
        self.assertEqual(str(caught.exception), table_io.NO_RETURNED)


if __name__ == "__main__":
    unittest.main()
