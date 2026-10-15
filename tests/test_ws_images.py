"""Картинки: логотип с белым стилизованным названием на тёмном фоне не должен проходить молча."""
from __future__ import annotations

import io
import random
import unittest
import zipfile
from pathlib import Path

import fitz
from PIL import Image, ImageDraw, ImageFilter, ImageFont

from anonymizer import imagetext, ocr
from tests.helpers import Env

from tests.font_support import system_font

ARIAL = system_font()
TIMES = system_font(serif=True)
READY = Path(TIMES).exists() and ocr.available()
BODY_BARE = "Компания Кооперация Энергия выпускает программное обеспечение."
BODY_LEGAL = "Заказчик: ООО «Кооперация Энергия», г. Москва."


def logo_png(size=44, fg=(255, 255, 255), lines=("Кооперация", "Энергия"), fmt="PNG", mark=True) -> bytes:
    w, h = 900, 300
    rnd = random.Random(3)
    img = Image.new("RGB", (w, h))
    px = img.load()
    for y in range(h):
        for x in range(w):
            n = rnd.randrange(-18, 18)
            px[x, y] = tuple(max(0, min(255, c + n)) for c in (int(10 + 30 * x / w), int(25 + 40 * x / w), int(80 + 60 * (1 - x / w))))
    img = img.filter(ImageFilter.GaussianBlur(1.2))
    d = ImageDraw.Draw(img)
    x0 = 40
    if mark:
        d.rectangle((40, 60, 200, 220), fill=(200, 30, 40))
        x0 = 240
    font = ImageFont.truetype(TIMES, size)
    for i, line in enumerate(lines):
        d.text((x0, 80 + i * size * 1.3), line, font=font, fill=fg)
    buf = io.BytesIO()
    img.save(buf, fmt, **({"quality": 85} if fmt == "JPEG" else {}))
    return buf.getvalue()


def pptx_with(png: bytes, body: str) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Презентация"
    box = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(6), Inches(1))
    box.text_frame.text = body
    slide.shapes.add_picture(io.BytesIO(png), Inches(1), Inches(1.5), width=Inches(6))
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def media(data: bytes) -> list[bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return [z.read(n) for n in z.namelist() if "/media/" in n]


def seen_words(data: bytes) -> str:
    doc = fitz.open()
    pix = fitz.Pixmap(data)
    page = doc.new_page(width=pix.width, height=pix.height)
    page.insert_image(page.rect, stream=data)
    return " ".join(w.text for w in ocr.ocr_page_words(page, scale=2.0)[0])


def text_of(f) -> str:
    return " ".join(m["text"] for m in f.messages)


@unittest.skipUnless(READY, "нужны Times New Roman и движок распознавания")
class LogoText(unittest.TestCase):
    def setUp(self):
        self.env = Env(image_mode="redact_text")
        self.addCleanup(self.env.close)

    def test_unknown_brand_in_standalone_png_is_flagged_loudly(self):
        job = self.env.anonymize(("logo.png", logo_png()))
        f = job.files[0]
        self.assertIn("Кооперация Энергия", text_of(f))
        self.assertIn("logo.png", text_of(f))
        self.assertNotEqual(f.status, "ok")

    def test_unknown_brand_in_pptx_is_flagged_loudly(self):
        job = self.env.anonymize(("deck.pptx", pptx_with(logo_png(), BODY_BARE)))
        f = job.files[0]
        self.assertIn("Кооперация Энергия", text_of(f))
        self.assertNotEqual(f.status, "ok")

    def test_known_organisation_split_over_two_lines_is_painted_and_restored(self):
        png = logo_png()
        job = self.env.anonymize(("deck.pptx", pptx_with(png, BODY_LEGAL)))
        f = job.files[0]
        (new,) = media(Path(f.out_path).read_bytes())
        seen = seen_words(new)
        self.assertNotIn("Кооперация", seen)
        self.assertNotIn("Энергия", seen)
        rf = self.env.restore(self.env.as_upload(job)).files[0]
        self.assertEqual(media(Path(rf.out_path).read_bytes()), [png])

    def test_known_organisation_misread_by_engine_is_still_painted(self):
        # Движок читает «Энергня» (подмена вместо реального OCR, чтобы тест не зависел от случайности распознавания).
        # Известное название из хранилища сопоставляется с допуском на ошибки.
        from unittest import mock
        env = Env(image_mode="redact_text", hide_terms=["Кооперация Энергия"])
        self.addCleanup(env.close)
        png = logo_png(size=16)
        real = ocr.provider()
        with mock.patch.object(ocr, "provider", lambda: MisreadingEngine(real)):
            job = env.anonymize(("small.png", png))
        f = job.files[0]
        self.assertTrue(f.out_path, text_of(f))
        seen = seen_words(Path(f.out_path).read_bytes())
        self.assertNotIn("Коопер", seen)
        self.assertNotIn("Энерг", seen)

    def test_logo_in_pdf_embedded_image_is_flagged(self):
        doc = fitz.open()
        page = doc.new_page()
        page.insert_font(fontname="ar", fontfile=ARIAL)
        page.insert_text((72, 72), "Отчёт за квартал", fontname="ar", fontsize=12)
        page.insert_image(fitz.Rect(72, 100, 472, 233), stream=logo_png())
        f = self.env.anonymize(("r.pdf", doc.tobytes())).files[0]
        self.assertIn("Кооперация Энергия", text_of(f))

    def test_ordinary_picture_text_is_not_flagged(self):
        png = logo_png(lines=("квартальный отчёт по продажам",), mark=False)
        f = self.env.anonymize(("plain.png", png)).files[0]
        self.assertNotIn("логотип", text_of(f).lower())


@unittest.skipUnless(READY, "нужен движок распознавания")
class Layers(unittest.TestCase):
    def test_multi_pass_returns_several_layers(self):
        doc = fitz.open()
        data = logo_png()
        pix = fitz.Pixmap(data)
        page = doc.new_page(width=pix.width, height=pix.height)
        page.insert_image(page.rect, stream=data)
        layers = ocr.ocr_image_layers(page, scale=2.0)
        self.assertGreaterEqual(len(layers), 4)
        joined = " ".join(w.text for layer in layers for w in layer)
        self.assertIn("Энергия", joined)


class Helpers(unittest.TestCase):
    def words(self, rows):
        return [ocr.Word(t, x0, y0, x1, y1, line) for t, x0, y0, x1, y1, line in rows]

    def test_brand_like_two_title_words_on_two_lines(self):
        ws = self.words([("Кооперация", 0, 0, 200, 40, 0), ("Энергия", 0, 50, 150, 90, 1)])
        self.assertEqual(imagetext.brand_phrases(ws, []), ["Кооперация Энергия"])

    def test_plain_sentence_is_not_brand_like(self):
        ws = self.words([(t, i * 60, 0, i * 60 + 50, 12, 0) for i, t in enumerate("отчёт по продажам за год вырос".split())])
        self.assertEqual(imagetext.brand_phrases(ws, []), [])

    def test_covered_phrase_is_not_reported(self):
        ws = self.words([("Кооперация", 0, 0, 200, 40, 0), ("Энергия", 0, 50, 150, 90, 1)])
        self.assertEqual(imagetext.brand_phrases(ws, [fitz.Rect(0, 0, 200, 100)]), [])

    def test_fuzzy_match_tolerates_ocr_errors(self):
        self.assertTrue(imagetext.close_enough("доверня", "доверия"))
        self.assertFalse(imagetext.close_enough("продажам", "доверия"))
        spans = imagetext.fuzzy_spans("Кооперацни Энергня и ещё слова", ["Кооперация Энергия"])
        self.assertEqual(spans, [(0, 18, "Кооперация Энергия")])


if __name__ == "__main__":
    unittest.main()


@unittest.skipUnless(READY, "нужны Times New Roman и движок распознавания")
class PdfEmbeddedLogo(unittest.TestCase):
    def test_known_org_in_pdf_picture_is_not_flagged_as_unknown(self):
        env = Env(image_mode="redact_text")
        self.addCleanup(env.close)
        doc = fitz.open()
        page = doc.new_page()
        page.insert_font(fontname="ar", fontfile=ARIAL)
        page.insert_text((72, 72), "Заказчик: ООО «Кооперация Энергия», г. Москва.", fontname="ar", fontsize=12)
        page.insert_image(fitz.Rect(72, 100, 472, 233), stream=logo_png())
        f = env.anonymize(("r.pdf", doc.tobytes())).files[0]
        self.assertNotIn("не опознала", text_of(f))


class MisreadingEngine:
    """Настоящий движок, но слово «Энергия» он читает как «Энергня» во всех проходах."""
    native_rotation = False

    def __init__(self, real):
        self.real = real

    def words(self, png):
        out = self.real.words(png)
        for w in out:
            w.text = w.text.replace("Энергия", "Энергня")
        return out


@unittest.skipUnless(READY, "нужны Times New Roman и движок распознавания")
class BrandPolicy(unittest.TestCase):
    def make(self, policy=None, **prefs):
        if policy:
            prefs["image_brand_policy"] = policy
        env = Env(image_mode="redact_text", **prefs)
        self.addCleanup(env.close)
        return env

    def test_default_is_paint_with_notice_review_and_exact_restore(self):
        env = self.make()
        png = logo_png()
        job = env.anonymize(("logo.png", png))
        f = job.files[0]
        out = Path(f.out_path).read_bytes()
        seen = seen_words(out)
        self.assertNotIn("Кооперация", seen)
        self.assertNotIn("Энергия", seen)
        text = text_of(f)
        self.assertIn("Кооперация Энергия", text)            # предупреждение сохраняется
        self.assertIn("Закрашено названий на логотипах: 1", text)
        rf = env.restore(env.as_upload(job)).files[0]
        self.assertEqual(Path(rf.out_path).read_bytes(), png)

    def test_painted_logo_text_gets_same_token_elsewhere(self):
        env = self.make()
        job = env.anonymize(("logo.png", logo_png()), ("note.txt", "Партнёр Кооперация Энергия подписал договор."))
        token = next(m for m in seen_words(Path(job.files[0].out_path).read_bytes()).split() if m.startswith("Company"))
        note = Path(job.files[1].out_path).read_text(encoding="utf-8")
        self.assertNotIn("Кооперация Энергия", note)
        self.assertIn(token.split("_")[0].rstrip("0123456789") , note)

    def test_pptx_paint_restores_exactly(self):
        env = self.make("paint")
        png = logo_png()
        job = env.anonymize(("deck.pptx", pptx_with(png, BODY_BARE)))
        f = job.files[0]
        (new,) = media(Path(f.out_path).read_bytes())
        self.assertNotIn("Энергия", seen_words(new))
        self.assertIn("Закрашено названий на логотипах", text_of(f))
        rf = env.restore(env.as_upload(job)).files[0]
        self.assertEqual(media(Path(rf.out_path).read_bytes()), [png])

    def test_warn_mode_leaves_pixels_and_warns(self):
        env = self.make("warn")
        f = env.anonymize(("logo.png", logo_png())).files[0]
        self.assertIn("Энергия", seen_words(Path(f.out_path).read_bytes()))
        self.assertIn("не опознала", text_of(f))
        self.assertNotIn("Закрашено названий", text_of(f))

    def test_off_mode_is_silent_about_logos(self):
        env = self.make("off")
        f = env.anonymize(("logo.png", logo_png())).files[0]
        self.assertNotIn("Кооперация Энергия", text_of(f))
        self.assertNotIn("Закрашено названий", text_of(f))

    def test_pdf_embedded_logo_is_painted(self):
        env = self.make()
        doc = fitz.open()
        page = doc.new_page()
        page.insert_font(fontname="ar", fontfile=ARIAL)
        page.insert_text((72, 72), "Отчёт за квартал", fontname="ar", fontsize=12)
        page.insert_image(fitz.Rect(72, 100, 472, 233), stream=logo_png())
        f = env.anonymize(("r.pdf", doc.tobytes())).files[0]
        out = fitz.open(stream=Path(f.out_path).read_bytes(), filetype="pdf")
        words = [w.text for w in ocr.ocr_page_words(out[0])[0]]
        self.assertNotIn("Энергия", " ".join(words))
        self.assertIn("Закрашено названий на логотипах", text_of(f))

    def test_single_word_and_table_text_is_not_painted(self):
        self.assertFalse(imagetext.paintable("Сбербанк", 1))
        self.assertFalse(imagetext.paintable("Кооперация Энергия", 6))
        self.assertTrue(imagetext.paintable("Кооперация Энергия", 1))

    def test_schema_option_has_label_and_validates(self):
        from anonymizer import settings_schema as S
        opt = S.BY_KEY["image_brand_policy"]
        self.assertEqual(opt.default, "paint")
        self.assertTrue(opt.label)
        self.assertEqual(S.coerce(opt, "warn"), "warn") if hasattr(S, "coerce") else None

    def test_fuzzy_only_known_name_is_painted_when_engine_misreads(self):
        from unittest import mock
        env = self.make("off", hide_terms=["Кооперация Энергия"])
        real = ocr.provider()
        png = logo_png()
        with mock.patch.object(ocr, "provider", lambda: MisreadingEngine(real)):
            job = env.anonymize(("logo.png", png))
        f = job.files[0]
        self.assertTrue(f.out_path, text_of(f))
        seen = seen_words(Path(f.out_path).read_bytes())
        self.assertNotIn("Кооперация", seen)
        rf = env.restore(env.as_upload(job)).files[0]
        self.assertEqual(Path(rf.out_path).read_bytes(), png)
