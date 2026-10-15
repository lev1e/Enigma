"""Умная таблица Excel: name и displayName меняются одинаково, проверка add-in не видит ложной утечки."""
import io
import re
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from anonymizer.addin import verify_document

from .helpers import Env

SEC = "Смирнов Алексей Петрович"


def table_xml(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return next(z.read(n).decode() for n in z.namelist() if n.startswith("xl/tables/"))


class XlsxTableNames(unittest.TestCase):
    def setUp(self):
        self.env = Env()
        self.addCleanup(self.env.close)

    def build(self, name, display):
        wb = Workbook()
        ws = wb.active
        ws.cell(1, 1, "Имя")
        ws.cell(1, 2, "Кто")
        ws.cell(2, 1, f"Директор {SEC}")
        ws.cell(2, 2, "x")
        ws.add_table(Table(name=name, displayName=display, ref="A1:B2"))
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    def attrs(self, data):
        xml = table_xml(data)
        return re.search(r' name="([^"]*)"', xml).group(1), re.search(r'displayName="([^"]*)"', xml).group(1)

    def test_plain_table_name_and_display_name_stay_consistent_and_verify_clean(self):
        job = self.env.anonymize(("t.xlsx", self.build("Tbl1", "Tbl1")))
        out = self.env.output_bytes(job)
        name, display = self.attrs(out)
        self.assertEqual(name, display)
        self.assertNotIn(SEC.split()[0], table_xml(out))
        with tempfile.TemporaryDirectory() as tmp:
            actual = Path(tmp) / "a.xlsx"
            actual.write_bytes(out)
            verdict = verify_document(self.env.output(job), actual, job.replaced)
        self.assertTrue(verdict["ok"], verdict)

    def test_table_named_after_a_person_is_renamed_in_both_attributes(self):
        job = self.env.anonymize(("t.xlsx", self.build("Смирнов_Алексей", "Смирнов_Алексей")))
        name, display = self.attrs(self.env.output_bytes(job))
        self.assertEqual(name, display)
        self.assertNotIn("Смирнов", name)
        restored = self.env.restore(self.env.as_upload(job))
        self.assertEqual(self.attrs(self.env.output_bytes(restored)), ("Смирнов_Алексей", "Смирнов_Алексей"))


if __name__ == "__main__":
    unittest.main()
