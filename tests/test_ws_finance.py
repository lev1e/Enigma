"""Finance masking (P0): default off, exact precision, verify, non-cell surfaces on Excel-style workbooks, notices, restore."""
from __future__ import annotations

import io
import re
import unittest
import zipfile
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from anonymizer import finance as F
from anonymizer.models import Settings
from anonymizer.service import Upload
from tests.helpers import Env
from tools import make_finance_realistic as R

ON = {"finance_numbers": True}


def book(rows, extra=None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for r in rows:
        ws.append(r)
    for k, v in (extra or {}).items():
        ws[k] = v
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def all_parts_text(data: bytes, skip=None) -> str:
    z = zipfile.ZipFile(io.BytesIO(data))
    return "\n".join(f"<<{n}>>\n" + z.read(n).decode("utf-8", "ignore") for n in z.namelist() if not (skip and skip.search(n)))


def forms(value) -> set[str]:
    """Как настоящая сумма может выглядеть в файле."""
    d = Decimal(str(value))
    plain = format(d.normalize(), "f")
    out = {plain, f"{int(d):,}".replace(",", " ") if d == int(d) else plain, f"{int(d):,}" if d == int(d) else plain,
           f"{int(d):,}".replace(",", " ") if d == int(d) else plain}
    return {o for o in out if len(o) >= 5}


def has_number(text: str, token: str) -> bool:
    return re.search(r"(?<![\d.,])" + re.escape(token) + r"(?![\d]|[.,]\d)", text) is not None


class Defaults(unittest.TestCase):
    def test_off_by_default(self):
        self.assertFalse(Settings().finance_numbers)
        env = Env()
        self.addCleanup(env.close)
        job = env.anonymize(("a.xlsx", book([["Выручка", 4_200_000], ["Расходы", 3_100_000], ["Прибыль", 1_100_000], ["Налог", 220_000]])))
        out = all_parts_text(env.output_bytes(job))
        self.assertIn("4200000", out)


class Precision(unittest.TestCase):
    def test_exact_sums_and_roundtrip(self):
        for _ in range(300):
            k = F.make_factor()
            vals = [Decimal(str(x)) for x in (1250.5, 0.0125, 4_200_000, 17, 999_999.99, 0.645)]
            scaled = [F.scale_value(v, k) for v in vals]
            # сумма масштабированных сходится с масштабированной суммой с точностью Excel (15 значащих цифр)
            self.assertEqual(F._cap_sig(sum(scaled)), F.scale_value(sum(vals), k), k)
            self.assertLessEqual(abs(sum(scaled) - F.scale_value(sum(vals), k)), Decimal("0.0000001"))
            for v, s in zip(vals, scaled):
                self.assertNotEqual(v, s)
                self.assertEqual(F.unscale_value(s, k), v.quantize(Decimal("0.01")) if abs(v) >= 100 else v.quantize(Decimal("0.000001")))

    def test_balance_check_stays_zero(self):
        data = R.build_all()["balance_sheet_ru"]["bytes"]
        env = Env()
        self.addCleanup(env.close)
        job = env.anonymize(("b.xlsx", data), options=ON)
        self.assertTrue(job.files[0].out_path, job.files[0].messages)
        wb = load_workbook(env.output(job))
        # суммарно по константам листа: активы - пассивы тождественны исходным, умноженным на k
        k = None
        orig = load_workbook(io.BytesIO(data))
        for ws in orig.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    n = wb[ws.title][c.coordinate].value
                    if isinstance(c.value, (int, float)) and c.value > 100000 and isinstance(n, (int, float)) and n != c.value:
                        ratio = Decimal(str(n)) / Decimal(str(c.value))
                        k = k or ratio
                        self.assertAlmostEqual(float(ratio), float(k), places=9)
        self.assertIsNotNone(k)


    def test_fifteen_digit_cap(self):
        k = (23135, 10000)
        big = Decimal("987654321098765")            # 15 цифр: после умножения больше 15 значащих -> усечение до 15
        scaled = F.scale_value(big, k)
        self.assertLessEqual(F._significant(scaled), 15)
        self.assertEqual(scaled, F._cap_sig(big * 23135 / 10000))
        back = F.unscale_value(scaled, k)
        self.assertLess(abs(back - big) / big, Decimal("1e-14"))
        self.assertEqual(F.scale_value(Decimal("0.0125"), k), Decimal("0.02891875"))
        self.assertNotEqual(F.scale_value(Decimal("1"), k), Decimal("1"))


# Книги, где намеренно остаются видимыми суммы, на которые опираются конфликтные формулы (поиск по значению, произведение сумм):
# обезличивание обязано назвать эти ячейки в предупреждениях, остальное - ни одной настоящей суммы.
CONFLICT_BOOKS = {"kpi_dashboard", "consolidated_multisheet"}
SKIP_PARTS = re.compile(r"(?i)^xl/(theme|styles)|docprops/|printersettings")


class Surfaces(unittest.TestCase):
    """Каждая синтетическая Excel-подобная книга: ни одна настоящая сумма не остаётся ни в одной части архива."""

    def test_no_secret_in_any_part_and_roundtrip(self):
        built = R.build_all()
        self.assertGreaterEqual(len(built), 10)
        report = []
        for name, item in built.items():
            with self.subTest(name):
                env = Env()
                self.addCleanup(env.close)
                job = env.anonymize((name + ".xlsx", item["bytes"]), options={**ON, "finance_text_amounts": True})
                f = job.files[0]
                self.assertTrue(f.out_path, (name, f.messages))
                text = all_parts_text(env.output_bytes(job), skip=SKIP_PARTS)
                leaked = [s for s in item["secrets"] if any(has_number(text, t) for t in forms(s))]
                report.append((name, f.status, leaked))
                if name in CONFLICT_BOOKS:
                    warns = " ".join(m["text"] for m in f.messages)
                    self.assertEqual(f.status, "attention", name)
                    self.assertRegex(warns, r"Лист .*, ячейка [A-Z]+\d+:", name)
                else:
                    self.assertFalse(leaked, (name, leaked))
                self.assertNotIn("не обезличены", " ".join(m["text"] for m in f.messages))
                back = env.restore(env.as_upload(job))
                self.assertTrue(back.files[0].out_path, back.files[0].messages)
                restored = load_workbook(env.output(back))
                for ref, value in item["money_cells"].items():
                    sheet, cell = ref.rsplit("!", 1)
                    got = restored[sheet.strip("'")][cell].value
                    self.assertEqual(Decimal(str(got)), Decimal(str(value)), (name, ref))


class VerifyOutput(unittest.TestCase):
    def setUp(self):
        self.env = Env()
        self.addCleanup(self.env.close)

    def test_equal_unscaled_number_does_not_block(self):
        rows = [["Выручка", 4_200_000], ["Себестоимость", 2_000_000], ["Расходы на аренду", 500_000], ["Чистая прибыль", 1_700_000],
                ["Акций, шт", 4_200_000], ["Код проекта", 2_000_000]]
        job = self.env.anonymize(("a.xlsx", book(rows)), options=ON)
        self.assertTrue(job.files[0].out_path, job.files[0].messages)
        self.assertNotEqual(job.files[0].status, "error")

    def test_injected_leak_is_detected(self):
        wb = Workbook()
        ws = wb.active
        for r in (["Выручка", 4_200_000], ["Себестоимость", 2_000_000], ["Расходы", 500_000], ["Прибыль", 1_700_000]):
            ws.append(r)
        out = io.BytesIO()
        wb.save(out)
        job = self.env.anonymize(("a.xlsx", out.getvalue()), options=ON)
        masked = self.env.output(job)
        rec = next(iter(self.env.vault.finance.values()))
        self.assertEqual(F.verify_output(masked, rec), [])
        bad = io.BytesIO()
        with zipfile.ZipFile(masked) as zin, zipfile.ZipFile(bad, "w") as zout:
            for n in zin.namelist():
                raw = zin.read(n)
                if n == "xl/worksheets/sheet1.xml":
                    raw = raw.replace(b"</worksheet>", b'<conditionalFormatting sqref="B1"><cfRule type="cellIs" priority="1" operator="greaterThan">'
                                                       b"<formula>4200000</formula></cfRule></conditionalFormatting></worksheet>")
                zout.writestr(n, raw)
        p = self.env.base / "bad.xlsx"
        p.write_bytes(bad.getvalue())
        leaks = F.verify_output(p, rec)
        self.assertTrue(any("sheet1.xml" in x for x in leaks), leaks)
        cell = io.BytesIO()
        with zipfile.ZipFile(masked) as zin, zipfile.ZipFile(cell, "w") as zout:
            for n in zin.namelist():
                raw = zin.read(n)
                if n == "xl/worksheets/sheet1.xml":
                    raw = re.sub(rb"(<c r=\"B1\"[^>]*><v>)[^<]*", rb"\g<1>4200000", raw)
                zout.writestr(n, raw)
        p.write_bytes(cell.getvalue())
        self.assertTrue(any("B1" in x for x in F.verify_output(p, rec)))


class Text(unittest.TestCase):
    def test_million_text_amount(self):
        pairs: list = []
        out = F._scale_text("Бюджет 4,2 млн ₽ плановая выручка", (23135, 10000), pairs)
        self.assertNotIn("4,2 млн", out)
        self.assertEqual(pairs[0][1], "4,2 млн ₽")
        self.assertNotEqual(F._scale_text("Порог 15 000 000 руб", (23135, 10000), []), "Порог 15 000 000 руб")
        self.assertEqual(F._scale_text("В 2024 году 3 филиала", (23135, 10000), []), "В 2024 году 3 филиала")

    def test_text_amount_in_cell_and_notices(self):
        rows = [["Выручка", 4_200_000], ["Себестоимость", 2_000_000], ["Расходы", 500_000], ["Прибыль", 1_700_000],
                ["Комментарий: план 4,2 млн ₽", None], ["Позиция", 1234567]]
        env = Env()
        self.addCleanup(env.close)
        off = env.anonymize(("a.xlsx", book(rows)), options=ON)
        texts = " ".join(m["text"] for m in off.files[0].messages)
        self.assertNotIn("не обезличены", texts)
        self.assertIn("S!B6", texts)                           # названа ячейка с неопознанным числом
        self.assertIn("S!A5", texts)                           # и ячейка с суммой в тексте
        self.assertIn("в тексте ячеек", texts)
        on = env.anonymize(("a.xlsx", book(rows)), options={**ON, "finance_text_amounts": True})
        self.assertNotIn("4,2 млн", all_parts_text(env.output_bytes(on)))


class PerSheet(unittest.TestCase):
    def test_per_sheet_factors_differ_without_cross_refs(self):
        wb = Workbook()
        a = wb.active
        a.title = "A"
        b = wb.create_sheet("B")
        for ws in (a, b):
            for i, label in enumerate(["Выручка", "Себестоимость", "Расходы", "Прибыль"], 1):
                ws.cell(i, 1, label)
                ws.cell(i, 2, 1_000_000 * i)
        out = io.BytesIO()
        wb.save(out)
        env = Env()
        self.addCleanup(env.close)
        job = env.anonymize(("p.xlsx", out.getvalue()), options={**ON, "finance_scale_per_sheet": True})
        rec = next(iter(env.vault.finance.values()))
        self.assertNotEqual(rec["sheets"][0]["k"], rec["sheets"][1]["k"])
        back = env.restore(env.as_upload(job))
        restored = load_workbook(env.output(back))
        self.assertEqual(restored["B"]["B3"].value, 3_000_000)
        self.assertEqual(restored["A"]["B2"].value, 2_000_000)


class NumbersAsText(unittest.TestCase):
    def test_numbers_stored_as_text_are_scaled(self):
        rows = [["Выручка", "4200000"], ["Себестоимость", "2000000"], ["Расходы", "500000"], ["Прибыль", 1_700_000]]
        env = Env()
        self.addCleanup(env.close)
        job = env.anonymize(("t.xlsx", book(rows)), options=ON)
        text = all_parts_text(env.output_bytes(job))
        self.assertFalse(has_number(text, "4200000"))
        back = env.restore(env.as_upload(job))
        self.assertEqual(load_workbook(env.output(back))["S"]["B1"].value, "4200000")


if __name__ == "__main__":
    unittest.main()
