"""«Письмо для Валерия Громова. Валерий Громов подтвердил.»: падежные формы одного человека — одна метка.

Раньше родительный «Сергея Петрова» читался как женская «Петрова Сергей» (имя на «-я» считалось женским), и
именительный «Сергей Петров» становился вторым человеком. «Валерия Громова» словарь знает как женщину, и
мужской «Валерий Громов» ни с чем не склеивался; одинокое «Громов» после такой пары оставалось в тексте."""
import io
import re
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook

from .helpers import Env


def txt(text: str) -> str:
    env = Env()
    try:
        job = env.anonymize(("t.txt", text))
        return Path(job.files[0].out_path).read_text("utf-8")
    finally:
        env.close()


def person_ids(out: str) -> set[str]:
    return {m.split("_")[0] for m in re.findall(r"Name\d+(?:_\d+)?", out)}


class CaseFormsShareOneLabelTests(unittest.TestCase):
    PEOPLE = [("Валерий", "Валерия", "Громов"), ("Сергей", "Сергея", "Петров"), ("Андрей", "Андрея", "Попов"),
              ("Юрий", "Юрия", "Волков"), ("Евгений", "Евгения", "Орлов"), ("Игорь", "Игоря", "Белов"),
              ("Дмитрий", "Дмитрия", "Кузнецов"), ("Илья", "Ильи", "Козлов"), ("Никита", "Никиты", "Громов"),
              ("Алексей", "Алексея", "Белов")]

    def test_nominative_after_genitive_after_a_full_stop(self):
        for nom, gen, surname in self.PEOPLE:
            with self.subTest(name=nom):
                out = txt(f"Письмо для {gen} {surname}а. {nom} {surname} подтвердил.")
                self.assertNotIn(surname, out)
                self.assertNotIn(nom[:4], out)
                self.assertEqual(len(person_ids(out)), 1, out)

    def test_every_case_initials_caps_and_lone_surname(self):
        text = ("Письмо для Валерия Громова. Валерий Громов подтвердил. Валерию Громову ответили. С Валерием Громовым "
                "говорили. О Валерии Громове писали. В. Громов тоже. Громов В.В. ушёл. ВАЛЕРИЙ ГРОМОВ тут. Потом Громов ушёл.")
        out = txt(text)
        self.assertNotRegex(out, "(?i)громов|валер")
        self.assertEqual(len(person_ids(out)), 1, out)

    def test_ambiguous_female_reading_does_not_split_the_man(self):
        for text in ("Валерия Громова подтвердила. Валерий Громов тоже.", "Валерий Громов тоже. Для Валерии Громова.",
                     "Письмо для Евгения Орлова. Евгений Орлов подтвердил. Орлов ушёл."):
            with self.subTest(text=text):
                out = txt(text)
                self.assertNotRegex(out, "(?i)громов|валер|орлов|евген")
                self.assertEqual(len(person_ids(out)), 1, out)

    def test_ambiguous_pair_alone_is_still_hidden_and_a_later_surname_follows(self):
        out = txt("Письмо для Валерия Громова. Громов подтвердил.")
        self.assertNotIn("Громов", out)
        self.assertEqual(len(person_ids(out)), 1, out)

    def test_woman_forms_stay_one_woman(self):
        out = txt("Валерию Громову передали. Валерия Громова подтвердила. С Валерией Громовой говорили. О Валерии Громовой писали.")
        self.assertNotRegex(out, "(?i)громов|валер")
        self.assertEqual(len(person_ids(out)), 1, out)

    def test_unisex_diminutive(self):
        out = txt("Письмо для Саши Громова. Саша Громов подтвердил.")
        self.assertNotRegex(out, "(?i)громов|саш")
        self.assertEqual(len(person_ids(out)), 1, out)

    def test_docx_paragraph_through_the_service(self):
        env = Env()
        d = Document()
        d.add_paragraph("Письмо для Валерия Громова. Валерий Громов подтвердил.")
        buf = io.BytesIO()
        d.save(buf)
        job = env.anonymize(("a.docx", buf.getvalue()))
        text = "\n".join(p.text for p in Document(str(job.files[0].out_path)).paragraphs)
        self.assertNotRegex(text, "(?i)громов|валер")
        self.assertEqual(len(person_ids(text)), 1, text)
        env.close()

    def test_xlsx_cell_through_the_service(self):
        env = Env()
        wb = Workbook()
        wb.active["A1"] = "Письмо для Сергея Петрова. Сергей Петров подтвердил."
        buf = io.BytesIO()
        wb.save(buf)
        job = env.anonymize(("a.xlsx", buf.getvalue()))
        value = load_workbook(str(job.files[0].out_path)).active["A1"].value
        self.assertNotRegex(value, "(?i)петров|серге")
        self.assertEqual(len(person_ids(value)), 1, value)
        env.close()

    def test_common_words_next_to_names_stay(self):
        out = txt("Письмо для Валерия Громова. Валерия подтвердила встречу в среду.")
        self.assertIn("встречу в среду", out)


if __name__ == "__main__":
    unittest.main()
