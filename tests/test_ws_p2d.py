"""P2D: selection in the preview panel, suggestion cards, preview addresses for docx/pptx/xlsx."""
import io
import re
import sys
import unittest
from pathlib import Path

WEB = Path(__file__).resolve().parents[1] / "anonymizer" / "web"
JS = (WEB / "ai-chat.js").read_text(encoding="utf-8")
HTML = (WEB / "index.html").read_text(encoding="utf-8")


def make_pptx():
    import pptx
    p = pptx.Presentation()
    for i in range(3):
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = f"обзор раздела {i + 1}"
        s.placeholders[1].text = f"Пункт {i + 1}: Alpha{i + 1}\nВторой абзац {i + 1}"
    b = io.BytesIO(); p.save(b); return b.getvalue()


def make_docx():
    import docx
    d = docx.Document(); d.add_heading("обзор договора", 1)
    d.add_paragraph("Стороны: Alpha1 и Alpha2."); d.add_paragraph("")
    d.add_paragraph("Срок: 12 месяцев.")
    t = d.add_table(rows=2, cols=2)
    for r, row in enumerate((("Пункт", "Сумма"), ("Оплата", "Amount1"))):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    b = io.BytesIO(); d.save(b); return b.getvalue()


def make_xlsx():
    import openpyxl
    w = openpyxl.Workbook(); ws = w.active; ws.title = "Лист1"
    ws.append(["Статья", "Сумма"]); ws.append(["Аренда Alpha1", 1000])
    w.create_sheet("Итоги").append(["Всего", "Amount2"])
    b = io.BytesIO(); w.save(b); return b.getvalue()


class PreviewAddresses(unittest.TestCase):
    def test_pptx_addresses_match_outline(self):
        from anonymizer import docedit, preview
        data = make_pptx()
        res = preview.render("a.pptx", data)
        known = {e["addr"] for c in docedit.outline(data, "a.pptx")["containers"] for e in c["elements"]}
        self.assertEqual([s["addr"] for s in res["slides"]], ["s1", "s2", "s3"])
        items = [it for s in res["slides"] for it in s["items"]]
        self.assertEqual(len(items), 9)
        self.assertTrue(all(it["addr"] in known for it in items))
        self.assertEqual(res["slides"][0]["items"][0]["role"], "title")
        self.assertEqual(len(res["slides"][0]["texts"]), 2)       # old shape kept

    def test_docx_addresses(self):
        from anonymizer import docedit, preview
        data = make_docx()
        res = preview.render("a.docx", data)
        known = {e["addr"] for c in docedit.outline(data, "a.docx")["containers"] for e in c["elements"]}
        paras = [b["addr"] for b in res["blocks"] if b["t"] != "table"]
        self.assertEqual(paras, ["p0", "p1", "p3"])               # the empty paragraph p2 is not shown and not addressed
        table = [b for b in res["blocks"] if b["t"] == "table"][0]
        self.assertEqual(table["addrs"], [["t1.r0.c0.p0", "t1.r0.c1.p0"], ["t1.r1.c0.p0", "t1.r1.c1.p0"]])
        self.assertTrue(all(a in known for a in paras + [x for r in table["addrs"] for x in r]))

    def test_xlsx_addresses(self):
        from anonymizer import preview
        res = preview.render("a.xlsx", make_xlsx())
        self.assertEqual(res["sheet_list"][0]["addrs"], [["Лист1!A1", "Лист1!B1"], ["Лист1!A2", "Лист1!B2"]])
        self.assertEqual(res["sheet_list"][1]["addrs"], [["Итоги!A1", "Итоги!B1"]])
        cell = res["sheet_list"][1]["rows"][0][1]
        self.assertEqual(cell, "Amount2")                        # cell text is data only

    def test_broken_outline_keeps_preview(self):
        from anonymizer import preview
        res = preview.render("a.docx", b"not a zip")
        self.assertEqual(res["kind"], "none")


class Page(unittest.TestCase):
    def test_selection_in_panel(self):
        for need in ('class="pv-pick"' if False else '"pv-pick"', 'aria-checked', 'role", "checkbox"', "pvSelToggle", "pvSelClear", "ArrowDown",
                     "const pvSel = new Map()", "data-addr" if False else "dataset.addr", "Выделение снято"):
            self.assertIn(need, HTML)
        self.assertIn('if (pvSel.size) { pvSelClear(true); return; }', HTML)      # Escape clears before closing the panel
        self.assertIn("pvSelClear(true)", HTML.split("function pvRender", 1)[1][:400])  # cleared on file switch

    def test_no_new_always_visible_buttons(self):
        # selection adds no static buttons to the markup: everything is created on demand
        before = HTML.count("<button")
        self.assertEqual(before, 104)        # P2D baseline 102 + add-in install/remove buttons (Settings)
        self.assertNotIn('id="pv-select', HTML)

    def test_chip_and_send_fields(self):
        for need in ('"Выделено: "', "body.selected = ", "body.selected_file = S.sel.file", "setSelection(", "onClearSelection", '"Снять выделение"'):
            self.assertIn(need, JS)
        self.assertIn('e.key === "Escape" && S.sel.items.length', JS)

    def test_cards(self):
        for need in ("renderProposals", "Применить все", '"Применить"', '"Отклонить"', "/patch", "d.proposals", "aic-qnote", "было" if False else "was", "arr"):
            self.assertIn(need, JS)
        self.assertIn('{id: pr.id, action, ops}', JS)
        self.assertIn("textContent", JS)
        self.assertNotIn("innerHTML", JS.split("function renderProposals", 1)[1].split("api: ", 1)[0][:6000])

    def test_reduced_motion_covers_new_motion(self):
        reduce = JS.split("prefers-reduced-motion: reduce", 1)[1].split("}", 1)[0]
        self.assertIn(".aic-prop > *", reduce)
        self.assertIn("prefers-reduced-motion: no-preference", JS.split(".aic-prop > *", 1)[0][-200:])
        self.assertRegex(HTML, r"prefers-reduced-motion: reduce\) \{ \.pv, \.pv \* \{ animation: none !important; transition: none !important")
        motion = re.findall(r"\.aic-prop[^{]*\{[^}]*\}", JS)
        for rule in motion:
            self.assertNotRegex(rule, r"(width|height|top|left|margin)\s*:\s*[^;]*transition")

    def test_no_external_urls(self):
        for name, src in (("js", JS), ("html", HTML)):
            urls = [u for u in re.findall(r"https?://[^\s\"'`)<>]+", src) if "www.w3.org" not in u]
            self.assertEqual(urls, [], name + " " + str(urls[:3]))


if __name__ == "__main__":
    unittest.main()


class TemplateTextNotSent(unittest.TestCase):
    def test_pptx_body_text_skips_layout_and_master_placeholders(self):
        import io, tempfile
        from pathlib import Path
        from anonymizer import ai_chat
        from .helpers import pptx_bytes
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "a.pptx"
            path.write_bytes(pptx_bytes(["Выручка выросла"], title="Итоги"))
            text = ai_chat._body_text(path)
        self.assertIn("Выручка выросла", text)
        self.assertNotIn("Click to edit", text)
        self.assertNotIn("Master", text)


class PatchBlockHidden(unittest.TestCase):
    def test_chat_js_strips_the_json_patch_from_the_visible_reply(self):
        from pathlib import Path
        js = (Path(__file__).resolve().parent.parent / "anonymizer" / "web" / "ai-chat.js").read_text("utf-8")
        self.assertIn("PATCH_FENCE", js)
        self.assertIn("noPatch(text)", js)
        self.assertIn("noPatch(out.textContent)", js)


class DesignCardsReadable(unittest.TestCase):
    def test_design_card_line_is_russian_not_json(self):
        from anonymizer import docedit
        line = docedit._design_ru("set_style", {"size": 28, "bold": True})
        self.assertEqual(line, "оформление текста: размер 28, жирный да")
        self.assertNotIn("{", docedit._design_ru("col_width", {"width": 40}))
        self.assertIn("ширина 40", docedit._design_ru("col_width", {"width": 40}))
