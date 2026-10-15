"""W10b: остаток замены картинок целиком. Проверки смотрят в байты ВЫХОДНЫХ файлов (рекурсивно: вложенные zip,
потоки PDF, data URI), а не в отчёт программы."""
from __future__ import annotations

import base64
import io
import os
import re
import struct
import unittest
import zipfile
from pathlib import Path

import fitz
from PIL import Image

from anonymizer import imagereplace
from anonymizer.service import Upload
from tests.helpers import docx_bytes
from tests.test_ws_w10 import W10Base, noise_image, windows

M_EXIF = b"SECRET-EXIF-GPS-55.7558N-37.6173E"
M_XMP = b"SECRET-XMP-CREATOR-Petrov"
M_IPTC = b"SECRET-IPTC-CAPTION"
M_ICC = b"SECRET-ICC-DESCRIPTION"
M_TRAIL = b"SECRET-TRAILER-AFTER-EOI"
M_THUMB_PIXELS = b"THUMBNAIL-MARKER-JPEG"
MARKERS = (M_EXIF, M_XMP, M_IPTC, M_ICC, M_TRAIL, M_THUMB_PIXELS)
META_SIGNATURES = (b"Exif\x00\x00", b"http://ns.adobe.com/xap", b"<x:xmpmeta", b"ICC_PROFILE", b"iCCP", b"tEXt", b"iTXt",
                   b"eXIf", b"Photoshop 3.0", b"GPSLatitude")


def _jpeg_segment(marker: int, payload: bytes) -> bytes:
    return b"\xff" + bytes([marker]) + struct.pack(">H", len(payload) + 2) + payload


def jpeg_rich(seed: int, size=(96, 80)) -> bytes:
    """JPEG с EXIF (GPS-строка и вложенная JPEG-миниатюра), XMP, IPTC, ICC-блоком и данными после EOI."""
    buf = io.BytesIO()
    noise_image(seed, size).save(buf, "JPEG", quality=90)
    base = buf.getvalue()
    thumb = io.BytesIO()
    noise_image(seed + 1000, (16, 12)).save(thumb, "JPEG")
    thumb_bytes = thumb.getvalue() + M_THUMB_PIXELS
    segs = (_jpeg_segment(0xE1, b"Exif\x00\x00II*\x00\x08\x00\x00\x00" + M_EXIF + b"GPSLatitude" + thumb_bytes)
            + _jpeg_segment(0xE1, b"http://ns.adobe.com/xap/1.0/\x00<x:xmpmeta>" + M_XMP + b"</x:xmpmeta>")
            + _jpeg_segment(0xED, b"Photoshop 3.0\x00" + M_IPTC)
            + _jpeg_segment(0xE2, b"ICC_PROFILE\x00\x01\x01" + M_ICC))
    return base[:2] + segs + base[2:] + M_TRAIL


def _chunk(kind: bytes, data: bytes) -> bytes:
    import zlib
    return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)


def png_rich(seed: int, size=(96, 80)) -> bytes:
    buf = io.BytesIO()
    noise_image(seed, size).save(buf, "PNG")
    base = buf.getvalue()
    iend = base.rindex(b"IEND") - 4
    extra = (_chunk(b"tEXt", b"Author\x00" + M_EXIF) + _chunk(b"iTXt", b"XML:com.adobe.xmp\x00\x00\x00\x00\x00<x:xmpmeta>" + M_XMP)
             + _chunk(b"eXIf", b"MM\x00*" + M_EXIF) + _chunk(b"iCCP", b"icc\x00\x00" + M_ICC))
    return base[:iend] + extra + base[iend:] + M_TRAIL


def tiff_rich(seed: int) -> bytes:
    buf = io.BytesIO()
    noise_image(seed).save(buf, "TIFF", description=M_EXIF.decode(), software=M_XMP.decode())
    return buf.getvalue()


def webp_rich(seed: int) -> bytes:
    buf = io.BytesIO()
    exif = Image.Exif()
    exif[0x010E] = M_EXIF.decode()
    noise_image(seed).save(buf, "WEBP", lossless=True, exif=exif, xmp=b"<x:xmpmeta>" + M_XMP + b"</x:xmpmeta>")
    return buf.getvalue()


def deep_blobs(data: bytes, depth: int = 0) -> list[bytes]:
    """Все байты файла и всё, что из него достаётся: части zip (вложенные тоже), потоки PDF, data URI."""
    out = [data]
    if depth > 6:
        return out
    if data[:4] == b"PK\x03\x04":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                for n in z.namelist():
                    out += deep_blobs(z.read(n), depth + 1)
        except zipfile.BadZipFile:
            pass
    elif data[:4] == b"%PDF":
        try:
            doc = fitz.open(stream=data)
            for x in range(1, doc.xref_length()):
                for getter in (doc.xref_stream_raw, doc.xref_stream):
                    try:
                        out.append(getter(x) or b"")
                    except Exception:
                        pass
            doc.close()
        except Exception:
            pass
    else:
        for m in re.finditer(rb"base64,([A-Za-z0-9+/=]{64,})", data):
            try:
                out += deep_blobs(base64.b64decode(m.group(1)), depth + 1)
            except Exception:
                pass
    return out


def build_zip(base: bytes, extra: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base)) as zin, zipfile.ZipFile(buf, "w") as zout:
        for i in zin.infolist():
            zout.writestr(i, zin.read(i.filename))
        for name, raw in extra.items():
            zout.writestr(name, raw)
    return buf.getvalue()


def plain_zip(files: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n, raw in files.items():
            z.writestr(n, raw)
    return buf.getvalue()


class ProofBase(W10Base):
    def prove_no_original(self, out: Path, originals: list[bytes]):
        blobs = deep_blobs(out.read_bytes())
        joined = b"\n".join(blobs)
        for raw in originals:
            for w in windows(raw, 6):
                self.assertNotIn(w, joined, "кусок оригинальной картинки остался в выходе")
        for marker in MARKERS:
            self.assertNotIn(marker, joined, f"метаданные остались: {marker!r}")
        for sig in META_SIGNATURES:
            self.assertNotIn(sig, joined, f"сигнатура метаданных осталась: {sig!r}")
        # миниатюра внутри JPEG: в каждом отдельном JPEG-потоке ровно один SOI-маркер
        for blob in blobs:
            if blob[:3] == b"\xff\xd8\xff" and len(blob) < 5_000_000:
                self.assertEqual(blob.count(b"\xff\xd8\xff"), 1, "во втором JPEG-потоке видна миниатюра")

    def prove_restore(self, out: Path, originals: list[bytes]):
        back = self.restored(out, out.name)
        joined = b"\n".join(deep_blobs(back.read_bytes()))
        for raw in originals:
            self.assertIn(raw, joined, "оригинал не вернулся побайтно")
        return back


class TestStandaloneRaw(ProofBase):
    def test_png_jpg_tiff_webp(self):
        for name, raw in (("a.png", png_rich(1)), ("b.jpg", jpeg_rich(2)), ("c.tif", tiff_rich(3)), ("d.tiff", tiff_rich(4)),
                          ("e.webp", webp_rich(5))):
            with self.subTest(name=name):
                job, out = self.run_one(name, raw)
                data = out.read_bytes()
                self.assertNotEqual(data, raw)
                self.prove_no_original(out, [raw])
                Image.open(io.BytesIO(data)).load()
                self.assertEqual(self.restored(out, out.name).read_bytes(), raw)

    def test_tiff_webp_not_refused(self):
        for name, raw in (("a.tif", tiff_rich(7)), ("b.webp", webp_rich(8))):
            f = self.env.anonymize((name, raw)).files[0]
            self.assertTrue(f.out_path, f"{name}: {f.messages}")

    def test_redact_mode_refuses_tiff(self):
        """В режиме «закрасить текст» TIFF/WebP по-прежнему не поддержаны: отказ, а не пропуск."""
        f = self.env.anonymize(("a.tif", tiff_rich(9)), options={"image_mode": "redact_text"}).files[0]
        self.assertFalse(f.out_path)


class TestContainers(ProofBase):
    def docx_with(self, extra):
        return build_zip(docx_bytes(["hello"]), extra)

    def test_docx_images_outside_media(self):
        png, jpg = png_rich(11), jpeg_rich(12)
        data = self.docx_with({"word/embeddings/image1.png": png, "word/charts/photo.jpeg": jpg,
                               "word/media/image7.png": png_rich(13)})
        job, out = self.run_one("o.docx", data)
        self.prove_no_original(out, [png, jpg])
        self.prove_restore(out, [png, jpg])

    def test_docx_nested_zip_with_images(self):
        png, jpg, tif = png_rich(21), jpeg_rich(22), tiff_rich(23)
        deep = plain_zip({"deep/pic.png": png_rich(24)})
        inner = plain_zip({"a/pic.png": png, "b/pic.jpg": jpg, "c/pic.tif": tif, "d/deep.zip": deep})
        data = self.docx_with({"word/embeddings/pack.zip": inner})
        f = self.env.anonymize(("z.docx", data)).files[0]
        if f.out_path:
            out = Path(f.out_path)
            self.prove_no_original(out, [png, jpg, tif, png_rich(24)])
            self.prove_restore(out, [png, jpg, tif, png_rich(24)])
        else:
            self.assertIn("картин", str(f.messages).lower())     # отказ допустим, молча пропустить нельзя

    def test_docx_with_embedded_workbook_images(self):
        png = png_rich(31)
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XImage
        wb = Workbook(); wb.active["A1"] = "x"
        wb.active.add_image(XImage(io.BytesIO(png)), "C3")
        wx = io.BytesIO(); wb.save(wx)
        data = self.docx_with({"word/embeddings/Microsoft_Excel_Sheet1.xlsx": wx.getvalue()})
        f = self.env.anonymize(("w.docx", data)).files[0]
        if f.out_path:
            out = Path(f.out_path)
            self.prove_no_original(out, [png])
            self.prove_restore(out, [png])
        else:
            self.assertIn("картин", str(f.messages).lower())

    def test_xlsx_chart_and_drawing_images(self):
        png, jpg = png_rich(41), jpeg_rich(42)
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.drawing.image import Image as XImage
        wb = Workbook(); ws = wb.active
        for i in range(1, 5):
            ws.append([f"r{i}", i * 10])
        ch = BarChart(); ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=4)); ws.add_chart(ch, "E2")
        ws.add_image(XImage(io.BytesIO(png)), "A8")
        b = io.BytesIO(); wb.save(b)
        data = build_zip(b.getvalue(), {"xl/charts/fill.jpeg": jpg})
        job, out = self.run_one("c.xlsx", data)
        self.prove_no_original(out, [png, jpg])
        self.prove_restore(out, [png, jpg])

    def test_pptx_embedded_workbook_and_chart_images(self):
        png, jpg = png_rich(51), jpeg_rich(52)
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XImage
        wb = Workbook(); wb.active["A1"] = "x"
        wb.active.add_image(XImage(io.BytesIO(png)), "C3")
        wx = io.BytesIO(); wb.save(wx)
        from pptx import Presentation
        prs = Presentation(); prs.slides.add_slide(prs.slide_layouts[6])
        pb = io.BytesIO(); prs.save(pb)
        data = build_zip(pb.getvalue(), {"ppt/embeddings/Microsoft_Excel_Sheet1.xlsx": wx.getvalue(),
                                         "ppt/charts/bg.jpg": jpg})
        f = self.env.anonymize(("e.pptx", data)).files[0]
        if f.out_path:
            out = Path(f.out_path)
            self.prove_no_original(out, [png, jpg])
            self.prove_restore(out, [png, jpg])
        else:
            self.assertIn("картин", str(f.messages).lower())

    def test_ole_bin_with_image_signature_refused(self):
        """Двоичный объект OLE нельзя разобрать: если в нём видна картинка, файл не выдаётся."""
        png = png_rich(61)
        ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64 + png + b"\x00" * 64
        data = self.docx_with({"word/embeddings/oleObject1.bin": ole})
        f = self.env.anonymize(("b.docx", data)).files[0]
        if f.out_path:
            joined = b"\n".join(deep_blobs(Path(f.out_path).read_bytes()))
            for w in windows(png, 6):
                self.assertNotIn(w, joined)
        else:
            self.assertIn("картин", str(f.messages).lower())

    def test_image_in_zip_fail_closed_when_store_full(self):
        os.environ["ANON_IMAGE_STORE_MB"] = "0.00001"
        try:
            data = self.docx_with({"word/embeddings/pack.zip": plain_zip({"a.png": png_rich(71)})})
            f = self.env.anonymize(("l.docx", data)).files[0]
            self.assertFalse(f.out_path)
        finally:
            del os.environ["ANON_IMAGE_STORE_MB"]


class TestHtmlAndPdf(ProofBase):
    def test_html_data_uri_png_jpg_webp(self):
        raws = [png_rich(81), jpeg_rich(82), webp_rich(83)]
        mimes = ["png", "jpeg", "webp"]
        body = "".join(f'<img src="data:image/{m};base64,{base64.b64encode(r).decode()}">' for m, r in zip(mimes, raws))
        job, out = self.run_one("h.html", f"<html><body><p>x</p>{body}</body></html>".encode())
        self.prove_no_original(out, raws)
        back = self.restored(out, out.name).read_text()
        for r in raws:
            self.assertIn(base64.b64encode(r).decode(), back)

    def test_pdf_with_rich_images(self):
        png, jpg = png_rich(91), jpeg_rich(92)
        doc = fitz.open(); page = doc.new_page()
        page.insert_text((72, 60), "Contract", fontsize=12)
        page.insert_image(fitz.Rect(72, 100, 272, 260), stream=png)
        page.insert_image(fitz.Rect(300, 100, 500, 260), stream=jpg)
        data = doc.tobytes(); doc.close()
        job, out = self.run_one("r.pdf", data)
        self.prove_no_original(out, [png, jpg, ])
        o = fitz.open(stream=data)
        raw_streams = [o.xref_stream_raw(i[0]) for i in o[0].get_images()]
        blob = b"\n".join(deep_blobs(out.read_bytes()))
        for s in raw_streams:
            for w in windows(s, 6):
                self.assertNotIn(w, blob)
        back = fitz.open(self.restored(out, out.name))
        got = [back.xref_stream(i[0]) for i in back[0].get_images()]
        for i in o[0].get_images():
            self.assertIn(o.xref_stream(i[0]), got, "поток картинки не вернулся побайтно")


class TestPlaceholderOcrCost(W10Base):
    def _scan(self, pages=3):
        doc = fitz.open()
        for i in range(pages):
            p = doc.new_page(); b = io.BytesIO(); noise_image(i, (600, 800)).save(b, "PNG")
            p.insert_image(p.rect, stream=b.getvalue())
        data = doc.tobytes(); doc.close()
        return data

    def test_placeholder_pages_not_ocrd_but_real_images_still_are(self):
        from anonymizer import ocr
        calls = {"n": 0}
        orig = ocr.ocr_image_layers

        def counted(*a, **k):
            calls["n"] += 1
            return orig(*a, **k)
        ocr.ocr_image_layers = counted
        try:
            data = self._scan(3)
            self.env.anonymize(("s.pdf", data))
            self.assertEqual(calls["n"], 0, "заглушки читаются распознаванием зря")
            calls["n"] = 0
            self.env.anonymize(("s2.pdf", data), options={"image_mode": "redact_text"})
            if ocr.available():
                self.assertGreaterEqual(calls["n"], 3, "настоящие картинки перестали читаться")
        finally:
            ocr.ocr_image_layers = orig


if __name__ == "__main__":
    unittest.main()
