"""S6B: картинки с логотипом: режим по умолчанию не оставляет пикселей, redact_text предупреждает уровнем warn."""
from __future__ import annotations

import io
import unittest
import zipfile

import fitz
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches as PInches

from anonymizer import imagereplace, settings_schema
from anonymizer.webapp import App  # noqa: F401  (проверка импорта)
from tests.helpers import Env

NOISE = bytes(range(256))


def logo_png() -> bytes:
    img = Image.new("RGB", (240, 90), (200, 30, 60))
    d = ImageDraw.Draw(img)
    d.rectangle((6, 6, 233, 83), outline=(255, 255, 255), width=3)
    d.text((24, 34), "VOLKOVA STROY LLC", fill=(255, 255, 255))
    return _png(img)


def _png(img) -> bytes:
    buf = io.BytesIO(); img.save(buf, "PNG"); return buf.getvalue()


def pptx_file(png: bytes) -> bytes:
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(png), PInches(1), PInches(1), PInches(3))
    out = io.BytesIO(); prs.save(out); return out.getvalue()


def docx_file(png: bytes) -> bytes:
    doc = Document(); doc.add_paragraph("Договор поставки")
    doc.add_picture(io.BytesIO(png), width=Inches(3))
    out = io.BytesIO(); doc.save(out); return out.getvalue()


def xlsx_file(png: bytes) -> bytes:
    wb = Workbook(); ws = wb.active; ws["A1"] = "Отчёт"
    ws.add_image(XlImage(io.BytesIO(png)), "C3")
    out = io.BytesIO(); wb.save(out); return out.getvalue()


def pdf_file(png: bytes) -> bytes:
    doc = fitz.open(); page = doc.new_page(); page.insert_text((72, 72), "Report")
    page.insert_image(fitz.Rect(72, 100, 312, 190), stream=png)
    return doc.tobytes()


def media_pixels(data: bytes, ext: str) -> list[bytes]:
    """Декодированные пиксели всех картинок внутри файла."""
    out = []
    if ext == "pdf":
        doc = fitz.open(stream=data, filetype="pdf")
        for page in doc:
            for info in page.get_images(full=True):
                pix = fitz.Pixmap(doc, info[0])
                if pix.n > 3: pix = fitz.Pixmap(fitz.csRGB, pix)
                out.append(pix.samples)
        return out
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for n in z.namelist():
            if "/media/" in n and n.lower().endswith(".png"):
                out.append(Image.open(io.BytesIO(z.read(n))).convert("RGB").tobytes())
    return out


def levels(job) -> list[tuple[str, str]]:
    return [(m["level"], m["text"]) for f in job.files for m in f.messages]


FORMATS = {"pptx": pptx_file, "docx": docx_file, "xlsx": xlsx_file, "pdf": pdf_file}


class DefaultReplaceKeepsNoPixels(unittest.TestCase):
    def test_all_formats(self):
        png = logo_png(); original = Image.open(io.BytesIO(png)).convert("RGB").tobytes()
        for ext, make in FORMATS.items():
            with self.subTest(ext), _env() as env:
                job = env.anonymize((f"doc.{ext}", make(png)))
                data = env.output_bytes(job)
                self.assertNotIn(original, media_pixels(data, ext), "оригинальные пиксели остались")
                self.assertFalse([t for lv, t in levels(job) if lv == "warn" and "остаются в файле видимыми" in t])


class RedactTextWarns(unittest.TestCase):
    def test_warn_level(self):
        png = logo_png()
        for ext in ("pptx", "docx", "xlsx"):
            with self.subTest(ext), _env(image_mode="redact_text") as env:
                job = env.anonymize((f"doc.{ext}", FORMATS[ext](png)))
                msgs = levels(job)
                self.assertTrue([1 for lv, t in msgs if lv == "warn" and "остаются в файле видимыми" in t], msgs)
                self.assertFalse([1 for lv, t in msgs if lv == "info" and "закрашены" in t and "Данные на картинках" in t])


class Schema(unittest.TestCase):
    def test_label_unsafe_and_default_first(self):
        opt = next(o for o in settings_schema.OPTIONS if o.key == "image_mode")
        self.assertEqual(opt.default, "replace"); self.assertEqual(opt.choices[0][0], "replace")
        self.assertIn("менее безопасно", dict(opt.choices)["redact_text"])

    def test_notice_texts(self):
        self.assertIn("векторно", imagereplace.WARN_VECTOR_LOGO)
        self.assertIn("видимыми", imagereplace.warn_kept_pixels(2))


class StoredFlag(unittest.TestCase):
    def test_state_exposes_flag(self):
        from anonymizer.webapp import App
        for mode, want in (("redact_text", True), ("replace", False)):
            with _env(image_mode=mode) as env:
                self.assertIs(_state(env)["image_mode_unsafe"], want)


class _env:
    def __init__(self, **p): self.e = Env(**p)
    def __enter__(self): return self.e
    def __exit__(self, *a): self.e.close()


def _state(env):
    from anonymizer import webapp
    app = webapp.App.__new__(webapp.App); app.vault = env.vault
    app.chats = type("C", (), {"__getattr__": lambda s, n: (lambda *a, **k: [])})()
    return webapp.App.state(app)


if __name__ == "__main__":
    unittest.main()
