"""Finance restore: scripted suite of AI-edit scenarios on a masked workbook.

Each scenario edits the masked file the way a model/tool would (openpyxl on the returned file), runs the real restore and checks
the numbers against what a human would expect in ORIGINAL units. Contract: every scenario is either restored correctly or
flagged in the report; a scenario with wrong numbers and status ``ok`` fails the suite.
"""
from __future__ import annotations

import io
import unittest
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from anonymizer.service import Upload
from tests.helpers import Env

PCT = "0.0%"
RUB = '#,##0 "₽"'
ORIG = {"Выручка": (1_250_000, 1_400_000, 1_525_000), "Себестоимость": (700_000, 770_000, 850_000),
        "Расходы на маркетинг": (120_000, 130_000, 140_000)}


def base_book() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "P&L"
    ws.append(["Статья", "Q1", "Q2", "Q3", "Итого, руб"])
    for label, vals in ORIG.items():
        ws.append([label, *vals])
    ws.insert_rows(4)
    ws["A4"] = "Валовая прибыль"
    for c in "BCD":
        ws[f"{c}4"] = f"={c}2-{c}3"
    # rows now: 2 Выручка, 3 Себестоимость, 4 Валовая прибыль, 5 Расходы на маркетинг
    ws["A6"] = "EBITDA"
    ws["A7"] = "Маржа EBITDA, %"
    ws["A8"], ws["B8"] = "Ставка налога, %", 0.2
    ws["A9"] = "Налог на прибыль"
    for c in "BCD":
        ws[f"{c}6"] = f"={c}4-{c}5"
        ws[f"{c}7"] = f"={c}6/{c}2"
        ws[f"{c}7"].number_format = PCT
        ws[f"{c}9"] = f"={c}6*$B$8"
    ws["B8"].number_format = "0%"
    for r in (2, 3, 5):
        ws[f"E{r}"] = f"=SUM(B{r}:D{r})"
        for c in "BCDE":
            ws[f"{c}{r}"].number_format = RUB
    prices = wb.create_sheet("Цены")
    prices.append(["Товар", "Цена, руб"])
    prices.append(["Товар А", 1250.5])
    prices.append(["Товар Б", 899])
    hidden = wb.create_sheet("Служебный")
    hidden["A1"], hidden["B1"] = "Резерв, руб", 777_333
    hidden.sheet_state = "hidden"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def row_of(ws, label, nth=1, col="A"):
    seen = 0
    for r in range(1, ws.max_row + 1):
        if ws[f"{col}{r}"].value == label:
            seen += 1
            if seen == nth:
                return r
    raise KeyError(label)


def at(wb, sheet, label, col, nth=1):
    ws = wb[sheet]
    return ws[f"{col}{row_of(ws, label, nth)}"].value


# -- сценарии: имя, правка(wb), ожидания [(лист, подпись, столбец, nth, значение)], режим ----------------------------------

def _k(wb):
    """Секретный коэффициент по замаскированной ячейке B2 (исходная выручка 1 250 000)."""
    return wb["P&L"]["B2"].value / 1_250_000


def _mul(ws, ref, k):
    ws[ref] = ws[ref].value * k


def s_noedit(wb):
    pass


def s_edit_one(wb):
    _mul(wb["P&L"], "B2", 1.1)


def s_edit_many(wb):
    ws = wb["P&L"]
    for ref in ("B2", "C2", "D2", "B3", "C5"):
        _mul(ws, ref, 0.9)


def s_formula_edit(wb):
    wb["P&L"]["B4"] = "=B2-B3-B5"


def s_paste_money(wb):
    ws = wb["P&L"]
    ws["B4"] = ws["B2"].value - ws["B3"].value


def s_paste_ratio(wb):
    wb["P&L"]["B7"] = 0.344


def s_insert_row(wb):
    wb["P&L"].insert_rows(3)
    wb["P&L"]["A3"] = "Скидки"


def s_delete_row(wb):
    ws = wb["P&L"]
    ws.delete_rows(row_of(ws, "Расходы на маркетинг"))


def s_new_row_label(wb):
    ws = wb["P&L"]
    ws["A11"], ws["B11"], ws["C11"] = "Прочие расходы", 100_000 * _k(wb), 120_000 * _k(wb)


def s_new_row_unknown_label(wb):
    ws = wb["P&L"]
    ws["A11"], ws["B11"] = "Фонд Альфа", 96_050 * _k(wb)


def s_new_row_unlabeled(wb):
    ws = wb["P&L"]
    ws["B12"], ws["C12"] = 500_000.0, 650_000.0


def s_new_sheet_label(wb):
    ws = wb.create_sheet("AI_Forecast")
    ws["A1"], ws["B1"] = "Выручка прогноз", 432_255 * _k(wb)


def s_new_sheet_unlabeled(wb):
    ws = wb.create_sheet("Расчёт")
    ws["A1"], ws["B1"] = "Foo", 123_456.0


def s_rename_sheet(wb):
    wb["P&L"].title = "ОПУ"


def s_rename_header(wb):
    wb["P&L"]["B1"] = "Q1 2024"


def s_add_formula_col(wb):
    ws = wb["P&L"]
    ws["F1"] = "Прогноз"
    ws["F2"] = "=D2*1.1"


def s_add_value_col(wb):
    ws = wb["P&L"]
    ws["F1"] = "Q4"
    ws["F2"], ws["F3"] = 1_600_000 * _k(wb), 900_000 * _k(wb)


def s_swap_rows(wb):
    ws = wb["P&L"]
    for c in "ABCD":
        ws[f"{c}2"].value, ws[f"{c}3"].value = ws[f"{c}3"].value, ws[f"{c}2"].value


def s_new_assumption_pct(wb):
    ws = wb.create_sheet("Assumptions")
    ws["A1"], ws["B1"] = "Рост выручки, %", 0.08
    ws["B1"].number_format = "0%"


def s_pct_cell(wb):
    ws = wb["P&L"]
    ws["A12"], ws["B12"] = "Доля маркетинга", 0.12
    ws["B12"].number_format = PCT


def s_format_only(wb):
    from openpyxl.styles import PatternFill
    wb["P&L"]["B2"].fill = PatternFill("solid", start_color="FFFF00")


def s_edit_tax(wb):
    wb["P&L"]["B8"] = 0.25


def s_clear_formula(wb):
    wb["P&L"]["B4"] = None


def s_formula_const(wb):
    wb["P&L"]["B6"] = "=B4-B5+500000"


def s_hidden_edit(wb):
    _mul(wb["Служебный"], "B1", 2)


def s_dup_labels(wb):
    ws = wb["P&L"]
    ws["A11"], ws["B11"] = "Прочее", 111_111.0 * _k(wb)
    ws["A12"], ws["B12"] = "Прочее", 222_222.0 * _k(wb)
    ws["B12"].value = ws["B12"].value * 2


def s_double_all(wb):
    ws = wb["P&L"]
    for r in (2, 3, 5):
        for c in "BCD":
            _mul(ws, f"{c}{r}", 2)


def s_insert_and_edit(wb):
    ws = wb["P&L"]
    k = _k(wb)
    ws.insert_rows(2)
    ws["A2"] = "Возвраты"
    ws["B2"] = 50_000 * k
    _mul(ws, "B3", 1.05)


def s_pasted_table(wb):
    ws = wb["P&L"]
    ws["A14"], ws["B14"] = "Прогноз", "Q4"
    ws["A15"], ws["B15"] = "Выручка", 7_000_000.0


def s_price_edit(wb):
    _mul(wb["Цены"], "B2", 1.2)


def s_remove_sheet(wb):
    del wb["Цены"]


SCENARIOS = [
    # (name, edit, expectations, mode)  expectations: (sheet, label, col, nth, value)
    ("no_edit", s_noedit, [("P&L", "Выручка", "B", 1, 1_250_000), ("Служебный", "Резерв, руб", "B", 1, 777_333)], "ok"),
    ("edit_one_input", s_edit_one, [("P&L", "Выручка", "B", 1, 1_375_000), ("P&L", "Выручка", "C", 1, 1_400_000)], "ok"),
    ("edit_many_inputs", s_edit_many, [("P&L", "Выручка", "B", 1, 1_125_000), ("P&L", "Выручка", "D", 1, 1_372_500),
                                       ("P&L", "Себестоимость", "B", 1, 630_000), ("P&L", "Расходы на маркетинг", "C", 1, 117_000),
                                       ("P&L", "Себестоимость", "C", 1, 770_000)], "ok"),
    ("formula_edit", s_formula_edit, [("P&L", "Валовая прибыль", "B", 1, "=B2-B3-B5"), ("P&L", "Выручка", "B", 1, 1_250_000)], "ok"),
    ("paste_money_over_formula", s_paste_money, [("P&L", "Валовая прибыль", "B", 1, 550_000)], "ok"),
    ("paste_ratio_over_formula", s_paste_ratio, [("P&L", "Маржа EBITDA, %", "B", 1, 0.344)], "ok"),
    ("insert_row", s_insert_row, [("P&L", "Выручка", "B", 1, 1_250_000), ("P&L", "Расходы на маркетинг", "C", 1, 130_000),
                                  ("P&L", "Себестоимость", "D", 1, 850_000)], "ok"),
    ("delete_row", s_delete_row, [("P&L", "Выручка", "B", 1, 1_250_000), ("P&L", "Себестоимость", "D", 1, 850_000)], "ok"),
    ("new_row_money_label", s_new_row_label, [("P&L", "Прочие расходы", "B", 1, 100_000), ("P&L", "Прочие расходы", "C", 1, 120_000)], "ok"),
    ("new_row_unknown_label_money_column", s_new_row_unknown_label, [("P&L", "Фонд Альфа", "B", 1, 96_050)], "ok"),
    ("new_unlabeled_row_money_column", s_new_row_unlabeled, [], "ok"),
    ("new_sheet_money_label", s_new_sheet_label, [("AI_Forecast", "Выручка прогноз", "B", 1, 432_255)], "ok"),
    ("new_sheet_unlabeled_numbers", s_new_sheet_unlabeled, [], "flag"),
    ("rename_sheet", s_rename_sheet, [("ОПУ", "Выручка", "B", 1, 1_250_000), ("ОПУ", "Расходы на маркетинг", "D", 1, 140_000)], "ok"),
    ("rename_column_header", s_rename_header, [("P&L", "Выручка", "B", 1, 1_250_000), ("P&L", "Выручка", "C", 1, 1_400_000)], "ok"),
    ("add_formula_column", s_add_formula_col, [("P&L", "Выручка", "D", 1, 1_525_000), ("P&L", "Выручка", "F", 1, "=D2*1.1")], "ok"),
    ("add_value_column", s_add_value_col, [("P&L", "Выручка", "F", 1, 1_600_000), ("P&L", "Себестоимость", "F", 1, 900_000)], "ok"),
    ("swap_rows", s_swap_rows, [("P&L", "Выручка", "B", 1, 1_250_000), ("P&L", "Себестоимость", "C", 1, 770_000)], "ok"),
    ("new_percent_assumption", s_new_assumption_pct, [("Assumptions", "Рост выручки, %", "B", 1, 0.08)], "ok"),
    ("new_percent_cell", s_pct_cell, [("P&L", "Доля маркетинга", "B", 1, 0.12)], "ok"),
    ("format_only", s_format_only, [("P&L", "Выручка", "B", 1, 1_250_000)], "ok"),
    ("edit_tax_rate", s_edit_tax, [("P&L", "Ставка налога, %", "B", 1, 0.25)], "ok"),
    ("clear_formula", s_clear_formula, [("P&L", "Выручка", "B", 1, 1_250_000)], "ok"),
    ("formula_with_money_constant", s_formula_const, [], "flag"),
    ("hidden_sheet_edit", s_hidden_edit, [("Служебный", "Резерв, руб", "B", 1, 1_554_666)], "ok"),
    ("duplicate_labels", s_dup_labels, [("P&L", "Прочее", "B", 1, 111_111), ("P&L", "Прочее", "B", 2, 444_444)], "ok"),
    ("double_all_inputs", s_double_all, [("P&L", "Выручка", "B", 1, 2_500_000), ("P&L", "Расходы на маркетинг", "D", 1, 280_000)], "ok"),
    ("insert_row_and_edit", s_insert_and_edit, [("P&L", "Возвраты", "B", 1, 50_000), ("P&L", "Выручка", "B", 1, 1_312_500),
                                                ("P&L", "Себестоимость", "B", 1, 700_000)], "ok"),
    ("pasted_table_same_label", s_pasted_table, [], "ok"),
    ("price_edit_other_sheet", s_price_edit, [("Цены", "Товар А", "B", 1, 1500.6)], "ok"),
    ("sheet_removed", s_remove_sheet, [("P&L", "Выручка", "B", 1, 1_250_000)], "ok"),
]


# -- вторая партия: правки, исход которых заранее неочевиден; контракт один: либо числа верны, либо файл помечен и не "ok" -----------

def s_text_number(wb):
    wb["P&L"]["B2"] = str(wb["P&L"]["B2"].value * 1.1)


def s_zero_out(wb):
    wb["P&L"]["B3"] = 0


def s_negate(wb):
    _mul(wb["P&L"], "B5", -1)


def s_insert_column(wb):
    wb["P&L"].insert_cols(2)


def s_total_row_formula(wb):
    ws = wb["P&L"]
    ws["A11"], ws["B11"] = "Итого расходы", "=B3+B5"


def s_cross_sheet_formula(wb):
    ws = wb.create_sheet("Assumptions")
    ws["A1"], ws["B1"] = "Выручка Q1 справочно", "='P&L'!B2"


def s_repoint_formula(wb):
    wb["P&L"]["B4"] = "=C2-C3"


def s_delete_hidden(wb):
    del wb["Служебный"]


def s_merge_cells(wb):
    wb["P&L"].merge_cells("G1:H1")


def s_comment(wb):
    from openpyxl.comments import Comment
    wb["P&L"]["B2"].comment = Comment("проверить", "AI")


def s_number_format(wb):
    wb["P&L"]["B2"].number_format = '0.0,,"M"'


def s_named_range(wb):
    from openpyxl.workbook.defined_name import DefinedName
    wb.defined_names["Revenue_Q1"] = DefinedName("Revenue_Q1", attr_text="'P&L'!$B$2")


def s_copy_sheet(wb):
    wb.copy_worksheet(wb["P&L"]).title = "P&L (копия)"


def s_add_sum_row(wb):
    ws = wb["P&L"]
    ws["A11"], ws["B11"] = "Итого затраты", "=SUM(B3:B3)+SUM(B5:B5)"


def s_label_changed(wb):
    wb["P&L"]["A2"] = "Выручка (скорр.)"


def s_two_sheets(wb):
    _mul(wb["P&L"], "B2", 1.1)
    _mul(wb["Цены"], "B3", 0.75)


def s_scale_by_formula(wb):
    wb["P&L"]["B2"] = "=B3*1.8"


def s_reorder_sheets(wb):
    wb.move_sheet("Цены", -1)


def s_append_total_value(wb):
    ws = wb["P&L"]
    ws["A11"], ws["B11"] = "Итого выручка", (ws["B2"].value + ws["C2"].value + ws["D2"].value)


def s_sum_as_value(wb):
    ws = wb["P&L"]
    ws["E2"] = ws["B2"].value + ws["C2"].value + ws["D2"].value


def s_swap_columns(wb):
    ws = wb["P&L"]
    for r in (2, 3, 5):
        ws[f"B{r}"].value, ws[f"C{r}"].value = ws[f"C{r}"].value, ws[f"B{r}"].value


def s_rounded_paste(wb):
    ws = wb["P&L"]
    ws["B2"] = round(ws["B2"].value, -3)


def s_new_sheet_percent_only(wb):
    ws = wb.create_sheet("Ratios")
    ws["A1"], ws["B1"] = "Маржа, %", 0.31
    ws["B1"].number_format = "0%"


SCENARIOS_2 = [
    ("text_number_in_input", s_text_number, [("P&L", "Выручка", "B", 1, 1_375_000)]),
    ("zero_out_cost", s_zero_out, [("P&L", "Себестоимость", "B", 1, 0), ("P&L", "Выручка", "B", 1, 1_250_000)]),
    ("negate_expense", s_negate, [("P&L", "Расходы на маркетинг", "B", 1, -120_000)]),
    ("insert_column", s_insert_column, [("P&L", "Выручка", "C", 1, 1_250_000), ("P&L", "Выручка", "D", 1, 1_400_000)]),
    ("total_row_formula", s_total_row_formula, [("P&L", "Итого расходы", "B", 1, "=B3+B5"), ("P&L", "Выручка", "B", 1, 1_250_000)]),
    ("cross_sheet_formula", s_cross_sheet_formula, [("Assumptions", "Выручка Q1 справочно", "B", 1, "='P&L'!B2"), ("P&L", "Выручка", "B", 1, 1_250_000)]),
    ("repoint_formula", s_repoint_formula, [("P&L", "Валовая прибыль", "B", 1, "=C2-C3")]),
    ("delete_hidden_sheet", s_delete_hidden, [("P&L", "Выручка", "B", 1, 1_250_000), ("P&L", "Себестоимость", "B", 1, 700_000)]),
    ("merge_cells", s_merge_cells, [("P&L", "Выручка", "B", 1, 1_250_000)]),
    ("add_comment", s_comment, [("P&L", "Выручка", "B", 1, 1_250_000)]),
    ("change_number_format", s_number_format, [("P&L", "Выручка", "B", 1, 1_250_000)]),
    ("add_named_range", s_named_range, [("P&L", "Выручка", "B", 1, 1_250_000)]),
    ("copy_sheet", s_copy_sheet, [("P&L", "Выручка", "B", 1, 1_250_000), ("P&L (копия)", "Выручка", "B", 1, 1_250_000)]),
    ("add_sum_row", s_add_sum_row, [("P&L", "Итого затраты", "B", 1, "=SUM(B3:B3)+SUM(B5:B5)")]),
    ("label_changed", s_label_changed, [("P&L", "Выручка (скорр.)", "B", 1, 1_250_000), ("P&L", "Себестоимость", "B", 1, 700_000)]),
    ("two_sheets_edit", s_two_sheets, [("P&L", "Выручка", "B", 1, 1_375_000), ("Цены", "Товар Б", "B", 1, 674.25)]),
    ("scale_input_by_formula", s_scale_by_formula, [("P&L", "Выручка", "B", 1, "=B3*1.8")]),
    ("reorder_sheets", s_reorder_sheets, [("P&L", "Выручка", "B", 1, 1_250_000), ("Цены", "Товар А", "B", 1, 1250.5)]),
    ("append_total_value", s_append_total_value, [("P&L", "Итого выручка", "B", 1, 4_175_000)]),
    ("sum_pasted_as_value", s_sum_as_value, [("P&L", "Выручка", "E", 1, 4_175_000)]),
    ("swap_columns", s_swap_columns, [("P&L", "Выручка", "B", 1, 1_400_000), ("P&L", "Выручка", "C", 1, 1_250_000)]),
    ("new_sheet_percent_only", s_new_sheet_percent_only, [("Ratios", "Маржа, %", "B", 1, 0.31)]),
]

NEAR = Decimal("0.5")      # единицы исходной валюты (для сумм до копеек достаточно 0.5 на 7 знаков)


def _close(a, b, tol):
    if isinstance(b, str):
        return a == b
    return a is not None and abs(Decimal(str(a)) - Decimal(str(b))) <= tol


class FinanceRestoreSuite(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.results: dict = {}

    def run_scenario(self, name, edit, expect, mode):
        env = Env()
        self.addCleanup(env.close)
        job = env.anonymize(("model.xlsx", base_book()), options={"finance_numbers": True})
        self.assertTrue(job.files[0].out_path, job.files[0].messages)
        masked = env.output_bytes(job)
        wb = load_workbook(io.BytesIO(masked))
        edit(wb)
        buf = io.BytesIO()
        wb.save(buf)
        back = env.restore(Upload("ai_result.xlsx", buf.getvalue()))
        outcome = back.files[0]
        self.assertTrue(outcome.out_path, outcome.messages)
        restored = load_workbook(env.output(back))
        wrong = []
        for sheet, label, col, nth, want in expect:
            try:
                got = at(restored, sheet, label, col, nth)
            except (KeyError, Exception) as exc:                 # noqa: BLE001
                got = f"missing:{exc}"
            tol = Decimal("0.01") if isinstance(want, str) or abs(Decimal(str(want))) < 5000 else NEAR
            if not _close(got, want, tol):
                wrong.append((sheet, label, col, want, got))
        flagged = any("Проверьте суммы" in m["text"] for m in outcome.messages)
        return wrong, flagged, outcome

    def test_ai_edit_contract(self):
        """Вторая партия: каждый неверный результат обязан быть помечен предупреждением и не иметь статус ok."""
        restored = flagged = 0
        report = []
        for name, edit, expect in SCENARIOS_2:
            with self.subTest(name):
                wrong, was_flagged, outcome = self.run_scenario(name, edit, expect, "safe")
                if wrong:
                    self.assertTrue(was_flagged and outcome.status != "ok", (name, wrong, outcome.status, [m["text"] for m in outcome.messages]))
                    flagged += 1
                    report.append((name, "flagged"))
                else:
                    restored += 1
                    report.append((name, "restored"))
        self.assertEqual(restored + flagged, len(SCENARIOS_2))
        self.assertGreaterEqual(len(SCENARIOS) + len(SCENARIOS_2), 40, report)
        self.assertGreaterEqual(restored, 12, report)

    def test_scenarios(self):
        correct = flagged_ok = 0
        report = []
        for name, edit, expect, mode in SCENARIOS:
            with self.subTest(name):
                wrong, flagged, outcome = self.run_scenario(name, edit, expect, mode)
                if mode == "flag":
                    self.assertTrue(flagged and outcome.status != "ok", (name, outcome.status, outcome.messages))
                    flagged_ok += 1
                    report.append((name, "flagged"))
                    continue
                self.assertFalse(wrong, (name, wrong, outcome.messages))
                correct += 1
                report.append((name, "restored" + (" +flags" if flagged else "")))
        total = len(SCENARIOS)
        self.assertGreaterEqual(total, 25)
        self.assertGreaterEqual(correct, 22, report)
        self.assertEqual(correct + flagged_ok, total, report)

    def test_unchanged_cells_identical(self):
        """Ячейки, которых модель не касалась, возвращаются тем же текстом числа, что был в оригинале."""
        import zipfile
        env = Env()
        self.addCleanup(env.close)
        orig = base_book()
        job = env.anonymize(("model.xlsx", orig), options={"finance_numbers": True})
        masked = env.output_bytes(job)
        wb = load_workbook(io.BytesIO(masked))
        _mul(wb["P&L"], "B2", 1.1)
        buf = io.BytesIO()
        wb.save(buf)
        back = env.restore(Upload("r.xlsx", buf.getvalue()))
        a = load_workbook(io.BytesIO(orig))
        b = load_workbook(env.output(back))
        for ws in a.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if (ws.title, cell.coordinate) == ("P&L", "B2"):
                        continue
                    self.assertEqual(b[ws.title][cell.coordinate].value, cell.value, (ws.title, cell.coordinate))


if __name__ == "__main__":
    unittest.main()
