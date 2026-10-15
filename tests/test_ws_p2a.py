"""P2A: движок правки Office-файлов (anonymizer/docedit*.py). Фикстуры строятся библиотеками, личных данных нет."""
from __future__ import annotations

import copy
import io
import random
import unittest
import zipfile

import docx
import openpyxl
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from anonymizer import docedit
from anonymizer.docedit import Opts, PatchRejected, apply_patch, diff_summary, outline

def _png() -> bytes:
    from PIL import Image
    out = io.BytesIO()
    Image.new("RGB", (4, 4), (200, 30, 30)).save(out, "PNG")
    return out.getvalue()


PNG = _png()


def make_pptx(chart: bool = True) -> bytes:
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Стратегия Company2 на 2026"
    tf = s1.placeholders[1].text_frame
    tf.text = "Руководитель проекта: Name1"
    p = tf.add_paragraph()
    r1 = p.add_run()
    r1.text = "Бюджет "
    r1.font.bold = True
    r1.font.size = Pt(20)
    r2 = p.add_run()
    r2.text = "согласован с Company4"
    tf.add_paragraph().text = "Третий абзац"
    tb = s1.shapes.add_textbox(Inches(1), Inches(5), Inches(3), Inches(1))
    tb.text_frame.text = "Подпись Name1"
    tb.text_frame.paragraphs[0].runs[0].font.name = "Arial"
    table = s1.shapes.add_table(2, 3, Inches(0.5), Inches(6), Inches(6), Inches(1)).table
    table.cell(0, 0).text = "Статья"
    table.cell(0, 1).text = "Сумма Company2"
    s1.notes_slide.notes_text_frame.text = "Заметка про Name1"
    s1.shapes.add_picture(io.BytesIO(PNG), Inches(7), Inches(5), Inches(1), Inches(1))
    if chart:
        cd = CategoryChartData()
        cd.categories = ["A", "B"]
        cd.add_series("S", (1, 2))
        s1.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1), Inches(3), Inches(2), cd)
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Итоги"
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def make_docx() -> bytes:
    d = docx.Document()
    d.add_heading("Договор с Company2", 1)
    p = d.add_paragraph()
    r = p.add_run("Исполнитель ")
    r.bold = True
    p.add_run("Name1 обязуется выполнить работы.")
    d.add_paragraph("Второй абзац без меток.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Пункт"
    t.cell(0, 1).text = "Сумма Company2"
    d.sections[0].header.paragraphs[0].text = "Шапка Company2"
    d.sections[0].footer.paragraphs[0].text = "Подвал"
    pic = d.add_paragraph()
    pic.add_run().add_picture(io.BytesIO(PNG))
    pic.add_run(" подпись к рисунку")
    fld = d.add_paragraph("Страница ")
    el = etree.SubElement(fld._p, docx.oxml.ns.qn("w:fldSimple"))
    el.set(docx.oxml.ns.qn("w:instr"), "PAGE")
    rev = d.add_paragraph()
    ins = etree.SubElement(rev._p, docx.oxml.ns.qn("w:ins"))
    ins.set(docx.oxml.ns.qn("w:id"), "1")
    ins.set(docx.oxml.ns.qn("w:author"), "x")
    rr = etree.SubElement(ins, docx.oxml.ns.qn("w:r"))
    etree.SubElement(rr, docx.oxml.ns.qn("w:t")).text = "правка рецензента"
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def make_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Бюджет"
    ws["A1"] = "Статья Company2"
    ws["A2"] = "Аренда"
    ws["B2"] = 1234.5
    ws["B3"] = 10
    ws["B4"] = "=SUM(B2:B3)"
    ws["A5"] = "Итого Name1"
    ws["B2"].font = openpyxl.styles.Font(bold=True, name="Calibri", size=12)
    ws2 = wb.create_sheet("Второй")
    ws2["A1"] = "Лист два"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def els(o):
    return [e for c in o["containers"] for e in c["elements"]]


def by_text(o, start):
    return next(e for e in els(o) if (e.get("text") or "").startswith(start))


def zip_parts(data):
    z = zipfile.ZipFile(io.BytesIO(data))
    return {i.filename: z.read(i) for i in z.infolist()}


def reject(tc, data, name, ops, code=None, opts=None):
    with tc.assertRaises(PatchRejected) as cm:
        apply_patch(data, name, {"ops": ops}, opts)
    if code:
        tc.assertEqual(cm.exception.code, code)
    return cm.exception


class PptxTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = make_pptx()
        cls.o = outline(cls.data, "d.pptx")

    def test_addresses_stable(self):
        a, b = outline(self.data, "d.pptx"), outline(self.data, "d.pptx")
        self.assertEqual(a, b)
        addrs = [e["addr"] for e in els(a)]
        self.assertEqual(len(addrs), len(set(addrs)))
        self.assertTrue(any(x.endswith(".notes.p0") for x in addrs))
        self.assertTrue(any(".r0.c1.p0" in x for x in addrs))
        # адрес фигуры = её id, а не позиция: удаление соседней фигуры не меняет адрес остальных
        victim = next(e for e in els(self.o) if e.get("text") == "Подпись Name1")
        target = by_text(self.o, "Сумма")
        shape = victim["addr"].rsplit(".", 1)[0]
        new, _ = apply_patch(self.data, "d.pptx", {"ops": [{"op": "delete", "at": shape}]})
        self.assertEqual(by_text(outline(new, "d.pptx"), "Сумма")["addr"], target["addr"])

    def test_non_editable_flagged_and_rejected(self):
        shapes = [e for e in els(self.o) if e["kind"] == "shape" and not e["editable"]]
        types = {e["type"] for e in shapes}
        self.assertIn("image", types)
        self.assertIn("chart", types)
        for s in shapes:
            self.assertIn("не редактируется", s["label"])
            reject(self, self.data, "d.pptx", [{"op": "delete", "at": s["addr"]}], "not_editable")
            reject(self, self.data, "d.pptx", [{"op": "set_style", "at": s["addr"], "bold": True}], "not_editable")

    def test_replace_keeps_first_run_format(self):
        e = by_text(self.o, "Бюджет")
        new, rep = apply_patch(self.data, "d.pptx", {"ops": [{"op": "replace_text", "at": e["addr"], "expect": e["fp"],
                                                              "text": "Бюджет утверждён Company4"}]})
        self.assertEqual(by_text(outline(new, "d.pptx"), "Бюджет")["text"], "Бюджет утверждён Company4")
        prs = Presentation(io.BytesIO(new))
        para = [p for sh in prs.slides[0].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs if p.text.startswith("Бюджет")][0]
        self.assertEqual(len(para.runs), 1)
        self.assertTrue(para.runs[0].font.bold)
        self.assertEqual(para.runs[0].font.size, Pt(20))
        self.assertEqual(rep["changed_parts"], ["ppt/slides/slide1.xml"])

    def test_unrelated_parts_identical(self):
        e = by_text(self.o, "Итоги")
        new, rep = apply_patch(self.data, "d.pptx", {"ops": [{"op": "replace_text", "at": e["addr"], "expect": e["fp"], "text": "Итоги года"}]})
        a, b = zip_parts(self.data), zip_parts(new)
        self.assertEqual(set(a), set(b))
        self.assertEqual(rep["changed_parts"], ["ppt/slides/slide2.xml"])
        for k in a:
            if k != "ppt/slides/slide2.xml":
                self.assertEqual(a[k], b[k], k)

    def test_ops_happy(self):
        e = by_text(self.o, "Третий")
        body = by_text(self.o, "Руководитель")
        sig = by_text(self.o, "Подпись")
        sig_shape = sig["addr"].rsplit(".", 1)[0]
        ops = [
            {"op": "insert_paragraph", "after": e["addr"], "text": "Четвёртый Name1"},
            {"op": "delete", "at": e["addr"], "expect": e["fp"]},
            {"op": "comment", "at": body["addr"], "text": "Слишком длинно"},
            {"op": "set_style", "at": body["addr"], "size": 18, "bold": True, "italic": False, "color": "#112233", "alignment": "center",
             "space_after": 6, "line_spacing": 110, "font": "Arial"},
            {"op": "move_shape", "at": sig_shape, "x": 100, "y": 400},
            {"op": "resize_shape", "at": sig_shape, "w": 200, "h": 40},
            {"op": "align", "ats": [sig_shape], "mode": "right", "to": "slide"},
            {"op": "set_layout", "slide": "s2", "layout": "Title and Content"},
        ]
        new, rep = apply_patch(self.data, "d.pptx", {"ops": ops})
        self.assertEqual(len(rep["applied"]), len(ops))
        self.assertEqual(rep["comments"][0]["text"], "Слишком длинно")
        o2 = outline(new, "d.pptx")
        self.assertIn("Четвёртый Name1", [x.get("text") for x in els(o2)])
        self.assertNotIn("Третий абзац", [x.get("text") for x in els(o2)])
        st = by_text(o2, "Руководитель")["style"]
        self.assertEqual((st["size"], st["bold"], st["color"], st["align"], st["font"]), (18.0, True, "112233", "center", "Arial"))
        g = next(x for x in els(o2) if x["addr"] == sig_shape)["geom"]
        self.assertEqual(g["w"], 200.0)
        self.assertAlmostEqual(g["x"] + g["w"], 720.0, delta=0.1)
        self.assertEqual(o2["containers"][1]["layout"], "Title and Content")
        Presentation(io.BytesIO(new))

    def test_table_col_width_and_bounds(self):
        tbl = next(e for e in els(self.o) if e.get("type") == "table")
        new, _ = apply_patch(self.data, "d.pptx", {"ops": [{"op": "col_width", "at": tbl["addr"], "col": 0, "width": 100}]})
        t2 = next(e for e in els(outline(new, "d.pptx")) if e.get("type") == "table")
        self.assertEqual(t2["cols"][0], 100.0)
        reject(self, self.data, "d.pptx", [{"op": "col_width", "at": tbl["addr"], "col": 0, "width": 900}], "out_of_bounds")

    def test_atomic_rejections(self):
        e = by_text(self.o, "Итоги")
        t1 = by_text(self.o, "Стратегия")
        good = {"op": "replace_text", "at": t1["addr"], "expect": t1["fp"], "text": "Стратегия Company2 на 2027"}
        sig = by_text(self.o, "Подпись")["addr"].rsplit(".", 1)[0]
        cases = [
            ({"op": "replace_text", "at": "s9.sh1.p0", "expect": "00000000", "text": "x"}, "unknown_address"),
            ({"op": "replace_text", "at": e["addr"], "expect": "deadbeef", "text": "x"}, "stale"),
            ({"op": "replace_text", "at": e["addr"], "expect": e["fp"], "text": "Итоги Company77"}, "label_unknown"),
            ({"op": "move_shape", "at": sig, "x": 5000, "y": 10}, "out_of_bounds"),
            ({"op": "set_style", "at": e["addr"], "font": "Comic Sans MS"}, "font_not_allowed"),
            ({"op": "set_layout", "slide": "s1", "layout": "Нет такого"}, "unknown_layout"),
            ({"op": "explode", "at": e["addr"]}, "unknown_op"),
            ({"op": "replace_text", "at": e["addr"], "expect": e["fp"], "text": "x" * 5000}, "text_too_long"),
            ({"op": "replace_text", "at": e["addr"], "expect": e["fp"], "text": "ab\x00c"}, "bad_text"),
            ({"op": "replace_text", "at": e["addr"], "expect": e["fp"], "text": "q", "extra": 1}, "bad_op"),
        ]
        before = bytes(self.data)
        for bad, code in cases:
            reject(self, self.data, "d.pptx", [good, bad], code)     # хорошая операция не должна просочиться
        self.assertEqual(before, self.data)
        reject(self, self.data, "d.pptx", [good] * 3, "duplicate_address")

    def test_limits(self):
        e = by_text(self.o, "Итоги")
        ops = [{"op": "comment", "at": e["addr"], "text": "c"}] * 5
        reject(self, self.data, "d.pptx", ops, "too_many_ops", Opts(max_ops=4))
        grow = {"op": "replace_text", "at": e["addr"], "expect": e["fp"], "text": "Итоги " + "очень длинно " * 20}
        reject(self, self.data, "d.pptx", [grow], "growth")
        apply_patch(self.data, "d.pptx", {"ops": [grow]}, Opts(max_growth=100, growth_floor=1000))

    def test_labels_protected(self):
        e = by_text(self.o, "Руководитель")
        for text, code in (("Руководитель проекта", "label_removed"), ("Руководитель Name1 name1", "label_unknown"),
                           ("Руководитель Name 1", "label_loose"), ("Руководитель Name1 и Name99", "label_unknown")):
            ex = reject(self, self.data, "d.pptx", [{"op": "replace_text", "at": e["addr"], "expect": e["fp"], "text": text}])
            self.assertIn(ex.code, {code, "label_loose", "label_removed"}, text)
        # метка, уже существующая в файле, использовать можно; удаление метки допустимо только через delete абзаца
        new, _ = apply_patch(self.data, "d.pptx", {"ops": [{"op": "replace_text", "at": e["addr"], "expect": e["fp"],
                                                            "text": "Name1 и Company2 — руководители"}]})
        reject(self, self.data, "d.pptx", [{"op": "comment", "at": e["addr"], "text": "см. Name42"}], "label_unknown")
        victim = by_text(self.o, "Подпись")
        apply_patch(self.data, "d.pptx", {"ops": [{"op": "delete", "at": victim["addr"].rsplit(".", 1)[0]}]})

    def test_stub_mode(self):
        st = outline(self.data, "d.pptx", text="stub")
        for a, b in zip(els(self.o), els(st)):
            self.assertEqual(a["addr"], b["addr"])
            if "text" in a:
                self.assertEqual(len(a["text"]), len(b["text"]))
                self.assertNotIn("fp", b)
        t = by_text(st, "xxxxxxxxxxx")
        self.assertIn("Name1", t["text"])
        self.assertNotIn("Руководитель", repr(st))

    def test_budget(self):
        sel = by_text(self.o, "Бюджет")["addr"]
        b = outline(self.data, "d.pptx", selected=sel, neighbours=1)
        texts = [e.get("text") for e in els(b) if e["kind"] == "paragraph"]
        self.assertIn("Третий абзац", texts)
        self.assertNotIn("Итоги", texts)
        self.assertTrue(b["containers"][1]["summary_only"])

    def test_diff_summary(self):
        e = by_text(self.o, "Итоги")
        patch = {"ops": [{"op": "replace_text", "at": e["addr"], "expect": e["fp"], "text": "Итоги года"}]}
        new, _ = apply_patch(self.data, "d.pptx", patch)
        ds = diff_summary(self.data, new, patch, "d.pptx")
        self.assertIn("было «Итоги» -> стало «Итоги года»", ds["items"][0]["text"])
        self.assertEqual(ds["preview_hint"]["addresses"], [e["addr"]])
        self.assertEqual(ds["preview_hint"]["containers"], ["s2"])

    def test_property_formatting_only_keeps_text(self):
        rnd = random.Random(7)
        o = outline(self.data, "d.pptx")
        texts = [(e["addr"], e["text"]) for e in els(o) if "text" in e]
        paras = [a for a, _ in texts]
        shapes = [e["addr"] for e in els(o) if e["kind"] == "shape" and e.get("type") == "text" and e["addr"].startswith("s1.sh") and ".notes" not in e["addr"]]
        for _ in range(25):
            ops = []
            for _ in range(rnd.randint(1, 4)):
                k = rnd.choice(["style", "move", "resize", "align", "col"])
                if k == "style":
                    ops.append({"op": "set_style", "at": rnd.choice(paras + shapes), **rnd.choice(
                        [{"bold": True}, {"size": rnd.randint(8, 40)}, {"color": "AA00BB"}, {"alignment": "right"}, {"font": "Calibri"},
                         {"space_before": 3, "line_spacing": 120}])})
                elif k == "move":
                    ops.append({"op": "move_shape", "at": rnd.choice(shapes), "x": rnd.randint(0, 100), "y": rnd.randint(0, 100)})
                elif k == "resize":
                    ops.append({"op": "resize_shape", "at": rnd.choice(shapes), "w": rnd.randint(50, 200), "h": rnd.randint(20, 80)})
                elif k == "align":
                    ops.append({"op": "align", "ats": [rnd.choice(shapes)], "mode": rnd.choice(["left", "center", "top", "middle"])})
                else:
                    tbl = next(e for e in els(o) if e.get("type") == "table")["addr"]
                    ops.append({"op": "col_width", "at": tbl, "col": rnd.randint(0, 2), "width": rnd.randint(40, 150)})
            try:
                new, _ = apply_patch(self.data, "d.pptx", {"ops": ops})
            except PatchRejected:
                continue          # допустимо: геометрия вышла за слайд; важно, что принятые патчи не трогают текст
            self.assertEqual(texts, [(e["addr"], e["text"]) for e in els(outline(new, "d.pptx")) if "text" in e])

    def test_group_child_and_xml_order(self):
        """Порядок дочерних элементов rPr соблюдён (схема), иначе PowerPoint просит восстановление."""
        e = by_text(self.o, "Руководитель")
        new, _ = apply_patch(self.data, "d.pptx", {"ops": [{"op": "set_style", "at": e["addr"], "color": "FF0000", "font": "Calibri", "size": 12}]})
        root = etree.fromstring(zip_parts(new)["ppt/slides/slide1.xml"])
        ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        for rpr in root.iter(ns + "rPr"):
            tags = [c.tag.replace(ns, "") for c in rpr]
            if "solidFill" in tags and "latin" in tags:
                self.assertLess(tags.index("solidFill"), tags.index("latin"))


class DocxTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = make_docx()
        cls.o = outline(cls.data, "d.docx")

    def test_addresses_and_flags(self):
        self.assertEqual(self.o, outline(self.data, "d.docx"))
        addrs = {e["addr"] for e in els(self.o)}
        self.assertTrue({"p0", "p1", "t1", "t1.r0.c1.p0", "h1.p0", "f1.p0"} <= addrs)
        fld = by_text(self.o, "Страница")
        self.assertFalse(fld["editable"])
        rev = next(e for e in els(self.o) if e["kind"] == "paragraph" and not e["editable"] and e["addr"] != fld["addr"])
        self.assertIn("рецензирование", rev["flags"][1])
        pic = by_text(self.o, " подпись")
        self.assertIn("has_object", pic["flags"])
        for e in (fld, rev):
            reject(self, self.data, "d.docx", [{"op": "replace_text", "at": e["addr"], "expect": e["fp"], "text": "x"}], "not_editable")

    def test_replace_keeps_format_and_object(self):
        e = by_text(self.o, "Исполнитель")
        new, rep = apply_patch(self.data, "d.docx", {"ops": [{"op": "replace_text", "at": e["addr"], "expect": e["fp"],
                                                              "text": "Исполнитель Name1 выполнит работы."}]})
        d = docx.Document(io.BytesIO(new))
        p = d.paragraphs[1]
        self.assertEqual(p.text, "Исполнитель Name1 выполнит работы.")
        self.assertTrue(p.runs[0].bold)
        pic = by_text(self.o, " подпись")
        new2, _ = apply_patch(self.data, "d.docx", {"ops": [{"op": "replace_text", "at": pic["addr"], "expect": pic["fp"], "text": "Другая подпись"}]})
        self.assertGreater(len(list(docx.Document(io.BytesIO(new2)).inline_shapes)), 0)
        self.assertEqual(rep["changed_parts"], ["word/document.xml"])

    def test_ops_happy(self):
        p2 = by_text(self.o, "Второй")
        cell = by_text(self.o, "Пункт")
        hdr = by_text(self.o, "Шапка")
        ops = [
            {"op": "insert_paragraph", "after": p2["addr"], "text": "Вставленный Company2"},
            {"op": "replace_text", "at": hdr["addr"], "expect": hdr["fp"], "text": "Шапка Company2 v2"},
            {"op": "set_style", "at": p2["addr"], "size": 14, "bold": True, "color": "336699", "alignment": "justify", "space_before": 6, "line_spacing": 150},
            {"op": "set_style", "at": "t1.r0.c0", "italic": True},
            {"op": "paragraph_style", "at": cell["addr"], "style": "Heading 2"},
            {"op": "col_width", "at": "t1", "col": 0, "width": 100},
            {"op": "comment", "at": cell["addr"], "text": "проверить"},
        ]
        new, rep = apply_patch(self.data, "d.docx", {"ops": ops})
        d = docx.Document(io.BytesIO(new))
        self.assertEqual(d.paragraphs[3].text, "Вставленный Company2")
        self.assertEqual(d.sections[0].header.paragraphs[0].text, "Шапка Company2 v2")
        self.assertEqual(d.paragraphs[2].runs[0].font.size, Pt(14))
        self.assertEqual(d.tables[0].cell(0, 0).paragraphs[0].style.name, "Heading 2")
        self.assertEqual(rep["comments"][0]["text"], "проверить")
        self.assertEqual(set(rep["changed_parts"]), {"word/document.xml", "word/header1.xml"})
        a, b = zip_parts(self.data), zip_parts(new)
        for k in a:
            if k not in rep["changed_parts"]:
                self.assertEqual(a[k], b[k], k)

    def test_delete_rules(self):
        p2 = by_text(self.o, "Второй")
        new, _ = apply_patch(self.data, "d.docx", {"ops": [{"op": "delete", "at": p2["addr"], "expect": p2["fp"]}]})
        self.assertNotIn("Второй абзац без меток.", [p.text for p in docx.Document(io.BytesIO(new)).paragraphs])
        cell = by_text(self.o, "Пункт")
        reject(self, self.data, "d.docx", [{"op": "delete", "at": cell["addr"], "expect": cell["fp"]}], "last_paragraph")
        reject(self, self.data, "d.docx", [{"op": "delete", "at": p2["addr"], "expect": "00000000"}], "stale")

    def test_labels_and_atomic(self):
        e = by_text(self.o, "Исполнитель")
        good = {"op": "comment", "at": e["addr"], "text": "ok"}
        for text, code in (("Исполнитель обязуется", "label_removed"), ("Исполнитель Name1 Name7 обязуется", "label_unknown")):
            reject(self, self.data, "d.docx", [good, {"op": "replace_text", "at": e["addr"], "expect": e["fp"], "text": text}], code)

    def test_property_formatting_only(self):
        rnd = random.Random(3)
        texts = [(e["addr"], e["text"]) for e in els(self.o) if "text" in e]
        paras = [a for a, _ in texts]
        for _ in range(20):
            ops = [{"op": "set_style", "at": rnd.choice(paras), **rnd.choice([{"bold": True}, {"size": 11}, {"color": "00AA00"}, {"alignment": "center"},
                                                                              {"space_after": 4}, {"font": "Calibri"}])} for _ in range(rnd.randint(1, 3))]
            new, _ = apply_patch(self.data, "d.docx", {"ops": ops})
            self.assertEqual(texts, [(e["addr"], e["text"]) for e in els(outline(new, "d.docx")) if "text" in e])
            docx.Document(io.BytesIO(new))

    def test_font_allowed_only_from_doc(self):
        p2 = by_text(self.o, "Второй")
        reject(self, self.data, "d.docx", [{"op": "set_style", "at": p2["addr"], "font": "Wingdings 7"}], "font_not_allowed")
        apply_patch(self.data, "d.docx", {"ops": [{"op": "set_style", "at": p2["addr"], "font": "Wingdings 7"}]}, Opts(safe_fonts=("Wingdings 7",)))


class XlsxTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = make_xlsx()
        cls.o = outline(cls.data, "d.xlsx")

    def test_addresses_stable(self):
        self.assertEqual(self.o, outline(self.data, "d.xlsx"))
        addrs = {e["addr"] for e in els(self.o)}
        self.assertTrue({"Бюджет!A1", "Бюджет!B2", "Бюджет!B4", "Второй!A1"} <= addrs)
        self.assertEqual(next(e for e in els(self.o) if e["addr"] == "Бюджет!B4")["text"], "=SUM(B2:B3)")
        self.assertEqual(next(e for e in els(self.o) if e["addr"] == "Бюджет!B2")["style"].get("bold"), True)

    def test_replace_text_and_scope(self):
        e = next(x for x in els(self.o) if x["addr"] == "Бюджет!A1")
        new, rep = apply_patch(self.data, "d.xlsx", {"ops": [{"op": "replace_text", "at": e["addr"], "expect": e["fp"], "text": "Статья Company2 (2026)"}]})
        wb = openpyxl.load_workbook(io.BytesIO(new))
        self.assertEqual(wb["Бюджет"]["A1"].value, "Статья Company2 (2026)")
        self.assertEqual(wb["Бюджет"]["B2"].font.bold, True)
        num = next(x for x in els(self.o) if x["addr"] == "Бюджет!B3")
        reject(self, self.data, "d.xlsx", [{"op": "replace_text", "at": num["addr"], "expect": num["fp"], "text": "x"}], "not_editable")
        reject(self, self.data, "d.xlsx", [{"op": "replace_text", "at": e["addr"], "expect": e["fp"], "text": "Статья"}], "label_removed")
        reject(self, self.data, "d.xlsx", [{"op": "insert_paragraph", "after": e["addr"], "text": "x"}], "unknown_op")

    def test_formula_stays_string_never_evaluated(self):
        new, rep = apply_patch(self.data, "d.xlsx", {"ops": [{"op": "set_cell", "at": "Бюджет!C2", "formula": "=B2*2"},
                                                              {"op": "set_cell", "at": "Бюджет!C3", "value": 5}]})
        wb = openpyxl.load_workbook(io.BytesIO(new))
        self.assertEqual(wb["Бюджет"]["C2"].value, "=B2*2")
        self.assertIsInstance(wb["Бюджет"]["C2"].value, str)
        self.assertEqual(wb["Бюджет"]["C3"].value, 5)
        wbv = openpyxl.load_workbook(io.BytesIO(new), data_only=True)
        self.assertIsNone(wbv["Бюджет"]["C2"].value)       # значение не вычислено и не подставлено
        self.assertIn(b"fullCalcOnLoad", zip_parts(new)["xl/workbook.xml"])
        for bad in ("=HYPERLINK(\"http://x\",1)", "='[evil.xlsx]S'!A1", "=WEBSERVICE(\"u\")"):
            reject(self, self.data, "d.xlsx", [{"op": "set_cell", "at": "Бюджет!C5", "formula": bad}], "bad_formula")
        f = next(x for x in els(self.o) if x["addr"] == "Бюджет!B4")
        reject(self, self.data, "d.xlsx", [{"op": "set_cell", "at": f["addr"], "formula": "=1"}], "expect_missing")
        reject(self, self.data, "d.xlsx", [{"op": "set_cell", "at": "Бюджет!C5", "value": "=SUM(1)"}], "bad_formula")

    def test_set_cell_labels(self):
        a5 = next(x for x in els(self.o) if x["addr"] == "Бюджет!A5")
        reject(self, self.data, "d.xlsx", [{"op": "set_cell", "at": a5["addr"], "expect": a5["fp"], "value": "Итого"}], "label_removed")
        reject(self, self.data, "d.xlsx", [{"op": "set_cell", "at": "Бюджет!C5", "value": "Name55"}], "label_unknown")
        apply_patch(self.data, "d.xlsx", {"ops": [{"op": "set_cell", "at": "Бюджет!C5", "value": "см. Name1"}]})

    def test_number_format_and_col_width(self):
        ops = [{"op": "number_format", "at": "Бюджет!B2:B3", "format": "#,##0.00"},
               {"op": "number_format", "at": "Бюджет!B2", "format": "0.0 \"руб.\""},
               {"op": "col_width", "at": "Бюджет!A:A", "width": 30}]
        new, _ = apply_patch(self.data, "d.xlsx", {"ops": ops})
        ws = openpyxl.load_workbook(io.BytesIO(new))["Бюджет"]
        self.assertEqual(ws["B3"].number_format, "#,##0.00")
        self.assertEqual(ws["B2"].number_format, '0.0 "руб."')
        self.assertEqual(ws["B2"].font.bold, True)
        self.assertEqual(ws.column_dimensions["A"].width, 30)
        self.assertEqual(ws["B2"].value, 1234.5)
        reject(self, self.data, "d.xlsx", [{"op": "number_format", "at": "Бюджет!B2", "format": "x" * 500}], "bad_format")
        reject(self, self.data, "d.xlsx", [{"op": "col_width", "at": "Бюджет!A", "width": 999}], "bad_number")

    def test_unrelated_parts_identical_and_property(self):
        a = zip_parts(self.data)
        new, rep = apply_patch(self.data, "d.xlsx", {"ops": [{"op": "number_format", "at": "Бюджет!B3", "format": "0.00%"}]})
        b = zip_parts(new)
        for k in a:
            if k not in rep["changed_parts"]:
                self.assertEqual(a[k], b[k], k)
        t1 = [(e["addr"], e["text"]) for e in els(self.o)]
        self.assertEqual(t1, [(e["addr"], e["text"]) for e in els(outline(new, "d.xlsx"))])

    def test_comment_only_changes_nothing(self):
        new, rep = apply_patch(self.data, "d.xlsx", {"ops": [{"op": "comment", "at": "Бюджет!B2", "text": "проверьте"}]})
        self.assertEqual(rep["changed_parts"], [])
        self.assertEqual(zip_parts(new).keys(), zip_parts(self.data).keys())
        self.assertEqual(rep["comments"], [{"at": "Бюджет!B2", "text": "проверьте"}])

    def test_diff_summary_xlsx(self):
        patch = {"ops": [{"op": "set_cell", "at": "Бюджет!C5", "value": "см. Name1"}]}
        new, _ = apply_patch(self.data, "d.xlsx", patch)
        ds = diff_summary(self.data, new, patch, "d.xlsx")
        self.assertEqual(ds["items"][0]["op"], "set_cell")


class GeneralTests(unittest.TestCase):
    def test_bad_inputs(self):
        with self.assertRaises(PatchRejected):
            outline(b"not a zip", "x.pptx")
        with self.assertRaises(PatchRejected):
            outline(make_pptx(False), "x.txt")
        for bad in ("{", "[]", {"ops": []}, {"ops": [], "zzz": 1}, {"v": 2, "ops": [{}]}, {"ops": [{"op": 1}]}):
            with self.assertRaises(PatchRejected):
                apply_patch(make_pptx(False), "x.pptx", bad)

    def test_patch_as_json_string(self):
        data = make_pptx(False)
        e = by_text(outline(data, "x.pptx"), "Итоги")
        import json
        new, _ = apply_patch(data, "x.pptx", json.dumps({"ops": [{"op": "replace_text", "at": e["addr"], "expect": e["fp"], "text": "Итоги 2"}]}))
        self.assertTrue(new)


if __name__ == "__main__":
    unittest.main()
