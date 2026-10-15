"""Workstream A (PDF/OCR): QR и штрихкоды, повёрнутые сканы, нечитаемый текст страниц, XFA/действия/EXIF/XMP/Subj."""
from __future__ import annotations

import hashlib
import io
import os
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path

import fitz

from anonymizer import ocr
from tests import pdf_probe as P
from tests.helpers import Env

HERE = Path(__file__).resolve().parent
ARIAL = Path(P.FONT).exists()
OCR_READY = ARIAL and ocr.available()
QR_PAYLOAD = "https://pay.example.ru/check?fio=Иванов Иван Петрович&inn=7707083893"
_MAKER: list = []


def make_code(payload: str, kind: str = "qr") -> bytes | None:
    """PNG с QR/Code128 (CoreImage); None, если на машине нечем сделать."""
    if not shutil.which("swiftc"):
        return None
    if not _MAKER:
        src = HERE / "qr_maker.swift"
        out = Path(tempfile.gettempdir()) / f"anon-qr-maker-{hashlib.sha1(src.read_bytes()).hexdigest()[:8]}"
        if not out.exists():
            done = subprocess.run(["swiftc", "-O", str(src), "-o", str(out)], capture_output=True)
            if done.returncode != 0:
                return None
        _MAKER.append(out)
    with tempfile.TemporaryDirectory() as tmp:
        png = Path(tmp) / "c.png"
        subprocess.run([str(_MAKER[0]), payload, str(png), kind], check=True, capture_output=True)
        return png.read_bytes()


def decode_all(page_or_png) -> list[str]:
    if isinstance(page_or_png, fitz.Page):
        return [b.payload for b in ocr.detect_barcodes(page_or_png, scale=3.0)]
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "x.png"
        path.write_bytes(page_or_png)
        return [b.payload for b in ocr.detect_barcodes_png(path)]


def texts(f) -> str:
    return " ".join(m["text"] for m in f.messages)


CODES = ARIAL and ocr.barcodes_available() and make_code("probe") is not None


@unittest.skipUnless(CODES and OCR_READY, "нужны Arial, OCR и поиск штрихкодов")
class BarcodePii(unittest.TestCase):
    def setUp(self):
        self.env = Env(image_mode="redact_text")
        self.addCleanup(self.env.close)

    def test_helper_sees_generated_qr(self):
        got = decode_all(make_code(QR_PAYLOAD))
        self.assertEqual(got, [QR_PAYLOAD])

    def test_qr_in_pdf_page_is_removed_or_file_blocked(self):
        doc = fitz.open()
        page = doc.new_page()
        P.put(page, (72, 100), "Договор поставки без персональных данных")
        page.insert_image(fitz.Rect(72, 200, 272, 400), stream=make_code(QR_PAYLOAD))
        job = self.env.anonymize(("qr.pdf", doc.tobytes()))
        f = job.files[0]
        self.assertNotIn("картинки без текста не менялись", texts(f))
        if not f.out_path:
            self.assertTrue(f.messages)
            return
        out = fitz.open(stream=Path(f.out_path).read_bytes(), filetype="pdf")
        for payload in decode_all(out[0]):
            self.assertNotIn("7707083893", payload)
            self.assertNotIn("Иванов", payload)
        self.assertIn("QR", texts(f) + " штрихкод")         # пользователь предупреждён
        self.assertNotEqual(f.status, "ok", texts(f))

    def test_qr_standalone_png(self):
        job = self.env.anonymize(("qr.png", make_code(QR_PAYLOAD)))
        f = job.files[0]
        if not f.out_path:
            self.assertTrue(f.messages)
            return
        for payload in decode_all(Path(f.out_path).read_bytes()):
            self.assertNotIn("7707083893", payload)

    def test_code128_with_inn_in_png(self):
        job = self.env.anonymize(("bar.png", make_code("7707083893", "code128")))
        f = job.files[0]
        if not f.out_path:
            self.assertTrue(f.messages)
            return
        self.assertNotIn("7707083893", decode_all(Path(f.out_path).read_bytes()))

    def test_qr_in_docx_image(self):
        from docx import Document
        d = Document()
        d.add_paragraph("Просто текст")
        d.add_picture(io.BytesIO(make_code(QR_PAYLOAD)))
        buf = io.BytesIO()
        d.save(buf)
        job = self.env.anonymize(("qr.docx", buf.getvalue()))
        f = job.files[0]
        if not f.out_path:
            self.assertTrue(f.messages)
            return
        import zipfile
        with zipfile.ZipFile(f.out_path) as z:
            for name in z.namelist():
                if name.startswith("word/media/"):
                    for payload in decode_all(z.read(name)):
                        self.assertNotIn("7707083893", payload)


def readable_in_any_orientation(data: bytes) -> str:
    """Весь текст, который читается на страницах результата при любом повороте картинки."""
    engine = ocr.provider()
    doc = fitz.open(stream=data, filetype="pdf")
    words = []
    for page in doc:
        for angle in (0, 90, 180, 270):
            m = fitz.Matrix(3, 3).prerotate(angle)
            pix = page.get_pixmap(matrix=m, alpha=False)
            with tempfile.TemporaryDirectory() as tmp:
                png = Path(tmp) / "p.png"
                pix.save(str(png))
                words.extend(w.text for w in engine.words(png))
    return " ".join(words)


SECRETS = ("Иванов", "Петрович", P.INN, "ivan.petrov")


@unittest.skipUnless(OCR_READY, "нужны Arial и движок распознавания")
class RotatedScans(unittest.TestCase):
    def setUp(self):
        self.env = Env(image_mode="redact_text")
        self.addCleanup(self.env.close)

    def scan(self, angle: int, mode: str) -> fitz.Document:
        src = fitz.open()
        page = src.new_page()
        P.put(page, (72, 100), P.LINE)
        P.put(page, (72, 130), "Телефон: " + P.PHONE)
        pix = page.get_pixmap(dpi=200)
        out = fitz.open()
        swap = mode == "image" and angle in (90, 270)
        w, h = (page.rect.height, page.rect.width) if swap else (page.rect.width, page.rect.height)
        target = out.new_page(width=w, height=h)
        if mode == "image":
            target.insert_image(target.rect, pixmap=pix, rotate=angle)
        else:
            target.insert_image(target.rect, pixmap=pix)
            target.set_rotation(angle)
        return out

    def check(self, angle: int, mode: str):
        job = self.env.anonymize(("scan.pdf", self.scan(angle, mode).tobytes()))
        f = job.files[0]
        if not f.out_path:                      # отказ выдать файл безопасен
            self.assertTrue(f.messages)
            return
        seen = readable_in_any_orientation(Path(f.out_path).read_bytes())
        for secret in SECRETS:
            self.assertNotIn(secret, seen, f"{mode} {angle}")

    def test_image_rotated_90(self):
        self.check(90, "image")

    def test_image_rotated_180(self):
        self.check(180, "image")

    def test_image_rotated_270(self):
        self.check(270, "image")

    def test_page_rotation_90(self):
        self.check(90, "page")

    def test_page_rotation_270(self):
        self.check(270, "page")

    def test_orientation_probe_path_for_engines_without_native_rotation(self):
        """Tesseract и Windows OCR читают только прямой текст: поворот ищется пробой четырёх поворотов."""
        ocr.FORCE_PROBE = True
        self.addCleanup(setattr, ocr, "FORCE_PROBE", False)
        page = self.scan(90, "image")[0]
        words, _ = ocr.ocr_page_words(page)
        joined = " ".join(w.text for w in words)
        self.assertIn("7707083893", joined)
        rotated = [w for w in words if w.text.startswith("7707083893")]
        self.assertTrue(rotated and rotated[0].rot in (90, 270))
        self.check(90, "image")

    def test_upright_scan_still_works(self):
        self.check(0, "image")


def outlined(text: str, size=14) -> fitz.Document:
    d = fitz.open()
    p = d.new_page()
    P.put(p, (72, 100), text, size=size)
    src = fitz.open("svg", p.get_svg_image(text_as_path=True).encode())
    return fitz.open("pdf", src.convert_to_pdf())


@unittest.skipUnless(OCR_READY, "нужны Arial и движок распознавания")
class UnreadablePageText(unittest.TestCase):
    """Страница показывает данные, а текстового слоя нет или он нечитаем: файл не выдаётся «чистым»."""

    def setUp(self):
        self.env = Env(image_mode="redact_text")
        self.addCleanup(self.env.close)

    def assert_hidden(self, doc: fitz.Document, secrets=("Иванов", "7707083893")):
        job = self.env.anonymize(("a.pdf", doc.tobytes()))
        f = job.files[0]
        if not f.out_path:
            self.assertTrue(f.messages)
            return f
        seen = readable_in_any_orientation(Path(f.out_path).read_bytes())
        for secret in secrets:
            self.assertNotIn(secret, seen)
        return f

    def test_font_without_tounicode(self):
        d = fitz.open()
        p = d.new_page()
        P.put(p, (72, 100), "Иванов Иван Петрович ИНН 7707083893")
        for font in p.get_fonts(full=True):
            d.xref_set_key(font[0], "ToUnicode", "null")
        broken = fitz.open("pdf", d.tobytes())
        self.assertNotIn("Иванов", broken[0].get_text())
        self.assertTrue(self.assert_hidden(broken).out_path)

    def test_few_outlined_glyphs(self):
        doc = outlined("Иванов Иван ИНН 7707083893")
        self.assertEqual(doc[0].get_text().strip(), "")
        self.assertTrue(self.assert_hidden(doc).out_path)

    def test_no_engine_for_unreadable_page_fails_closed(self):
        d = fitz.open()
        p = d.new_page()
        P.put(p, (72, 100), "Иванов Иван Петрович ИНН 7707083893")
        for font in p.get_fonts(full=True):
            d.xref_set_key(font[0], "ToUnicode", "null")
        broken = fitz.open("pdf", d.tobytes())
        saved = (ocr._PROVIDER, ocr._PROVIDER_SET)
        ocr._PROVIDER, ocr._PROVIDER_SET = None, True
        self.addCleanup(lambda: setattr(ocr, "_PROVIDER", saved[0]) or setattr(ocr, "_PROVIDER_SET", saved[1]))
        job = self.env.anonymize(("a.pdf", broken.tobytes()))
        self.assertFalse(job.files[0].out_path)


if __name__ == "__main__":
    unittest.main()
