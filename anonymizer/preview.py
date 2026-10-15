"""Просмотр файла чата для боковой панели (S11A).

Результат — только данные (JSON): страница строит DOM через textContent, поэтому содержимое файла никогда не
становится разметкой и не выполняется. HTML/SVG показываются исходным текстом. Сети нет, файл не покидает компьютер.
Всё, что не удалось разобрать или слишком велико, даёт {"kind": "none", "reason": ...} (закрыто по умолчанию).
"""

from __future__ import annotations

import base64
import io
import os
import re

from .tokens import TOKEN_RE

MAX_BYTES = 8 * 1024 * 1024        # больше — «слишком большой для просмотра»
MAX_LINES = 1500
MAX_LINE_CHARS = 2000
MAX_ROWS = 100
MAX_COLS = 24
MAX_SHEETS = 12
MAX_BLOCKS = 400
MAX_SLIDES = 60
MAX_IMAGE_BYTES = 6 * 1024 * 1024
TEXT_EXT = {".txt", ".md", ".markdown", ".log", ".json", ".xml", ".yaml", ".yml", ".ini", ".cfg", ".rtf",
            ".html", ".htm", ".svg", ".tex", ".rst", ".sql"}
TABLE_EXT = {".csv", ".tsv"}
IMAGE_MIME = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".gif": "image/gif",
              ".webp": "image/webp", ".bmp": "image/bmp"}
INVISIBLE = dict.fromkeys(map(ord, "​‌‍⁠﻿­"), None)


def _none(reason: str, **extra) -> dict:
    return {"kind": "none", "reason": reason, **extra}


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16") if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig", "cp1251"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, ValueError):
            continue
    return data.decode("utf-8", "replace")


def _cell(value, marks: bool):
    """Ячейка: строка или список сегментов [[текст, 0|1], ...], если в ней есть метки."""
    text = ("" if value is None else str(value)).translate(INVISIBLE)
    if len(text) > MAX_LINE_CHARS:
        text = text[:MAX_LINE_CHARS] + "…"
    return _segments(text) if marks else text


def _segments(text: str):
    found = list(TOKEN_RE.finditer(text))
    if not found:
        return text
    out, at = [], 0
    for m in found:
        if m.start() > at:
            out.append([text[at:m.start()], 0])
        out.append([m.group(0), 1])
        at = m.end()
    if at < len(text):
        out.append([text[at:], 0])
    return out


LANG = {".md": "markdown", ".markdown": "markdown", ".json": "json", ".xml": "xml", ".html": "xml", ".htm": "xml", ".svg": "xml",
        ".yaml": "yaml", ".yml": "yaml", ".ini": "ini", ".cfg": "ini", ".sql": "sql"}


def _text(data: bytes, marks: bool, ext: str = "") -> dict:
    lines = _decode(data).splitlines()
    total = len(lines)
    shown = [_cell(line, marks) for line in lines[:MAX_LINES]]
    res = {"kind": "text", "lines": shown, "total": total, "truncated": total > MAX_LINES}
    if ext in LANG:
        res["lang"] = LANG[ext]
    return res


def _table_rows(rows, marks: bool, total: int | None = None) -> dict:
    out, wide = [], 0
    for row in rows:
        row = list(row)
        wide = max(wide, len(row))
        out.append([_cell(c, marks) for c in row[:MAX_COLS]])
        if len(out) >= MAX_ROWS:
            break
    return {"kind": "table", "rows": out, "cols_cut": wide > MAX_COLS, "limit": MAX_ROWS,
            "total": total if total is not None else None}


def _csv(data: bytes, ext: str, marks: bool) -> dict:
    import csv
    text = _decode(data)
    rows, total = [], 0
    reader = csv.reader(io.StringIO(text), delimiter="\t" if ext == ".tsv" else (";" if text[:2000].count(";") > text[:2000].count(",") else ","))
    try:
        for row in reader:
            total += 1
            if len(rows) < MAX_ROWS:
                rows.append(row)
    except csv.Error:
        return _none("Не удалось разобрать таблицу")
    return _table_rows(rows, marks, total)


def _outline(data: bytes, name: str) -> list:
    """Контейнеры docedit (те же адреса, что получает модель). Любой сбой: адресов нет, просмотр работает как раньше."""
    try:
        from . import docedit
        return docedit.outline(data, name)["containers"]
    except Exception:
        return []


def _xlsx(data: bytes, marks: bool, name: str = "") -> dict:
    import openpyxl
    cont = _outline(data, name or "x.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets = []
        for ws in wb.worksheets[:MAX_SHEETS]:
            rows, total = [], 0
            for row in ws.iter_rows(min_row=1, min_col=1, values_only=True):
                total += 1
                if len(rows) < MAX_ROWS:
                    rows.append(row)
                if total >= 20000:
                    break
            one = _table_rows(rows, marks, total)
            one["sheet"] = ws.title
            known = {e["addr"] for c in cont if c["addr"] == ws.title for e in c["elements"] if e.get("editable", True)}
            if known:
                from openpyxl.utils import get_column_letter
                one["addrs"] = [[(f"{ws.title}!{get_column_letter(j + 1)}{i + 1}" if f"{ws.title}!{get_column_letter(j + 1)}{i + 1}" in known else "")
                                 for j in range(len(r))] for i, r in enumerate(one["rows"])]
            sheets.append(one)
        res = dict(sheets[0]) if sheets else _table_rows([], marks, 0)
        res["sheets"] = len(wb.worksheets)
        res["sheet_list"] = [{k: v for k, v in sh.items() if k != "kind"} for sh in sheets]
        return res
    finally:
        wb.close()


def _docx(data: bytes, marks: bool, name: str = "") -> dict:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    document = docx.Document(io.BytesIO(data))
    els = {e["addr"]: e for c in _outline(data, name or "x.docx") if c["addr"] == "body" for e in c["elements"]}
    paras = iter([a for a, e in els.items() if re.fullmatch(r"p\d+", a) and (e.get("text") or "").strip()])
    tabs = iter([a for a in els if re.fullmatch(r"t\d+", a)])
    blocks = []
    for child in document.element.body.iterchildren():
        if len(blocks) >= MAX_BLOCKS:
            break
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text
            if not text.strip():
                continue
            style = (para.style.name if para.style is not None else "") or ""
            level = 0
            if style.lower().startswith(("heading", "заголовок")) or style.lower() == "title":
                digits = "".join(ch for ch in style if ch.isdigit())
                level = min(int(digits), 3) if digits else 1
            block = {"t": "h" if level else "p", "level": level, "text": _cell(text, marks)}
            addr = next(paras, None)
            if addr and els[addr].get("editable", True):
                block["addr"] = addr
            blocks.append(block)
        elif tag == "tbl":
            table = Table(child, document)
            rows, addrs = [], []
            tname = next(tabs, None)
            for r, row in enumerate(table.rows[:MAX_ROWS]):
                rows.append([_cell(c.text, marks) for c in row.cells[:MAX_COLS]])
                addrs.append([(f"{tname}.r{r}.c{k}.p0" if tname and f"{tname}.r{r}.c{k}.p0" in els and els[f"{tname}.r{r}.c{k}.p0"].get("editable", True) else "")
                              for k in range(len(rows[-1]))])
            block = {"t": "table", "rows": rows}
            if any(a for row in addrs for a in row):
                block["addrs"] = addrs
            blocks.append(block)
    return {"kind": "doc", "blocks": blocks, "truncated": len(blocks) >= MAX_BLOCKS}


def _pptx(data: bytes, marks: bool, name: str = "") -> dict:
    import pptx
    deck = pptx.Presentation(io.BytesIO(data))
    cont = {c["addr"]: c for c in _outline(data, name or "x.pptx")}
    slides = []
    for number, slide in enumerate(deck.slides, 1):
        if number > MAX_SLIDES:
            break
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                texts.append(_cell(shape.text_frame.text, marks))
        one = {"n": number, "texts": texts}
        c = cont.get(f"s{number}")
        if c:
            one["addr"] = c["addr"]
            one["items"] = [{"addr": e["addr"], "role": e.get("role") or "", "text": _cell(e.get("text") or "", marks)}
                            for e in c["elements"] if e.get("kind") == "paragraph" and (e.get("text") or "").strip() and e.get("editable", True)]
        slides.append(one)
    return {"kind": "slides", "slides": slides, "total": len(deck.slides)}


def _pdf(data: bytes, page: int) -> dict:
    import fitz
    document = fitz.open(stream=data, filetype="pdf")
    try:
        count = document.page_count
        if count == 0:
            return _none("В PDF нет страниц")
        page = max(1, min(page, count))
        pixmap = document[page - 1].get_pixmap(matrix=fitz.Matrix(1.4, 1.4), alpha=False)
        png = pixmap.tobytes("png")
        if len(png) > MAX_IMAGE_BYTES:
            return _none("Страница слишком большая для просмотра")
        return {"kind": "pdf", "page": page, "pages": count,
                "image": "data:image/png;base64," + base64.b64encode(png).decode("ascii")}
    finally:
        document.close()


def render(name: str, data: bytes, page: int = 1, marks: bool = True) -> dict:
    """Содержимое для панели. `marks` — подсвечивать метки (только для обезличенных версий)."""
    if len(data) > MAX_BYTES:
        return _none("Файл слишком большой для просмотра", size=len(data))
    ext = os.path.splitext(name or "")[1].lower()
    try:
        if ext in IMAGE_MIME:
            if len(data) > MAX_IMAGE_BYTES:
                return _none("Файл слишком большой для просмотра", size=len(data))
            return {"kind": "image", "image": f"data:{IMAGE_MIME[ext]};base64," + base64.b64encode(data).decode("ascii")}
        if ext == ".pdf":
            return _pdf(data, page)
        if ext == ".docx":
            return _docx(data, marks, name)
        if ext == ".pptx":
            return _pptx(data, marks, name)
        if ext in (".xlsx", ".xlsm"):
            return _xlsx(data, marks, name)
        if ext in TABLE_EXT:
            return _csv(data, ext, marks)
        if ext in TEXT_EXT or not ext:
            return _text(data, marks, ext)
    except Exception:   # битый или необычный файл: просмотра нет, остальное работает
        return _none("Этот файл не удалось показать")
    return _none("Просмотр этого формата недоступен")
