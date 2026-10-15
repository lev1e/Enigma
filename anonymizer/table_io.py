"""Таблица замен чата как файл: выгрузка в Excel/CSV, загрузка правок и «вернуть заново».

Выгрузка содержит настоящие значения, поэтому живёт только во временной папке и отдаётся человеку. Загрузка сперва
строит план (ничего не меняя), потом по метке плана применяет его одним шагом, который отменяется целиком.
Показываемые человеку строки плана содержат только коды меток, без значений.
"""

from __future__ import annotations

import base64
import binascii
import csv
import hashlib
import io
import json
import re
import threading
import time

from .tokens import EMAIL_TOKEN_RE, KIND_BY_CODE, TOKEN_RE, format_token, parse_match, token_base

MAX_TABLE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20000
PLAN_TTL = 600
MAX_PLANS = 8
HEADERS = ("Код", "Вид", "Значение", "Скрывать", "Объединить с")
EXPIRED = "Просмотр устарел. Загрузите таблицу заново."
NO_RETURNED = "Файл, который вернул ИИ, в чате не сохранён: добавьте его снова."
_LOCK = threading.Lock()


# -- выгрузка -----------------------------------------------------------------------------

def export_rows(app, chat_id: str) -> list[list[str]]:
    """Строки без заголовка: по одной на написание метки. Слитые метки идут своими строками с целью слияния."""
    vault = app.vault
    out: list[list[str]] = []
    for row in app.chats.table(chat_id, vault):
        label = row["label"]
        for v in row["variants"]:
            out.append([v["token"], label, v["value"], "нет" if v["kept"] else "да", ""])
        for source in row["merged"]:
            entity = vault.entities.get(source)
            if not entity or entity.get("secret"):
                continue
            kind = KIND_BY_CODE[entity["kind"]]
            number = int(re.search(r"\d+", source).group())
            kept = app.chats.keep_set(chat_id, vault.prefs)
            for i, text in enumerate(entity["spellings"], start=1):
                edited = entity.get("edits", {}).get(str(i), text)
                out.append([format_token(kind, number, i), label, edited, "нет" if text.lower() in kept else "да", row["base"]])
    return out


def export_file(app, chat_id: str, fmt: str) -> tuple[bytes, str]:
    if app.chats.get(chat_id) is None:
        raise KeyError("Чата нет")
    if fmt not in ("xlsx", "csv"):
        raise ValueError("Неизвестный формат таблицы")
    rows = export_rows(app, chat_id)
    if fmt == "csv":
        buffer = io.StringIO(newline="")
        writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
        writer.writerow(HEADERS)
        writer.writerows(rows)
        return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8"), "замены.csv"
    from openpyxl import Workbook
    from openpyxl.styles import Font
    book = Workbook()
    sheet = book.active
    sheet.title = "Замены"
    sheet.append(list(HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for values in rows:
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"          # значение, похожее на формулу, остаётся текстом
    for column, width in zip("ABCDE", (16, 26, 44, 12, 16)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue(), "замены.xlsx"


# -- разбор файла -------------------------------------------------------------------------

def _norm(text) -> str:
    return re.sub(r"\s+", "", str(text or "")).casefold().replace("ё", "е")


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _read_xlsx(data: bytes) -> list[tuple[int, list[str]]]:
    from openpyxl import load_workbook
    try:
        book = load_workbook(io.BytesIO(data), data_only=True)
        sheet = book.worksheets[0]
        return [(n, [_cell(v) for v in values]) for n, values in enumerate(sheet.iter_rows(values_only=True), start=1)]
    except Exception:
        raise ValueError("Не удалось открыть файл. Загрузите таблицу .xlsx или .csv, которую скачали здесь.") from None


def _read_csv(data: bytes) -> list[tuple[int, list[str]]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = data.decode("cp1251")
        except UnicodeDecodeError:
            raise ValueError("Не удалось прочитать файл. Загрузите таблицу .xlsx или .csv, которую скачали здесь.") from None
    if "\x00" in text:
        raise ValueError("Не удалось прочитать файл. Загрузите таблицу .xlsx или .csv, которую скачали здесь.")
    first = next((line for line in text.splitlines() if line.strip()), "")
    delimiter = max((";", ",", "\t"), key=lambda d: first.count(d))
    rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
    return [(n, [c.strip() for c in values]) for n, values in enumerate(rows, start=1)]


def parse_table(name: str, data: bytes) -> list[dict]:
    """Строки данных: {row, code, value, hide, merge}; None в поле — столбца нет."""
    if not data:
        raise ValueError("Файл пустой.")
    if len(data) > MAX_TABLE_BYTES:
        raise ValueError("Файл больше 5 МБ. Загрузите таблицу, которую скачали здесь.")
    raw = _read_xlsx(data) if data[:2] == b"PK" or str(name).lower().endswith(".xlsx") else _read_csv(data)
    if len(raw) > MAX_ROWS:
        raise ValueError("В таблице слишком много строк.")
    header_at = next((i for i, (_, cells) in enumerate(raw) if any(cells)), None)
    if header_at is None:
        raise ValueError("Таблица пустая.")
    columns: dict[str, int] = {}
    for i, title in enumerate(raw[header_at][1]):
        columns.setdefault(_norm(title), i)
    code_at, value_at = columns.get("код"), columns.get("значение")
    if code_at is None or value_at is None:
        raise ValueError("Не нашли столбцы «Код» и «Значение». Скачайте таблицу из программы и правьте её.")
    hide_at, merge_at = columns.get("скрывать"), columns.get("объединитьс")
    rows = []
    for number, cells in raw[header_at + 1:]:
        if not any(cells):
            continue
        pick = lambda at: (cells[at] if at is not None and at < len(cells) else None)
        rows.append({"row": number, "code": pick(code_at) or "", "value": pick(value_at) or "",
                     "hide": pick(hide_at), "merge": pick(merge_at) or ""})
    return rows


def _parse_code(text: str):
    """(вид, номер, вариант) или None."""
    text = (text or "").strip()
    match = EMAIL_TOKEN_RE.fullmatch(text)
    if match:
        return parse_match(match, email=True)
    match = TOKEN_RE.fullmatch(text)
    return parse_match(match) if match else None


def _hide_flag(text):
    word = _norm(text)
    if word in ("да", "yes", "true", "1", "y"):
        return True
    if word in ("нет", "no", "false", "0", "n"):
        return False
    return None


# -- план ---------------------------------------------------------------------------------

def _plans(app) -> dict:
    if not hasattr(app, "table_plans"):
        app.table_plans = {}
    return app.table_plans


def build_plan(app, chat_id: str, rows: list[dict]) -> dict:
    vault = app.vault
    table = {r["base"]: r for r in app.chats.table(chat_id, vault)}
    plan = {"edits": [], "terms": [], "merges": [], "ai_changes": [], "restore_changes": [], "merges_text": [],
            "unchanged": 0, "unrecognized": []}
    seen_terms: dict[str, bool] = {}
    seen_merge: dict[str, str] = {}

    def bad(row, reason):
        plan["unrecognized"].append({"row": row["row"], "reason": reason})

    for row in rows:
        parsed = _parse_code(row["code"])
        if parsed is None:
            bad(row, "Не код метки" if row["code"] else "Нет кода")
            continue
        kind, number, variant = parsed
        source = token_base(kind, number)
        code = format_token(kind, number, variant)
        entity = vault.entities.get(source)
        if entity is None:
            bad(row, "Кода нет в хранилище")
            continue
        final = vault._final(source)
        if final not in table:
            bad(row, "Кода нет в этом чате")
            continue
        merge_to = None
        if row["merge"]:
            target = _parse_code(row["merge"])
            target_base = token_base(target[0], target[1]) if target else None
            if target_base is None or target_base not in vault.entities or vault._final(target_base) not in table:
                bad(row, "Нет такого кода для объединения в этом чате")
                continue
            merge_to = vault._final(target_base)
            if KIND_BY_CODE[vault.entities[merge_to]["kind"]] != KIND_BY_CODE[vault.entities[final]["kind"]]:
                bad(row, "Объединять можно только метки одного вида")
                continue
        if source != final:                                   # уже слитая метка: проверяется только цель слияния
            if merge_to is not None and merge_to != final:
                bad(row, "Код уже объединён с другим")
            else:
                plan["unchanged"] += 1
            continue
        current = vault.describe(final, app.chats.keep_set(chat_id, vault.prefs))
        if current["secret"]:
            bad(row, "Пароли и ключи не меняются")
            continue
        if not 1 <= variant <= len(current["variants"]):
            bad(row, "Нет такого написания")
            continue
        cur = current["variants"][variant - 1]
        value = row["value"].strip()
        if not value:
            bad(row, "Пустое значение")
            continue
        hide = _hide_flag(row["hide"]) if row["hide"] else None
        if row["hide"] and hide is None:
            bad(row, "Непонятно, скрывать или нет (нужно «да» или «нет»)")
            continue
        term_key = cur["original"].lower()
        if hide is not None and hide != (not cur["kept"]) and term_key in seen_terms and seen_terms[term_key] != hide:
            bad(row, "Противоречит другой строке")
            continue
        if merge_to is not None and merge_to != final and seen_merge.get(final, merge_to) != merge_to:
            bad(row, "Объединение противоречит другой строке")
            continue
        changed = False
        if hide is not None and hide != (not cur["kept"]) and term_key not in seen_terms:
            seen_terms[term_key] = hide
            plan["terms"].append(("hide_terms" if hide else "keep_terms", cur["original"], final))
            plan["ai_changes"].append(f"{code}: " + ("не скрывать -> скрывать" if hide else "скрывать -> не скрывать"))
            changed = True
        if value != cur["value"]:
            plan["edits"].append((final, variant, value, cur["original"]))
            plan["restore_changes"].append(f"{code}: значение изменится")
            changed = True
        if merge_to is not None and merge_to != final and final not in seen_merge:
            seen_merge[final] = merge_to
            plan["merges"].append((final, merge_to))
            plan["merges_text"].append(f"{final}: объединить с {merge_to}")
            changed = True
        if not changed:
            plan["unchanged"] += 1
    return plan


def plan_text(plan: dict) -> str:
    text = (f"Изменит файл для ИИ: {len(plan['ai_changes'])}. Изменит только возврат: {len(plan['restore_changes'])}.")
    if plan["merges"]:
        text += f" Объединит метки: {len(plan['merges'])}."
    bad = plan["unrecognized"]
    if bad:
        shown = ", ".join(str(item["row"]) for item in bad[:10]) + ("…" if len(bad) > 10 else "")
        text += f" Не распознано строк: {len(bad)} (номера: {shown})."
    return text


def preview(app, chat_id: str, name: str, data: bytes) -> dict:
    if app.chats.get(chat_id) is None:
        raise KeyError("Чата нет")
    plan = build_plan(app, chat_id, parse_table(name, data))
    token = base64.urlsafe_b64encode(hashlib.sha256(f"{time.time_ns()}{chat_id}".encode()).digest())[:16].decode()
    now = time.time()
    with _LOCK:
        plans = _plans(app)
        for key in [k for k, v in plans.items() if now - v["time"] > PLAN_TTL]:
            del plans[key]
        while len(plans) >= MAX_PLANS:
            del plans[min(plans, key=lambda k: plans[k]["time"])]
        plans[token] = {"time": now, "chat": chat_id, "plan": plan}
    return {"token": token, "ai_changes": plan["ai_changes"], "restore_changes": plan["restore_changes"],
            "merges": plan["merges_text"], "unchanged": plan["unchanged"], "unrecognized": plan["unrecognized"],
            "text": plan_text(plan)}


def apply_plan(app, chat_id: str, token: str) -> dict:
    if app.chats.get(chat_id) is None:
        raise KeyError("Чата нет")
    with _LOCK:
        entry = _plans(app).pop(str(token or ""), None)
    if entry is None or entry["chat"] != chat_id or time.time() - entry["time"] > PLAN_TTL:
        raise ValueError(EXPIRED)
    vault, plan = app.vault, entry["plan"]
    items: list[dict] = []
    done = {"ai": 0, "restore": 0, "merges": 0}
    failed = 0
    for list_name, term, base in plan["terms"]:
        try:
            info = app.chats.add_term(chat_id, list_name, term)     # решения таблицы действуют только в этом чате
        except ValueError:
            failed += 1
            continue
        items.append({"of": "keep" if list_name == "keep_terms" else "hide", "base": base, "term": term,
                      "was_other": info["was_other"], "was_listed": info["was_listed"]})
        done["ai"] += 1
    for base, variant, value, original in plan["edits"]:
        try:
            previous = vault.entities[vault._final(base)].get("edits", {}).get(str(variant))
            vault.edit_value(base, variant, value)
        except (ValueError, KeyError):
            failed += 1
            continue
        items.append({"of": "edit", "base": vault._final(base), "variant": variant, "previous": previous})
        done["restore"] += 1
    for source, target in plan["merges"]:
        try:
            vault.merge(source, target)
        except (ValueError, KeyError):
            failed += 1
            continue
        items.append({"of": "merge", "base": source})
        done["merges"] += 1
    ai_changed = bool(done["ai"] or done["merges"])
    if items:
        app.chats.add_step(chat_id, "edit", action="import", ai=done["ai"], restore=done["restore"],
                           merges=done["merges"], ai_changed=ai_changed)
    return {"applied": done, "skipped": failed, "ai_changed": ai_changed, "undo": {"of": "batch", "items": items},
            "table": app.chats.table(chat_id, vault)}


def import_table(app, chat_id: str, name: str, data: bytes, apply: bool = False, token: str | None = None) -> dict:
    if apply:
        return apply_plan(app, chat_id, token or "")
    return preview(app, chat_id, name, data)


def decode_upload(body: dict) -> bytes:
    try:
        data = base64.b64decode(str(body.get("data") or ""), validate=False)
    except (binascii.Error, ValueError):
        raise ValueError("Файл не удалось прочитать. Загрузите таблицу заново.") from None
    return data


# -- «вернуть заново» ---------------------------------------------------------------------

def _value_of(vault, base: str, variant: int):
    """Что вернулось бы вместо метки сейчас; то же, что `Vault.lookup`, но без отметок использования."""
    entity = vault.entities.get(base)
    for _ in range(50):
        if entity is None:
            return None
        if entity.get("merged_into"):
            mapped = entity.get("vmap", {}).get(str(variant))
            base, entity = entity["merged_into"], vault.entities.get(entity["merged_into"])
            if mapped is None:
                return None
            variant = int(mapped)
            continue
        if entity.get("secret") or not 1 <= variant <= len(entity["spellings"]):
            return None
        edited = entity.get("edits", {}).get(str(variant))
        return edited if edited is not None else entity["spellings"][variant - 1]
    return None


def signature(vault, bases) -> str:
    """Отпечаток значений, которые подставил бы возврат для этих меток. Значений в нём не прочитать."""
    parts = []
    for base in sorted(set(bases)):
        entity = vault.entities.get(base)
        count = len(entity["spellings"]) if entity else 0
        parts.append([base, [_value_of(vault, base, i) for i in range(1, count + 1)]])
    return hashlib.sha256(json.dumps(parts, ensure_ascii=False).encode("utf-8")).hexdigest()[:24]


def stale_restored(app, chat_id: str) -> list[dict]:
    chat = app.chats.get(chat_id)
    out = []
    for record in (chat or {"files": []})["files"]:
        if "restored" in record.get("roles", {}) and record.get("rsig") and record.get("rtokens") is not None \
                and signature(app.vault, record["rtokens"]) != record["rsig"]:
            out.append(record)
    return out


def restore_state(app, chat_id: str) -> dict:
    stale = stale_restored(app, chat_id)
    ready = [r for r in stale if "ret" in r.get("roles", {})]
    reason = ""
    if stale and not ready:
        reason = NO_RETURNED
    elif len(ready) < len(stale):
        reason = "Часть файлов, которые вернул ИИ, в чате не сохранена: они останутся как есть."
    return {"stale": len(stale), "redo": bool(ready), "reason": reason}
