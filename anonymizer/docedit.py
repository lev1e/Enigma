"""Движок правки pptx/docx/xlsx по патчу (P2A): outline, apply_patch, diff_summary.

Публичный API (подключается оркестратором, веб-слой сюда не входит):

    outline(data, name, opts=None, *, text=None, selected=None, neighbours=1) -> dict
    apply_patch(data, name, patch, opts=None) -> (new_bytes, report)      # PatchRejected при любой причине отказа
    diff_summary(old, new, patch, name, opts=None) -> {"items": [...], "preview_hint": {...}}

Работает с любым docx/pptx/xlsx, обезличен он или нет. Метки (Name1, Org2), если есть, защищены: не пропадают, не
искажаются, новые не придумываются. Патч применяется атомарно: либо весь, либо ничего. Формулы xlsx хранятся строкой
и никогда не вычисляются. Комментарии (op comment) возвращаются в report["comments"] и остаются внутри приложения.
"""

from __future__ import annotations

import io
import json
import os

from .docedit_core import Elem, Opts, Package, PatchRejected, check_labels, fingerprint, labels_of  # noqa: F401
from .docedit_pptx import PptxHandler

__all__ = ["Opts", "PatchRejected", "outline", "apply_patch", "diff_summary", "budget_outline", "fingerprint", "SUPPORTED_EXT"]

SUPPORTED_EXT = (".pptx", ".docx", ".xlsx")

BASE_SCHEMA = {
    "replace_text": ({"at", "expect", "text"}, set()),
    "insert_paragraph": ({"after", "text"}, set()),
    "delete": ({"at"}, {"expect"}),
    "comment": ({"at", "text"}, set()),
}
STYLE_KEYS = {"font", "size", "bold", "italic", "color", "alignment", "space_before", "space_after", "line_spacing"}
PPTX_SCHEMA = {
    **BASE_SCHEMA,
    "set_style": ({"at"}, STYLE_KEYS),
    "move_shape": ({"at", "x", "y"}, set()),
    "resize_shape": ({"at", "w", "h"}, set()),
    "align": ({"ats", "mode"}, {"to"}),
    "set_layout": ({"slide", "layout"}, set()),
    "col_width": ({"at", "col", "width"}, set()),
}
PptxHandler.SCHEMA = PPTX_SCHEMA


def _opts(opts) -> Opts:
    if opts is None:
        return Opts()
    if isinstance(opts, dict):
        return Opts(**opts)
    return opts


def _handler(pkg: Package, name: str, opts: Opts):
    ext = os.path.splitext(name.lower())[1]
    if ext == ".pptx":
        return PptxHandler(pkg, opts)
    if ext == ".docx":
        from .docedit_docx import DocxHandler
        return DocxHandler(pkg, opts)
    if ext == ".xlsx":
        from .docedit_xlsx import XlsxHandler
        return XlsxHandler(pkg, opts)
    raise PatchRejected("bad_format", "Поддерживаются только pptx, docx, xlsx.")


def _open(data: bytes, name: str, opts: Opts):
    pkg = Package(data)
    try:
        return pkg, _handler(pkg, name, opts)
    except PatchRejected:
        raise
    except Exception as exc:        # любая неожиданность разбора = файл не обрабатывается (закрыто по умолчанию)
        raise PatchRejected("bad_file", "Не удалось разобрать структуру файла.") from exc


# ---- outline ------------------------------------------------------------------------------------------------

def outline(data: bytes, name: str, opts=None, *, text: str | None = None, selected=None, neighbours: int = 1) -> dict:
    """Дерево элементов с устойчивыми адресами. text="stub": та же структура, тексты заменены заглушками (для правок оформления).
    selected: адрес или список адресов — вернуть бюджетный срез (выбранное + соседи), остальное свёрнуто."""
    o = _opts(opts)
    stub = (text or o.text) == "stub"
    pkg, h = _open(data, name, o)
    out = h.outline(stub)
    out["name"] = name
    out["text_mode"] = "stub" if stub else "full"
    out["stats"] = _stats(out)
    if out["stats"]["elements"] > o.max_elements and selected is None:
        raise PatchRejected("too_big", "Слишком много элементов: выберите часть документа.")
    if selected:
        out = budget_outline(out, selected, neighbours)
    return out


def _stats(out: dict) -> dict:
    n = ed = 0
    for c in out["containers"]:
        for e in c["elements"]:
            n += 1
            ed += 1 if e.get("editable") else 0
    return {"containers": len(out["containers"]), "elements": n, "editable": ed}


def budget_outline(out: dict, selected, neighbours: int = 1, max_elements: int = 300) -> dict:
    """Выбранные элементы + `neighbours` абзацев вокруг в том же контейнере; остальные контейнеры сводятся к заголовку."""
    sel = [selected] if isinstance(selected, str) else list(selected)
    res = {k: v for k, v in out.items() if k != "containers"}
    res["containers"], res["budget"] = [], {"selected": sel, "neighbours": neighbours}
    for c in out["containers"]:
        els = c["elements"]
        addrs = [e["addr"] for e in els]
        keep: set[int] = set()
        for s in sel:
            if s == c["addr"]:
                keep.update(range(len(els)))
            for idx, a in enumerate(addrs):
                if a == s or a.startswith(s + ".") or a.startswith(s + "/"):
                    keep.add(idx)
        if keep:
            para_idx = [i for i, e in enumerate(els) if e["kind"] in ("paragraph", "cell")]
            for i in sorted(keep):
                if els[i]["kind"] in ("paragraph", "cell") and i in para_idx:
                    pos = para_idx.index(i)
                    for j in para_idx[max(0, pos - neighbours): pos + neighbours + 1]:
                        keep.add(j)
            for i in list(keep):          # заголовок фигуры, к которой относятся абзацы
                if els[i]["kind"] in ("paragraph", "cell"):
                    for j in range(i, -1, -1):
                        if els[j]["kind"] == "shape" and els[i]["addr"].startswith(els[j]["addr"] + "."):
                            keep.add(j)
                            break
            picked = [els[i] for i in sorted(keep)][:max_elements]
            res["containers"].append({**{k: v for k, v in c.items() if k != "elements"}, "elements": picked,
                                      "omitted": len(els) - len(picked)})
        else:
            res["containers"].append({**{k: v for k, v in c.items() if k != "elements"}, "elements": [], "omitted": len(els),
                                      "summary_only": True})
    return res


# ---- применение патча ----------------------------------------------------------------------------------------

class Ctx:
    """Состояние применения: лимиты, накопители, проверка «ещё жив ли узел»."""

    def __init__(self, pkg: Package, handler, opts: Opts, known: set[str]):
        self.pkg, self.h, self.opts, self.known = pkg, handler, opts, known
        self.applied: list[dict] = []
        self.comments: list[dict] = []
        self.old_chars = self.new_chars = self.replaced = self.inserted = 0
        self.dead: set = set()
        self.removed_labels: set[str] = set()

    def live(self, e: Elem, i: int) -> None:
        node = e.node
        top = node
        while top is not None:
            if top in self.dead:
                raise PatchRejected("gone", "Элемент уже удалён предыдущей операцией.", i, e.addr)
            parent = top.getparent()
            if parent is None:
                break
            top = parent
        if e.part and top is not self.pkg.xml(e.part):
            raise PatchRejected("gone", "Элемент уже удалён предыдущей операцией.", i, e.addr)

    def kill(self, e: Elem) -> None:
        self.dead.add(e.node)

    def expect(self, e: Elem, op: dict, i: int) -> None:
        exp = op.get("expect")
        if not isinstance(exp, str) or not exp:
            raise PatchRejected("expect_missing", "Нужен отпечаток текущего текста (expect).", i, e.addr)
        if exp != fingerprint(e.text or ""):
            raise PatchRejected("stale", "Устарело: текст уже изменён.", i, e.addr)

    def grow(self, old: int, new: int) -> None:
        self.old_chars += old
        self.new_chars += new
        self.replaced += 1

    def insert(self, n: int) -> None:
        self.inserted += n

    def check_font(self, font, doc_fonts: set[str], i: int, at: str) -> None:
        if not isinstance(font, str) or not font or len(font) > 80:
            raise PatchRejected("font_not_allowed", "Шрифт указан неверно.", i, at)
        allowed = {f.casefold() for f in doc_fonts} | {f.casefold() for f in self.opts.safe_fonts} | {"+mj-lt", "+mn-lt"}
        if font.casefold() not in allowed:
            raise PatchRejected("font_not_allowed", "Шрифт не используется в документе или его теме.", i, at)

    def done(self, i: int, op: dict, at: str, **info) -> None:
        self.removed_labels.update(info.pop("removed_labels", []) or [])
        self.applied.append({"index": i, "op": op["op"], "at": at, **{k: v for k, v in info.items()
                                                                       if k in ("kind", "old", "new", "props")}})

    def finish(self) -> None:
        o = self.opts
        if self.inserted > o.max_total_insert:
            raise PatchRejected("size_explosion", "Слишком много вставленного текста.")
        limit = self.old_chars * (1 + o.max_growth) + o.growth_floor * self.replaced
        if self.new_chars > limit:
            raise PatchRejected("growth", f"Текст вырос больше чем на {int(o.max_growth * 100)} %.")


def _load_patch(patch) -> dict:
    if isinstance(patch, (str, bytes)):
        try:
            patch = json.loads(patch)
        except ValueError as exc:
            raise PatchRejected("bad_patch", "Патч не разобран как JSON.") from exc
    if not isinstance(patch, dict):
        raise PatchRejected("bad_patch", "Патч должен быть объектом.")
    extra = set(patch) - {"v", "base", "note", "ops"}
    if extra:
        raise PatchRejected("bad_patch", "Неизвестные поля патча.")
    if patch.get("v", 1) != 1:
        raise PatchRejected("bad_patch", "Неизвестная версия патча.")
    if not isinstance(patch.get("ops"), list):
        raise PatchRejected("bad_patch", "Нужен список ops.")
    return patch


def _validate_schema(handler, ops: list, opts: Opts) -> None:
    if not ops:
        raise PatchRejected("bad_patch", "Пустой патч.")
    if len(ops) > opts.max_ops:
        raise PatchRejected("too_many_ops", f"Больше {opts.max_ops} операций.")
    schema = handler.SCHEMA
    seen: dict[str, set] = {}
    for i, op in enumerate(ops):
        if not isinstance(op, dict) or not isinstance(op.get("op"), str):
            raise PatchRejected("bad_op", "Операция должна быть объектом с полем op.", i)
        name = op["op"]
        if name not in schema or name not in handler.OPS:
            raise PatchRejected("unknown_op", f"Операция {name[:30]!r} не поддерживается для этого формата.", i)
        req, opt = schema[name]
        keys = set(op) - {"op"}
        if not req <= keys or keys - req - opt:
            raise PatchRejected("bad_op", f"Недопустимые или пропущенные поля в {name}.", i)
        if name in ("replace_text", "delete"):
            at = op.get("at")
            if not isinstance(at, str):
                raise PatchRejected("bad_type", "at должен быть строкой.", i)
            if at in seen.setdefault("text", set()):
                raise PatchRejected("duplicate_address", "Две операции над одним адресом.", i, at)
            seen["text"].add(at)


def apply_patch(data: bytes, name: str, patch, opts=None):
    """Применить патч целиком или не применять вовсе. Возвращает (новые байты, отчёт); иначе PatchRejected."""
    o = _opts(opts)
    patch = _load_patch(patch)
    pkg, h = _open(data, name, o)
    ops = patch["ops"]
    _validate_schema(h, ops, o)
    known = set(o.known_labels) | {l for t in h.all_texts() for l in labels_of(t)}
    ctx = Ctx(pkg, h, o, known)
    try:
        for i, op in enumerate(ops):
            getattr(h, "op_" + op["op"])(ctx, i, op)
        ctx.finish()
        h.finalize(ctx) if hasattr(h, "finalize") else None
        new = pkg.serialize()
        _verify(data, new, name, o, known, h)
    except PatchRejected:
        raise
    except (KeyError, ValueError, TypeError, AttributeError, IndexError) as exc:
        raise PatchRejected("internal", f"Операция не может быть применена ({type(exc).__name__}).") from exc
    report = {"ok": True, "format": h.fmt, "ops": len(ops), "applied": ctx.applied, "comments": ctx.comments,
              "changed_parts": sorted(pkg.dirty | pkg.removed), "removed_labels": sorted(ctx.removed_labels),
              "note": patch.get("note", ""), "warnings": getattr(h, "warnings", [])}
    return new, report


def _verify(old: bytes, new: bytes, name: str, o: Opts, known: set[str], h) -> None:
    """Результат должен открываться теми же библиотеками, а метки в нём — быть только известными."""
    ext = os.path.splitext(name.lower())[1]
    try:
        if ext == ".pptx":
            from pptx import Presentation
            Presentation(io.BytesIO(new))
        elif ext == ".docx":
            import docx
            docx.Document(io.BytesIO(new))
        else:
            import openpyxl
            openpyxl.load_workbook(io.BytesIO(new))
    except Exception as exc:
        raise PatchRejected("reopen_failed", "Результат не открывается: правка не применена.") from exc
    pkg2, h2 = _open(new, name, o)
    pkg1 = Package(old)
    if len(h2.containers) != len(h.containers) or (ext == ".xlsx" and [c["addr"] for c in h2.containers] != [c["addr"] for c in h.containers]):
        raise PatchRejected("structure_changed", "Число слайдов, листов или секций изменилось.")
    for t in h2.all_texts():
        stray = labels_of(t) - known
        if stray:
            raise PatchRejected("label_unknown", "В результате появились неизвестные метки.")
    allowed_gone = {"xl/calcChain.xml"}
    if set(pkg1.raw) - set(pkg2.raw) - allowed_gone:
        raise PatchRejected("structure_changed", "Из файла пропали части.")


# ---- сводка для карточек ----------------------------------------------------------------------------------------

def _clip(s, n=160) -> str:
    s = s or ""
    return s if len(s) <= n else s[: n - 1] + "…"


def _where(addr: str, container: dict | None, el: dict | None) -> str:
    role = {"title": "Заголовок", "body": "Текст", "table cell": "Таблица", "note": "Заметки", "subtitle": "Подзаголовок",
            "footer": "Колонтитул"}.get((el or {}).get("role", ""), "Текст")
    kind = (container or {}).get("kind")
    if kind == "slide":
        return f"Слайд {container['addr'][1:]} · {role}"
    if kind == "sheet":
        return f"Лист «{container.get('title') or container['addr']}» · ячейка {addr.rpartition('!')[2]}"
    return f"{role} {addr}"


_KIND_RU = {"set_style": "оформление текста", "paragraph_style": "стиль абзаца", "col_width": "ширина колонки", "set_layout": "макет слайда",
            "number_format": "формат числа", "set_cell": "значение ячейки"}
_PROP_RU = {"size": "размер", "bold": "жирный", "italic": "курсив", "underline": "подчёркнутый", "color": "цвет", "font": "шрифт", "align": "выравнивание",
            "width": "ширина", "col": "колонка", "layout": "макет", "style": "стиль", "format": "формат", "fill": "заливка"}


def _design_ru(kind: str, props: dict) -> str:
    """Русская строка карточки для операции оформления: «оформление текста: размер 28, жирный: да»."""
    def val(v):
        return "да" if v is True else "нет" if v is False else str(v)
    bits = [f"{_PROP_RU.get(kk, kk)} {val(vv)}" for kk, vv in props.items() if isinstance(vv, (str, int, float, bool))]
    return _KIND_RU.get(kind, "оформление") + (": " + ", ".join(bits) if bits else "")


def diff_summary(old: bytes, new: bytes, patch, name: str, opts=None) -> dict:
    """Русские пункты «было -> стало» для карточек + preview_hint (адреса затронутых элементов в НОВОМ файле)."""
    o = _opts(opts)
    patch = _load_patch(patch)
    old_o, new_o = outline(old, name, o), outline(new, name, o)
    old_map = {e["addr"]: (c, e) for c in old_o["containers"] for e in c["elements"]}
    new_map = {e["addr"]: (c, e) for c in new_o["containers"] for e in c["elements"]}
    new_by_text: dict[str, list[str]] = {}
    for a, (c, e) in new_map.items():
        if e.get("text"):
            new_by_text.setdefault(e["text"], []).append(a)
    items, hint = [], []

    def locate(addr, text=None):
        if text and text in new_by_text:
            same = [a for a in new_by_text[text] if a == addr]
            return same[0] if same else new_by_text[text][0]
        return addr if addr in new_map else None

    for i, op in enumerate(patch["ops"]):
        k = op.get("op")
        at = op.get("at") or op.get("after") or op.get("slide") or ""
        c, e = old_map.get(at, (None, None))
        where = _where(at, c, e)
        if k == "replace_text":
            items.append({"index": i, "op": k, "at": at, "was": e and e.get("text"), "now": op.get("text"),
                          "text": f"{where}: было «{_clip(e and e.get('text'))}» -> стало «{_clip(op.get('text'))}»"})
            hint.append(locate(at, op.get("text")))
        elif k == "insert_paragraph":
            items.append({"index": i, "op": k, "at": at, "was": None, "now": op.get("text"),
                          "text": f"{where}: было (нет абзаца) -> стало «{_clip(op.get('text'))}»"})
            hint.append(locate(at, op.get("text")))
        elif k == "delete":
            items.append({"index": i, "op": k, "at": at, "was": e and e.get("text"), "now": None,
                          "text": f"{where}: было «{_clip((e or {}).get('text') or (e or {}).get('name'))}» -> стало (удалено)"})
        elif k == "comment":
            items.append({"index": i, "op": k, "at": at, "text": f"Комментарий ИИ ({where}): {_clip(op.get('text'), 300)}"})
            hint.append(at if at in new_map else None)
        else:
            props = {kk: vv for kk, vv in op.items() if kk not in ("op", "at", "ats")}
            was = (e or {}).get("style") or (e or {}).get("geom")
            items.append({"index": i, "op": k, "at": at, "was": was, "now": props,
                          "text": f"{where}: {_design_ru(k, props)}"})
            hint.extend(op.get("ats") or [at if at in new_map else None])
    addrs = [a for a in dict.fromkeys(hint) if a]
    conts = list(dict.fromkeys(next((c["addr"] for c in new_o["containers"] for e in c["elements"] if e["addr"] == a), None) for a in addrs))
    return {"items": items, "preview_hint": {"addresses": addrs, "containers": [c for c in conts if c]}}
